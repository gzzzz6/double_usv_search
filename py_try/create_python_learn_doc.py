from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


OUT = Path(r"F:\pythonprojects\py_try\python_learn.docx")


def set_font(run, name="Microsoft YaHei", size=None, color=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    if bold is not None:
        run.bold = bold


def shade(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_heading(doc, text, level=1, color=(31, 78, 121)):
    paragraph = doc.add_heading(level=level)
    run = paragraph.add_run(text)
    set_font(run, size=15 if level == 1 else 13, color=color, bold=True)
    return paragraph


def add_body(doc, text, color=(0, 0, 0), bold=False):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_font(run, color=color, bold=bold)
    return paragraph


def add_code(doc, text):
    paragraph = doc.add_paragraph(style="CodeBlock")
    shade(paragraph, "F2F2F2")
    run = paragraph.add_run(text)
    set_font(run, name="Consolas", size=10, color=(64, 64, 64))
    return paragraph


def add_key(doc, text):
    paragraph = doc.add_paragraph(style="KeyPoint")
    shade(paragraph, "FFF2CC")
    run = paragraph.add_run(text)
    set_font(run, color=(192, 80, 77), bold=True)
    return paragraph


def build_doc():
    doc = Document()

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(11)

    for style_name in ["Title", "Heading 1", "Heading 2"]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    code_style = styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code_style.font.size = Pt(10)

    key_style = styles.add_style("KeyPoint", WD_STYLE_TYPE.PARAGRAPH)
    key_style.font.name = "Microsoft YaHei"
    key_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    key_style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Python 学习笔记：第一阶段 · 第一节")
    set_font(run, size=20, color=(31, 78, 121), bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("主题：Python 程序如何运行、print 输出、变量与基础数据类型")
    set_font(run, size=11, color=(89, 89, 89))

    add_heading(doc, "一、Python 程序如何运行")
    add_body(
        doc,
        "Python 文件通常以 .py 结尾。运行程序时，Python 解释器会从文件第一行开始，按照从上到下的顺序逐行执行代码。",
    )
    add_key(doc, "核心理解：Python 是解释执行的语言。你写好代码后，通常不需要手动编译，直接交给 Python 解释器运行。")
    add_code(doc, 'print("第一行")\nprint("第二行")\nprint("第三行")')
    add_body(doc, "上面的代码会依次输出“第一行”“第二行”“第三行”。代码顺序会直接影响程序运行结果。", color=(89, 89, 89))

    add_heading(doc, "二、print() 输出")
    add_body(doc, "print() 用来把内容显示到屏幕上。它是初学 Python 时最常用的函数之一。")
    add_code(doc, 'print("Hello, Python!")\nprint(123)\nprint(3.14)')
    add_key(doc, "注意：字符串需要加引号，数字通常不需要加引号。")
    add_code(doc, 'print(100)     # 这是数字 100\nprint("100")   # 这是字符串 "100"')
    add_body(doc, "两行代码显示出来都像 100，但它们的类型不同：数字可以直接参与数学运算，字符串本质上是文本。", color=(89, 89, 89))

    add_heading(doc, "三、变量是什么")
    add_body(doc, "变量可以理解为“给一个值起名字”。程序后续可以通过这个名字继续使用这个值。")
    add_code(doc, 'age = 18\nname = "Alice"\nheight = 1.68')
    add_body(doc, '这里 age 表示 18，name 表示 "Alice"，height 表示 1.68。')
    add_key(doc, "核心理解：变量名是给人和程序看的标签，变量中保存的是具体的值。")
    add_code(doc, "price = 10\ncount = 3\ntotal = price * count\nprint(total)")
    add_body(doc, "这段程序会输出 30，因为 total 保存了 price * count 的计算结果。", color=(89, 89, 89))

    add_heading(doc, "四、基础数据类型")
    add_body(doc, "第一阶段先重点掌握 4 种基础类型：int、float、str、bool。")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, text in enumerate(["类型", "含义", "示例", "说明"]):
        paragraph = table.rows[0].cells[i].paragraphs[0]
        run = paragraph.add_run(text)
        set_font(run, color=(255, 255, 255), bold=True)
        shade(paragraph, "4472C4")

    rows = [
        ("int", "整数", "18", "没有小数部分的数字"),
        ("float", "小数", "1.75", "带小数部分的数字"),
        ("str", "字符串", '"Python"', "文本，需要加引号"),
        ("bool", "布尔值", "True / False", "表示真或假，首字母必须大写"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            run = cells[i].paragraphs[0].add_run(text)
            set_font(run)

    add_code(doc, 'age = 18              # int\nheight = 1.75         # float\nname = "Tom"          # str\nis_student = True     # bool')
    add_key(doc, '注意：True 和 False 是布尔值；"True" 和 "False" 是字符串。它们不是同一种东西。')

    add_heading(doc, "五、使用 type() 查看类型")
    add_body(doc, "type() 可以查看一个值或变量的数据类型。")
    add_code(doc, 'print(type(18))\nprint(type(3.14))\nprint(type("hello"))\nprint(type(True))')
    add_body(doc, "输出结果类似 <class 'int'>、<class 'float'>、<class 'str'>、<class 'bool'>。")
    add_key(doc, "学习建议：遇到不确定的数据时，可以先用 type() 看它到底是什么类型。")

    add_heading(doc, "六、字符串拼接与类型转换")
    add_body(doc, "字符串可以用 + 进行拼接。")
    add_code(doc, 'first_name = "Tom"\nlast_name = "Smith"\nfull_name = first_name + " " + last_name\nprint(full_name)')
    add_body(doc, "但是字符串不能直接和数字相加。")
    add_code(doc, 'age = 18\nprint("年龄是" + age)   # 这会报错')
    add_body(doc, "如果确实要拼接，需要先把数字转换成字符串。")
    add_code(doc, 'age = 18\nprint("年龄是" + str(age))')
    add_key(doc, 'str(age) 的作用是把数字 18 转成字符串 "18"。')

    add_heading(doc, "七、f-string：推荐的格式化输出方式")
    add_body(doc, "f-string 是 Python 中非常常用的字符串格式化方式。它可以把变量自然地放进字符串中。")
    add_code(doc, 'name = "张三"\nage = 20\nheight = 1.75\nprint(f"大家好，我叫{name}，今年{age}岁，身高{height}米。")')
    add_key(doc, "推荐习惯：需要把变量放进句子里输出时，优先使用 f-string。")

    add_heading(doc, "八、本节小结")
    for item in [
        "Python 程序从上到下逐行执行。",
        "print() 用于输出内容。",
        "变量是给值起的名字，可以保存数字、文本、布尔值等。",
        "基础类型包括 int、float、str、bool。",
        "type() 可以查看数据类型。",
        "字符串和数字不能直接用 + 拼接，需要类型转换。",
        "f-string 是更推荐的变量输出方式。",
    ]:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_font(run)

    add_heading(doc, "九、复习题")
    for question in [
        'print(100) 和 print("100") 的区别是什么？',
        'True 和 "True" 的区别是什么？',
        '"年龄是" + age 为什么会报错？',
        'f"我的年龄是{age}" 中的大括号有什么作用？',
        "变量 total = price * count 中，total 保存的是什么？",
    ]:
        paragraph = doc.add_paragraph(style="List Number")
        run = paragraph.add_run(question)
        set_font(run)

    doc.add_page_break()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Python 学习笔记：第二阶段 · 第一节")
    set_font(run, size=20, color=(31, 78, 121), bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("主题：核心数据结构入门与列表 list")
    set_font(run, size=11, color=(89, 89, 89))

    add_heading(doc, "一、为什么需要数据结构")
    add_body(
        doc,
        "第一阶段主要处理单个值，比如一个姓名、一个年龄、一个成绩。但真实程序经常需要处理一组数据，例如多个学生、多个成绩、多个商品。",
    )
    add_code(doc, 'name = "张三"\nage = 20\n\nstudents = ["张三", "李四", "王五"]\nscores = [88, 92, 75]')
    add_key(doc, "核心理解：数据结构就是组织数据的方式。列表适合保存一组有顺序的数据。")

    add_heading(doc, "二、Python 常见核心数据结构")
    add_body(doc, "第二阶段会逐步学习 4 种常见结构：list、tuple、dict、set。")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, text in enumerate(["结构", "中文名", "基本写法", "典型用途"]):
        paragraph = table.rows[0].cells[i].paragraphs[0]
        run = paragraph.add_run(text)
        set_font(run, color=(255, 255, 255), bold=True)
        shade(paragraph, "4472C4")

    rows = [
        ("list", "列表", "[1, 2, 3]", "保存一组有顺序、可修改的数据"),
        ("tuple", "元组", "(1, 2, 3)", "保存一组有顺序、通常不修改的数据"),
        ("dict", "字典", '{"name": "张三"}', "保存键值对，适合描述对象信息"),
        ("set", "集合", "{1, 2, 3}", "保存不重复的数据，适合去重"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            run = cells[i].paragraphs[0].add_run(text)
            set_font(run)

    add_heading(doc, "三、列表 list 是什么")
    add_body(doc, "列表用中括号 [] 表示，可以在一个变量中保存多个元素。列表中的元素有固定顺序。")
    add_code(doc, 'names = ["张三", "李四", "王五"]\nscores = [88, 92, 75, 60]')
    add_body(doc, "列表中的每一项叫做元素。names 这个列表有 3 个元素，scores 这个列表有 4 个元素。")
    add_key(doc, "核心理解：列表不是一个值，而是一组值的容器。")

    add_heading(doc, "四、下标 index")
    add_body(doc, "列表中的每个元素都有位置编号，这个编号叫下标。Python 的下标从 0 开始。")
    add_code(doc, 'names = ["张三", "李四", "王五"]\n\nprint(names[0])  # 张三\nprint(names[1])  # 李四\nprint(names[2])  # 王五')
    add_key(doc, "注意：第一个元素的下标是 0，不是 1。")
    add_body(doc, "如果访问不存在的下标，例如 names[3]，程序会报错，因为这个列表只有 0、1、2 三个有效下标。", color=(89, 89, 89))

    add_heading(doc, "五、修改列表元素")
    add_body(doc, "列表是可变的。可以通过下标找到某个位置，然后给它重新赋值。")
    add_code(doc, 'names = ["张三", "李四", "王五"]\nnames[1] = "赵六"\nprint(names)')
    add_body(doc, '输出结果是 ["张三", "赵六", "王五"]。原来的 "李四" 被替换成了 "赵六"。')

    add_heading(doc, "六、添加元素 append()")
    add_body(doc, "append() 用来在列表末尾添加一个新元素。")
    add_code(doc, 'fruits = ["苹果", "香蕉", "橘子"]\nfruits.append("西瓜")\nprint(fruits)')
    add_key(doc, "append() 会修改原来的列表，而不是创建一个全新的列表。")

    add_heading(doc, "七、删除元素 remove() 和 pop()")
    add_body(doc, "remove() 按元素的值删除，pop() 按下标删除。")
    add_code(doc, 'fruits = ["苹果", "香蕉", "橘子"]\nfruits.remove("香蕉")\nprint(fruits)')
    add_code(doc, 'names = ["张三", "李四", "王五"]\nnames.pop(1)\nprint(names)')
    add_body(doc, "如果 pop() 不写下标，默认删除最后一个元素。")
    add_key(doc, "选择方法时先想清楚：你是知道要删除的值，还是知道它的位置。")

    add_heading(doc, "八、列表长度 len()")
    add_body(doc, "len() 可以得到列表中元素的数量。")
    add_code(doc, 'names = ["张三", "李四", "王五"]\nprint(len(names))  # 3')
    add_body(doc, "len() 常用于统计数量，也经常和 sum() 一起计算平均值。")

    add_heading(doc, "九、遍历列表 for")
    add_body(doc, "遍历就是把列表中的元素一个一个取出来处理。for 循环是遍历列表最常用的方式。")
    add_code(doc, 'names = ["张三", "李四", "王五"]\n\nfor name in names:\n    print(name)')
    add_key(doc, "核心理解：for name in names 的意思是：每次从 names 里取出一个元素，临时叫它 name。")

    add_heading(doc, "十、带编号遍历 enumerate()")
    add_body(doc, "如果遍历时既需要元素，也需要它是第几个，可以使用 enumerate()。")
    add_code(doc, 'scores = [88, 92, 75]\n\nfor i, score in enumerate(scores, start=1):\n    print(f"第{i}个成绩是：{score}")')
    add_body(doc, "start=1 表示编号从 1 开始，更适合输出给用户看。")

    add_heading(doc, "十一、列表统计")
    add_body(doc, "对于数字列表，可以使用 sum()、max()、min()、len() 做常见统计。")
    add_code(doc, 'scores = [88, 92, 75, 60, 100]\n\ntotal = sum(scores)\nhighest = max(scores)\nlowest = min(scores)\naverage = total / len(scores)\n\nprint(f"总分：{total}")\nprint(f"最高分：{highest}")\nprint(f"最低分：{lowest}")\nprint(f"平均分：{average}")')
    add_key(doc, "平均值的基本公式：总和 / 数量，也就是 sum(scores) / len(scores)。")

    add_heading(doc, "十二、循环累加与计数")
    add_body(doc, "除了使用内置函数，也可以用循环自己统计。这样能帮助理解程序是如何一步步计算结果的。")
    add_code(doc, 'scores = [88, 45, 92, 59, 76, 100]\ncount = 0\n\nfor score in scores:\n    if score >= 60:\n        count = count + 1\n\nprint(f"及格人数：{count}")')
    add_body(doc, "count 初始为 0。每遇到一个及格成绩，就让 count 增加 1。循环结束后，count 中保存的就是及格人数。")

    add_heading(doc, "十三、本节小结")
    for item in [
        "列表 list 用于保存一组有顺序的数据。",
        "列表下标从 0 开始。",
        "可以通过 names[0] 访问元素，也可以通过 names[1] = 新值 修改元素。",
        "append() 用于添加元素，remove() 和 pop() 用于删除元素。",
        "len() 可以统计列表长度。",
        "for 循环可以遍历列表。",
        "enumerate() 可以在遍历时同时得到编号和元素。",
        "sum()、max()、min() 常用于数字列表统计。",
    ]:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_font(run)

    add_heading(doc, "十四、复习题")
    for question in [
        "为什么 names[0] 表示第一个元素？",
        "append() 和 remove() 分别有什么作用？",
        "remove() 和 pop() 的区别是什么？",
        "len(scores) 得到的是什么？",
        "for score in scores 中，score 每次代表什么？",
        "enumerate(scores, start=1) 为什么适合输出“第几个成绩”？",
        "如何计算成绩列表的平均分？",
    ]:
        paragraph = doc.add_paragraph(style="List Number")
        run = paragraph.add_run(question)
        set_font(run)

    doc.add_page_break()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Python 学习笔记：第二阶段 · 第二节")
    set_font(run, size=20, color=(31, 78, 121), bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("主题：字典 dict")
    set_font(run, size=11, color=(89, 89, 89))

    add_heading(doc, "一、字典 dict 是什么")
    add_body(
        doc,
        "字典用来保存“带标签的数据”。列表主要靠位置访问元素，字典主要靠键 key 访问对应的值 value。",
    )
    add_code(doc, 'student = {\n    "name": "张三",\n    "age": 20,\n    "score": 88\n}')
    add_body(doc, '这里 "name"、"age"、"score" 是键；"张三"、20、88 是对应的值。')
    add_key(doc, "核心理解：字典保存的是 key: value 这样的键值对。通过 key 可以快速找到 value。")

    add_heading(doc, "二、列表和字典的区别")
    add_body(doc, "列表适合保存一组同类数据，字典适合描述一个对象的多个属性。")
    add_code(doc, 'scores = [88, 92, 75]\n\nstudent = {\n    "name": "张三",\n    "age": 20,\n    "score": 88\n}')
    add_body(doc, "scores 是一组成绩；student 是一个学生的信息。")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, text in enumerate(["结构", "访问方式", "适合场景", "示例"]):
        paragraph = table.rows[0].cells[i].paragraphs[0]
        run = paragraph.add_run(text)
        set_font(run, color=(255, 255, 255), bold=True)
        shade(paragraph, "4472C4")

    rows = [
        ("list", "通过下标访问", "一组有顺序的数据", "scores[0]"),
        ("dict", "通过键访问", "一个对象的多项信息", 'student["name"]'),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            run = cells[i].paragraphs[0].add_run(text)
            set_font(run)

    add_heading(doc, "三、读取字典中的值")
    add_body(doc, "读取字典时，用中括号加 key。")
    add_code(doc, 'student = {\n    "name": "张三",\n    "age": 20,\n    "score": 88\n}\n\nprint(student["name"])\nprint(student["age"])\nprint(student["score"])')
    add_key(doc, "注意：字典不能像列表那样用 student[0] 访问。字典要用 key 访问。")

    add_heading(doc, "四、修改字典中的值")
    add_body(doc, "如果 key 已经存在，对它重新赋值就会修改原来的 value。")
    add_code(doc, 'student = {\n    "name": "张三",\n    "age": 20,\n    "score": 88\n}\n\nstudent["score"] = 95\nprint(student)')
    add_body(doc, '执行后，"score" 对应的值会从 88 改成 95。')

    add_heading(doc, "五、添加新的键值对")
    add_body(doc, "如果 key 不存在，对它赋值就会添加一组新的键值对。")
    add_code(doc, 'student = {\n    "name": "张三",\n    "age": 20\n}\n\nstudent["score"] = 88\nprint(student)')
    add_key(doc, "同样是 student[key] = value：key 存在时是修改，key 不存在时是添加。")

    add_heading(doc, "六、删除键值对 pop()")
    add_body(doc, "pop() 可以根据 key 删除一组键值对。")
    add_code(doc, 'student = {\n    "name": "张三",\n    "age": 20,\n    "score": 88\n}\n\nstudent.pop("age")\nprint(student)')
    add_body(doc, '执行后，"age": 20 会被删除。')

    add_heading(doc, "七、判断 key 是否存在")
    add_body(doc, "使用 in 可以判断某个 key 是否在字典中。")
    add_code(doc, 'student = {\n    "name": "张三",\n    "age": 20\n}\n\nprint("name" in student)\nprint("score" in student)')
    add_body(doc, '输出结果是 True 和 False，因为 "name" 存在，"score" 不存在。')
    add_code(doc, 'if "score" in student:\n    print(student["score"])\nelse:\n    print("没有成绩")')

    add_heading(doc, "八、get() 安全读取")
    add_body(doc, "如果直接读取不存在的 key，会出现 KeyError。get() 可以更安全地读取字典。")
    add_code(doc, 'student = {\n    "name": "张三",\n    "age": 20\n}\n\nprint(student.get("score"))\nprint(student.get("phone", "暂无手机号"))')
    add_key(doc, 'get(key, 默认值) 的含义是：如果 key 存在，返回对应的值；如果不存在，返回默认值。')

    add_heading(doc, "九、遍历字典")
    add_body(doc, "遍历字典常见有三种方式：遍历 key、遍历 value、同时遍历 key 和 value。")
    add_code(doc, 'student = {\n    "name": "张三",\n    "age": 20,\n    "score": 88\n}\n\nfor key in student:\n    print(key)')
    add_code(doc, 'for value in student.values():\n    print(value)')
    add_code(doc, 'for key, value in student.items():\n    print(f"{key}: {value}")')
    add_key(doc, "最常用的是 items()，因为它可以同时拿到 key 和 value。")

    add_heading(doc, "十、列表里面放字典")
    add_body(doc, "真实程序中，经常用“列表 + 字典”保存一组对象。列表表示多个对象，字典表示每个对象的详细信息。")
    add_code(doc, 'students = [\n    {"name": "张三", "score": 88},\n    {"name": "李四", "score": 92},\n    {"name": "王五", "score": 75}\n]')
    add_body(doc, "students 是一个列表，列表里的每个元素都是一个字典。每个字典代表一个学生。")
    add_code(doc, 'for student in students:\n    print(f\'{student["name"]}的成绩是：{student["score"]}\')')

    add_heading(doc, "十一、多个学生统计平均分")
    add_body(doc, "统计多个学生平均分时，可以遍历学生列表，把每个学生字典里的 score 累加起来。")
    add_code(doc, 'students = [\n    {"name": "张三", "score": 88},\n    {"name": "李四", "score": 92},\n    {"name": "王五", "score": 75}\n]\n\ntotal = 0\nfor student in students:\n    total = total + student["score"]\n\naverage = total / len(students)\nprint(f"平均分：{average}")')
    add_key(doc, "建议使用 len(students)，不要把人数写死成 3。这样学生数量变化时，代码仍然正确。")

    add_heading(doc, "十二、常见错误")
    add_body(doc, "1. 把数字写成字符串。年龄应该写成 20，而不是 \"20\"，除非你明确只想把它当文本。")
    add_code(doc, '"age": 20      # 推荐\n"age": "20"    # 这是字符串，不适合计算')
    add_body(doc, "2. 直接读取不存在的 key。")
    add_code(doc, 'print(student["phone"])  # 如果 phone 不存在，会报 KeyError')
    add_body(doc, "更稳妥的写法：")
    add_code(doc, 'print(student.get("phone", "暂无手机号"))')
    add_body(doc, "3. f-string 中引号混乱。")
    add_code(doc, 'print(f\'{student["name"]}的成绩是：{student["score"]}\')')

    add_heading(doc, "十三、本节小结")
    for item in [
        "字典 dict 用于保存 key: value 键值对。",
        "列表靠下标访问，字典靠 key 访问。",
        "student['name'] 可以读取 name 对应的值。",
        "key 存在时赋值是修改，key 不存在时赋值是添加。",
        "pop(key) 可以删除键值对。",
        "in 可以判断 key 是否存在。",
        "get() 可以安全读取不存在的 key，并提供默认值。",
        "items() 可以同时遍历 key 和 value。",
        "列表中可以放多个字典，用来表示多个对象。",
    ]:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_font(run)

    add_heading(doc, "十四、复习题")
    for question in [
        "字典中的 key 和 value 分别是什么意思？",
        "student['score'] 的作用是什么？",
        "student['city'] = '北京' 在 city 不存在时会发生什么？",
        "pop('age') 的作用是什么？",
        "student.get('phone', '暂无手机号') 的含义是什么？",
        "for key, value in student.items() 中，key 和 value 每次分别代表什么？",
        "为什么多个学生适合用“列表里面放字典”的结构？",
    ]:
        paragraph = doc.add_paragraph(style="List Number")
        run = paragraph.add_run(question)
        set_font(run)

    doc.add_page_break()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Python 学习笔记：第二阶段 · 第三节")
    set_font(run, size=20, color=(31, 78, 121), bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("主题：元组 tuple 和集合 set")
    set_font(run, size=11, color=(89, 89, 89))

    add_heading(doc, "一、元组 tuple 是什么")
    add_body(doc, "元组和列表类似，都可以保存一组有顺序的数据。区别是：列表可以修改，元组不能修改。")
    add_code(doc, 'names_list = ["张三", "李四", "王五"]\nnames_tuple = ("张三", "李四", "王五")')
    add_body(doc, "列表使用中括号 []，元组使用小括号 ()。")
    add_key(doc, "核心理解：tuple 是有顺序、不可修改的数据容器。")

    add_heading(doc, "二、访问元组元素")
    add_body(doc, "元组和列表一样，可以通过下标访问元素，下标也从 0 开始。")
    add_code(doc, 'point = (10, 20)\n\nprint(point[0])\nprint(point[1])')
    add_body(doc, "point[0] 是 10，point[1] 是 20。")

    add_heading(doc, "三、元组不可修改")
    add_body(doc, "元组创建后，里面的元素不能被重新赋值。")
    add_code(doc, 'names = ("张三", "李四", "王五")\nnames[1] = "赵六"  # 这会报错')
    add_body(doc, "这类错误通常是 TypeError，因为 tuple 不支持元素修改。")
    add_key(doc, "当一组数据不希望被程序意外改动时，可以考虑用元组。")

    add_heading(doc, "四、元组的常见使用场景")
    add_body(doc, "元组常用于表达固定不变的一组值，例如坐标、颜色、固定配置等。")
    add_code(doc, 'point = (10, 20)\ncolor = (255, 0, 0)\nweekdays = ("周一", "周二", "周三", "周四", "周五")')

    add_heading(doc, "五、单元素元组")
    add_body(doc, "只有一个元素的元组必须写逗号。这个点很容易出错。")
    add_code(doc, 'a = (10)\nb = (10,)\n\nprint(type(a))\nprint(type(b))')
    add_body(doc, "a 的类型是 int，b 的类型才是 tuple。")
    add_key(doc, "判断单元素元组的关键不是小括号，而是逗号。")

    add_heading(doc, "六、遍历元组")
    add_body(doc, "元组也可以使用 for 循环遍历。")
    add_code(doc, 'weekdays = ("周一", "周二", "周三", "周四", "周五")\n\nfor day in weekdays:\n    print(day)')

    add_heading(doc, "七、集合 set 是什么")
    add_body(doc, "集合用来保存一组不重复的数据。集合使用大括号 {}，但它不是字典，因为集合里面没有 key: value。")
    add_code(doc, 'numbers = {1, 2, 3, 4}\nwords = {"Python", "Java", "C++"}')
    add_key(doc, "核心理解：set 是无顺序、不重复的数据容器。")

    add_heading(doc, "八、集合自动去重")
    add_body(doc, "集合最常见的用途是去重。重复元素放进集合后，只会保留一份。")
    add_code(doc, 'numbers = {1, 2, 2, 3, 3, 4}\nprint(numbers)')
    add_body(doc, "输出结果会类似 {1, 2, 3, 4}。重复的 2 和 3 被自动去掉了。")
    add_code(doc, 'names = ["张三", "李四", "张三", "王五", "李四"]\nunique_names = set(names)\nprint(unique_names)')
    add_key(doc, "注意：集合无顺序，所以去重后的输出顺序不一定和原列表一致。")

    add_heading(doc, "九、集合添加和删除")
    add_body(doc, "add() 用于添加元素，remove() 和 discard() 都可以删除元素。")
    add_code(doc, 'numbers = {1, 2, 3}\nnumbers.add(4)\nnumbers.remove(2)\nprint(numbers)')
    add_body(doc, "remove() 删除不存在的元素会报错，discard() 删除不存在的元素不会报错。")
    add_code(doc, 'numbers = {1, 2, 3}\nnumbers.discard(100)\nprint(numbers)')
    add_key(doc, "不确定元素是否存在时，优先使用 discard()。")

    add_heading(doc, "十、集合运算")
    add_body(doc, "集合支持交集、并集、差集等数学集合运算。")
    add_code(doc, 'a = {1, 2, 3, 4}\nb = {3, 4, 5, 6}\n\nprint(a & b)  # 交集：两个集合都有的元素\nprint(a | b)  # 并集：两个集合合并后的全部元素\nprint(a - b)  # 差集：a 有，b 没有的元素')

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, text in enumerate(["运算", "符号", "含义", "示例结果"]):
        paragraph = table.rows[0].cells[i].paragraphs[0]
        run = paragraph.add_run(text)
        set_font(run, color=(255, 255, 255), bold=True)
        shade(paragraph, "4472C4")

    rows = [
        ("交集", "&", "两个集合都有的元素", "{3, 4}"),
        ("并集", "|", "两个集合合并后的全部元素", "{1, 2, 3, 4, 5, 6}"),
        ("差集", "-", "左边集合有、右边集合没有的元素", "{1, 2}"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            run = cells[i].paragraphs[0].add_run(text)
            set_font(run)

    add_heading(doc, "十一、空集合")
    add_body(doc, "空集合不能写成 {}，因为 {} 表示空字典。")
    add_code(doc, 'empty_dict = {}\nempty_set = set()\n\nprint(type(empty_dict))\nprint(type(empty_set))')
    add_key(doc, "创建空集合必须使用 set()。")

    add_heading(doc, "十二、list、tuple、dict、set 对比")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for i, text in enumerate(["结构", "中文名", "是否有顺序", "是否可修改", "主要特点"]):
        paragraph = table.rows[0].cells[i].paragraphs[0]
        run = paragraph.add_run(text)
        set_font(run, color=(255, 255, 255), bold=True)
        shade(paragraph, "4472C4")

    rows = [
        ("list", "列表", "是", "是", "按下标访问，适合一组数据"),
        ("tuple", "元组", "是", "否", "适合固定不变的数据"),
        ("dict", "字典", "按 key 管理", "是", "key-value 键值对"),
        ("set", "集合", "否", "是", "不重复，适合去重和集合运算"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            run = cells[i].paragraphs[0].add_run(text)
            set_font(run)

    add_heading(doc, "十三、本节小结")
    for item in [
        "元组 tuple 和列表类似，但元组不可修改。",
        "元组可以通过下标访问，也可以用 for 遍历。",
        "单元素元组必须写成 (10,)，逗号不能省略。",
        "集合 set 用于保存不重复的数据。",
        "集合是无顺序的，不能依赖输出顺序。",
        "set(list) 可以快速去重。",
        "add() 添加集合元素，remove() 和 discard() 删除集合元素。",
        "discard() 删除不存在的元素不会报错。",
        "集合支持交集 &、并集 |、差集 -。",
        "空集合必须写 set()，{} 是空字典。",
    ]:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_font(run)

    add_heading(doc, "十四、复习题")
    for question in [
        "tuple 和 list 最大的区别是什么？",
        "为什么 (10) 不是元组，而 (10,) 是元组？",
        "什么时候适合用 tuple？",
        "set 的两个核心特点是什么？",
        "为什么 set(names) 可以去重？",
        "remove() 和 discard() 删除集合元素时有什么区别？",
        "a & b、a | b、a - b 分别表示什么？",
        "{} 和 set() 分别创建什么对象？",
    ]:
        paragraph = doc.add_paragraph(style="List Number")
        run = paragraph.add_run(question)
        set_font(run)

    doc.add_page_break()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Python 学习笔记：第二阶段 · 第四节")
    set_font(run, size=20, color=(31, 78, 121), bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("主题：切片、排序、列表推导式")
    set_font(run, size=11, color=(89, 89, 89))

    add_heading(doc, "一、切片 slice 是什么")
    add_body(doc, "切片用于从列表、字符串、元组等序列中取出一部分数据。")
    add_code(doc, "序列[start:end]")
    add_body(doc, "含义是：从 start 开始取，到 end 之前结束。包含 start，不包含 end。")
    add_code(doc, "nums = [10, 20, 30, 40, 50]\nprint(nums[1:4])")
    add_body(doc, "输出结果是 [20, 30, 40]，因为取到了下标 1、2、3，没有取下标 4。")
    add_key(doc, "核心理解：切片左闭右开，也就是包含左边界，不包含右边界。")

    add_heading(doc, "二、切片常见写法")
    add_body(doc, "切片可以省略 start 或 end。省略 start 表示从开头开始，省略 end 表示取到最后。")
    add_code(doc, "nums = [10, 20, 30, 40, 50, 60]\n\nprint(nums[:3])   # [10, 20, 30]\nprint(nums[3:])   # [40, 50, 60]\nprint(nums[:])    # 复制整个列表")
    add_body(doc, "负数下标可以从右往左数。")
    add_code(doc, "print(nums[-1])   # 最后一个元素\nprint(nums[-2])   # 倒数第二个元素")
    add_body(doc, "倒序列表可以使用 [::-1]。")
    add_code(doc, "print(nums[::-1])")

    add_heading(doc, "三、切片步长")
    add_body(doc, "完整切片语法是：")
    add_code(doc, "序列[start:end:step]")
    add_body(doc, "step 表示步长，也就是每隔几个位置取一次。")
    add_code(doc, "nums = [10, 20, 30, 40, 50, 60]\n\nprint(nums[::2])   # [10, 30, 50]\nprint(nums[1::2])  # [20, 40, 60]")
    add_key(doc, "step 为 2 时，表示每隔 2 个下标取一个元素。step 为 -1 时，常用于倒序。")

    add_heading(doc, "四、sorted() 排序")
    add_body(doc, "sorted() 会返回一个排序后的新列表，不会修改原列表。")
    add_code(doc, "scores = [88, 60, 100, 75, 92]\n\nnew_scores = sorted(scores)\n\nprint(scores)\nprint(new_scores)")
    add_body(doc, "scores 仍然是原来的顺序，new_scores 是升序排序后的新列表。")
    add_key(doc, "当你想保留原列表时，优先使用 sorted()。")

    add_heading(doc, "五、sort() 排序")
    add_body(doc, "列表自身的 .sort() 方法会直接修改原列表，不会返回一个新列表。")
    add_code(doc, "scores = [88, 60, 100, 75, 92]\n\nscores.sort()\nprint(scores)")
    add_body(doc, "执行后，scores 自己变成 [60, 75, 88, 92, 100]。")
    add_key(doc, "sort() 是原地排序，会改变原列表。")

    add_heading(doc, "六、降序排序")
    add_body(doc, "无论是 sorted() 还是 sort()，都可以通过 reverse=True 实现降序。")
    add_code(doc, "scores = [88, 60, 100, 75, 92]\n\nprint(sorted(scores, reverse=True))\n\nscores.sort(reverse=True)\nprint(scores)")
    add_body(doc, "reverse=True 表示反向排序，也就是从大到小。")

    add_heading(doc, "七、按字典字段排序")
    add_body(doc, "列表中放字典时，经常需要按字典里的某个字段排序，例如按学生成绩排序。")
    add_code(doc, 'students = [\n    {"name": "张三", "score": 88},\n    {"name": "李四", "score": 92},\n    {"name": "王五", "score": 75}\n]\n\nstudents_sorted = sorted(students, key=lambda student: student["score"], reverse=True)\nprint(students_sorted)')
    add_body(doc, '这里 key=lambda student: student["score"] 的意思是：排序时，每个学生用 score 作为排序依据。')
    add_key(doc, "lambda 现在可以先理解为一个临时小函数，用来告诉 sorted() 按什么规则排序。")

    add_heading(doc, "八、列表推导式")
    add_body(doc, "列表推导式用于快速生成新列表。它常常可以替代“创建空列表 + for 循环 + append”的写法。")
    add_code(doc, "nums = [1, 2, 3, 4, 5]\n\nsquares = []\nfor n in nums:\n    squares.append(n * n)\n\nprint(squares)")
    add_body(doc, "上面的普通循环可以改写成列表推导式：")
    add_code(doc, "nums = [1, 2, 3, 4, 5]\nsquares = [n * n for n in nums]\nprint(squares)")
    add_key(doc, "核心格式：[新元素表达式 for 临时变量 in 原列表]")

    add_heading(doc, "九、带条件的列表推导式")
    add_body(doc, "列表推导式可以加 if 条件，用来筛选元素。")
    add_code(doc, "nums = [1, 2, 3, 4, 5, 6]\n\neven_nums = [n for n in nums if n % 2 == 0]\nprint(even_nums)")
    add_body(doc, "意思是：从 nums 中依次取出 n，如果 n 是偶数，就把 n 放进新列表。")
    add_code(doc, "scores = [88, 45, 92, 59, 76, 100]\nqualified = [score for score in scores if score >= 60]\nprint(qualified)")
    add_key(doc, "带条件的核心格式：[新元素表达式 for 临时变量 in 原列表 if 条件]")

    add_heading(doc, "十、普通循环和列表推导式对比")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, text in enumerate(["目标", "普通循环", "列表推导式"]):
        paragraph = table.rows[0].cells[i].paragraphs[0]
        run = paragraph.add_run(text)
        set_font(run, color=(255, 255, 255), bold=True)
        shade(paragraph, "4472C4")

    rows = [
        ("生成平方数", "for n in nums: append(n*n)", "[n * n for n in nums]"),
        ("筛选偶数", "if n % 2 == 0: append(n)", "[n for n in nums if n % 2 == 0]"),
        ("筛选及格成绩", "if score >= 60: append(score)", "[score for score in scores if score >= 60]"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            run = cells[i].paragraphs[0].add_run(text)
            set_font(run)

    add_heading(doc, "十一、常见错误")
    add_body(doc, "1. 切片右边界不包含。")
    add_code(doc, "nums = [10, 20, 30, 40, 50]\nprint(nums[1:4])  # 取到 40，不取 50")
    add_body(doc, "2. 混淆 sorted() 和 sort()。")
    add_code(doc, "new_scores = sorted(scores)  # 返回新列表，不改原列表\nscores.sort()               # 修改原列表")
    add_body(doc, "3. 题目要求列表推导式时，不要再写普通 for append。")
    add_code(doc, "qualified = [score for score in scores if score >= 60]")

    add_heading(doc, "十二、本节小结")
    for item in [
        "切片用于从序列中取出一部分数据。",
        "切片语法是 sequence[start:end:step]。",
        "切片包含 start，不包含 end。",
        "nums[:3] 从开头取到下标 3 之前，nums[3:] 从下标 3 取到最后。",
        "nums[::-1] 可以倒序。",
        "sorted() 返回新列表，不修改原列表。",
        "list.sort() 会直接修改原列表。",
        "reverse=True 表示降序排序。",
        "排序字典列表时，可以使用 key=lambda student: student['score']。",
        "列表推导式可以简洁生成新列表。",
        "带条件的列表推导式可以筛选数据。",
    ]:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_font(run)

    add_heading(doc, "十三、复习题")
    for question in [
        "nums[1:4] 会取到哪些下标？",
        "为什么切片说“包含 start，不包含 end”？",
        "nums[:3]、nums[3:]、nums[::-1] 分别是什么意思？",
        "sorted(scores) 和 scores.sort() 的区别是什么？",
        "reverse=True 的作用是什么？",
        "sorted(students, key=lambda student: student['score']) 中 key 的作用是什么？",
        "[n * n for n in nums] 这句代码的含义是什么？",
        "[score for score in scores if score >= 60] 这句代码如何理解？",
    ]:
        paragraph = doc.add_paragraph(style="List Number")
        run = paragraph.add_run(question)
        set_font(run)

    doc.add_page_break()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Python 学习笔记：第三阶段 · 第一节")
    set_font(run, size=20, color=(31, 78, 121), bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("主题：函数基础")
    set_font(run, size=11, color=(89, 89, 89))

    add_heading(doc, "一、函数是什么")
    add_body(doc, "函数可以理解为“给一段代码起名字”。把一段经常使用的逻辑封装成函数后，需要时直接调用它。")
    add_code(doc, "def add(a, b):\n    return a + b\n\nresult = add(10, 20)\nprint(result)")
    add_body(doc, "这里 add 是函数名，a 和 b 是参数，return a + b 会把计算结果返回给调用者。")
    add_key(doc, "核心理解：函数用于封装重复逻辑，让代码更清晰、更容易复用。")

    add_heading(doc, "二、函数基本语法")
    add_body(doc, "定义函数使用 def，函数体必须缩进。")
    add_code(doc, "def 函数名(参数):\n    函数体\n    return 返回值")
    add_code(doc, 'def say_hello():\n    print("你好，Python")\n\nsay_hello()')
    add_body(doc, "调用函数时必须写括号。say_hello() 会执行函数；只写 say_hello 不会执行函数。")
    add_key(doc, "缩进属于函数体的一部分。缩进结束后，就回到函数外部。")

    add_heading(doc, "三、无参数函数")
    add_body(doc, "如果函数不需要外部输入，可以不写参数。")
    add_code(doc, 'def say_hello():\n    print("你好，Python")\n\nsay_hello()')
    add_body(doc, "这个函数每次调用都会输出固定内容。")

    add_heading(doc, "四、参数是什么")
    add_body(doc, "参数是函数接收的输入。调用函数时传入不同参数，函数就可以处理不同数据。")
    add_code(doc, 'def greet(name):\n    print(f"你好，{name}")\n\ngreet("张三")\ngreet("李四")')
    add_body(doc, '第一次调用时，name 是 "张三"；第二次调用时，name 是 "李四"。')
    add_key(doc, "参数让函数从“只能做固定事情”变成“可以处理不同数据”。")

    add_heading(doc, "五、多个参数")
    add_body(doc, "函数可以接收多个参数，参数之间用逗号分隔。")
    add_code(doc, "def add(a, b):\n    return a + b\n\nprint(add(10, 20))\nprint(add(3, 5))")
    add_body(doc, "调用函数时，实参会按顺序传给形参。")

    add_heading(doc, "六、return 返回值")
    add_body(doc, "return 用来把函数的计算结果返回给调用者。返回后的结果可以继续保存、计算或传给其他函数。")
    add_code(doc, "def add(a, b):\n    return a + b\n\nresult = add(10, 20)\nprint(result)")
    add_key(doc, "print() 是显示给人看；return 是返回给程序继续用。")
    add_body(doc, "如果函数没有 return，它默认返回 None。")
    add_code(doc, "def add(a, b):\n    print(a + b)\n\nresult = add(10, 20)\nprint(result)  # None")

    add_heading(doc, "七、函数处理列表")
    add_body(doc, "函数的参数也可以是列表。比如计算平均分、统计及格人数。")
    add_code(doc, "def calc_average(scores):\n    return sum(scores) / len(scores)\n\nscores = [88, 92, 75]\naverage = calc_average(scores)\nprint(average)")
    add_code(doc, "def count_pass(scores):\n    count = 0\n    for score in scores:\n        if score >= 60:\n            count = count + 1\n    return count\n\nscores = [88, 45, 92, 59, 76]\nprint(count_pass(scores))")
    add_body(doc, "这类函数的优点是：以后换一组成绩，也可以直接复用函数。")

    add_heading(doc, "八、默认参数")
    add_body(doc, "参数可以设置默认值。调用时如果没有传这个参数，就使用默认值。")
    add_code(doc, 'def introduce(name, city="北京"):\n    print(f"我叫{name}，来自{city}")\n\nintroduce("张三")\nintroduce("李四", "上海")')
    add_body(doc, '第一次调用没有传 city，所以使用默认值 "北京"；第二次传了 "上海"，所以使用传入值。')

    add_heading(doc, "九、关键字参数")
    add_body(doc, "调用函数时可以指定参数名，这种方式叫关键字参数。使用关键字参数时，参数顺序可以打乱。")
    add_code(doc, 'def show_info(name, age, score):\n    print(f"姓名：{name}，年龄：{age}，成绩：{score}")\n\nshow_info(score=88, name="张三", age=20)')
    add_key(doc, "关键字参数能让调用更清楚，也能减少参数顺序写错的风险。")

    add_heading(doc, "十、局部变量")
    add_body(doc, "函数内部定义的变量通常只在函数内部有效，这种变量叫局部变量。")
    add_code(doc, "def test():\n    x = 10\n    print(x)\n\ntest()\nprint(x)  # 函数外部无法直接访问 x")
    add_body(doc, "如果需要把函数内部的结果拿到外面使用，应该用 return 返回。")

    add_heading(doc, "十一、常见错误")
    add_body(doc, "1. 定义了函数但忘记调用。")
    add_code(doc, 'def say_hello():\n    print("你好，Python")\n\nsay_hello   # 不会执行\nsay_hello() # 会执行')
    add_body(doc, "2. 把 print() 当成 return。")
    add_code(doc, "def add(a, b):\n    print(a + b)\n\nresult = add(10, 20)  # result 是 None")
    add_body(doc, "3. 函数名拼写不清晰。函数名应该表达它做什么，比如 calc_average 表示计算平均值。")
    add_code(doc, "def calc_average(scores):\n    return sum(scores) / len(scores)")

    add_heading(doc, "十二、本节小结")
    for item in [
        "函数用于封装一段代码，让代码可以复用。",
        "定义函数使用 def，调用函数需要写函数名和括号。",
        "参数是函数的输入。",
        "函数可以有多个参数。",
        "return 会把结果返回给程序继续使用。",
        "print() 只是把结果显示出来，不等于返回值。",
        "没有 return 的函数默认返回 None。",
        "默认参数可以在调用时省略。",
        "关键字参数可以指定参数名，调用顺序可以打乱。",
        "函数内部定义的变量通常是局部变量。",
    ]:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_font(run)

    add_heading(doc, "十三、复习题")
    for question in [
        "函数的作用是什么？",
        "def say_hello(): 中 def 表示什么？",
        "参数的作用是什么？",
        "print() 和 return 的区别是什么？",
        "没有 return 的函数默认返回什么？",
        "默认参数 introduce(name, city='北京') 如何理解？",
        "关键字参数 show_info(score=88, name='张三', age=20) 有什么好处？",
        "为什么函数内部的局部变量不能直接在函数外部使用？",
    ]:
        paragraph = doc.add_paragraph(style="List Number")
        run = paragraph.add_run(question)
        set_font(run)

    doc.add_page_break()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Python 学习笔记：第三阶段 · 第二节")
    set_font(run, size=20, color=(31, 78, 121), bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("主题：函数进阶：作用域、参数类型、模块化思维")
    set_font(run, size=11, color=(89, 89, 89))

    add_heading(doc, "一、作用域是什么")
    add_body(doc, "作用域就是变量可以被访问的范围。变量定义在什么位置，决定了它能在哪些地方使用。")
    add_code(doc, "x = 100\n\ndef test():\n    y = 200\n    print(x)\n    print(y)\n\ntest()")
    add_body(doc, "这里 x 定义在函数外，是全局变量；y 定义在函数内，是局部变量。")
    add_key(doc, "核心理解：函数内部可以读取外面的全局变量，但函数外面不能直接访问函数里的局部变量。")

    add_heading(doc, "二、全局变量和局部变量")
    add_body(doc, "定义在函数外的变量通常叫全局变量，定义在函数内的变量通常叫局部变量。")
    add_code(doc, 'school = "清华大学"\n\ndef show_school():\n    print(school)\n\ndef test_local():\n    local_num = 100\n    print(local_num)\n\nshow_school()\ntest_local()')
    add_body(doc, "show_school() 可以读取全局变量 school；test_local() 中的 local_num 只在这个函数内部有效。")
    add_key(doc, "写函数时，尽量让函数内部自己处理局部变量；需要交给外部继续使用的结果，用 return 返回。")

    add_heading(doc, "三、同名变量与遮蔽")
    add_body(doc, "如果函数内部和函数外部有同名变量，函数内部会优先使用自己的局部变量。")
    add_code(doc, 'name = "张三"\n\ndef show_name():\n    name = "李四"\n    print(name)\n\nshow_name()\nprint(name)')
    add_body(doc, '运行后会先输出 "李四"，再输出 "张三"。说明函数内部的 name 不会直接修改外部的同名变量。')
    add_key(doc, "可以先把它理解成：函数内部的同名变量会“遮住”外部变量。")

    add_heading(doc, "四、位置参数")
    add_body(doc, "位置参数就是按顺序传值。调用时第一个实参给第一个形参，第二个实参给第二个形参。")
    add_code(doc, "def multiply(a, b):\n    return a * b\n\nprint(multiply(3, 4))")
    add_body(doc, "这里 3 会传给 a，4 会传给 b。参数顺序写反，结果也可能变化。")

    add_heading(doc, "五、默认参数")
    add_body(doc, "默认参数就是参数先带一个默认值。调用时如果没传，就使用默认值。")
    add_code(doc, "def power(num, exponent=2):\n    return num ** exponent\n\nprint(power(3))\nprint(power(2, 3))")
    add_body(doc, "power(3) 相当于计算 3 的平方；power(2, 3) 则把 exponent 改成 3，计算 2 的 3 次方。")
    add_key(doc, "默认参数很适合“多数时候用默认值，少数时候再改”的场景。")

    add_heading(doc, "六、关键字参数")
    add_body(doc, "关键字参数是在调用函数时明确写出参数名。这样代码更清楚，参数顺序也可以变化。")
    add_code(doc, 'def show_student(name, age, city):\n    print(f"{name}，{age}，{city}")\n\nshow_student(age=18, city="北京", name="张三")')
    add_body(doc, "调用时虽然顺序变了，但因为写了参数名，Python 仍然能正确对应。")
    add_key(doc, "当参数较多时，关键字参数通常更容易读，也更不容易写错。")

    add_heading(doc, "七、一个函数只做一件事")
    add_body(doc, "写函数时，尽量让每个函数只负责一个明确任务。这样代码更容易理解、修改和复用。")
    add_code(doc, "def calc_total(scores):\n    return sum(scores)\n\ndef calc_average(scores):\n    return sum(scores) / len(scores)\n\ndef count_pass(scores):\n    count = 0\n    for score in scores:\n        if score >= 60:\n            count = count + 1\n    return count")
    add_body(doc, "这里三个函数分别负责总分、平均分、及格人数，每个函数的职责都很清楚。")
    add_key(doc, "不要把很多不相干的逻辑都塞进同一个函数里。函数短一点、职责单一点，后面更好维护。")

    add_heading(doc, "八、模块化思维")
    add_body(doc, "模块化思维就是把一个大问题拆成多个小功能，每个小功能写成一个函数。")
    add_code(doc, "scores = [88, 45, 92, 76, 59]\n\nprint(calc_total(scores))\nprint(calc_average(scores))\nprint(count_pass(scores))")
    add_body(doc, "“成绩统计”本来是一个整体任务，但拆成多个小函数后，代码会更整齐，也更容易逐个检查。")
    add_key(doc, "以后写稍微复杂一点的程序时，先想“这个任务能拆成几个小函数”。")

    add_heading(doc, "九、函数之间可以配合")
    add_body(doc, "函数不只是单独使用，它们还可以互相配合。一个函数可以调用另一个函数。")
    add_code(doc, "def calc_total(scores):\n    return sum(scores)\n\ndef calc_average(scores):\n    total = calc_total(scores)\n    return total / len(scores)\n\nscores = [88, 92, 75]\nprint(calc_average(scores))")
    add_body(doc, "calc_average() 先调用 calc_total() 得到总分，再继续计算平均分。这样可以减少重复代码。")
    add_key(doc, "当一个小功能已经写好时，后面的函数可以直接复用它，而不是再重新写一遍。")

    add_heading(doc, "十、常见错误")
    add_body(doc, "1. 在函数外直接访问局部变量。")
    add_code(doc, "def test_local():\n    local_num = 100\n    print(local_num)\n\nprint(local_num)  # 报错")
    add_body(doc, "2. 忘记 return，导致函数结果拿不到外面继续使用。")
    add_code(doc, "def multiply(a, b):\n    a * b\n\nresult = multiply(3, 4)\nprint(result)  # None")
    add_body(doc, "3. 关键字参数把参数名写错。")
    add_code(doc, 'def show_student(name, age, city):\n    print(name, age, city)\n\nshow_student(nam="张三", age=18, city="北京")')

    add_heading(doc, "十一、本节小结")
    for item in [
        "作用域表示变量可以被访问的范围。",
        "函数外定义的通常是全局变量，函数内定义的通常是局部变量。",
        "函数内部可以读取全局变量，但函数外不能直接访问局部变量。",
        "如果内外有同名变量，函数内部会优先使用自己的局部变量。",
        "位置参数按顺序传值。",
        "默认参数在没有传值时会使用默认值。",
        "关键字参数可以明确指定参数名。",
        "一个函数尽量只做一件事。",
        "模块化思维就是把大任务拆成多个小函数。",
        "函数之间可以互相调用，减少重复代码。",
    ]:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_font(run)

    add_heading(doc, "十二、复习题")
    for question in [
        "什么是作用域？",
        "全局变量和局部变量有什么区别？",
        "为什么函数外不能直接访问 local_num 这样的局部变量？",
        "如果函数内外都有同名变量，程序会优先使用哪一个？",
        "位置参数和关键字参数的区别是什么？",
        "power(num, exponent=2) 中的 =2 表示什么？",
        "为什么说一个函数尽量只做一件事？",
        "“模块化思维”可以怎样帮助你组织代码？",
    ]:
        paragraph = doc.add_paragraph(style="List Number")
        run = paragraph.add_run(question)
        set_font(run)

    doc.add_page_break()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Python 学习笔记：第三阶段 · 第三节")
    set_font(run, size=20, color=(31, 78, 121), bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("主题：模块与 import")
    set_font(run, size=11, color=(89, 89, 89))

    add_heading(doc, "一、什么是模块")
    add_body(doc, "在 Python 里，一个 .py 文件就可以看成一个模块。模块可以理解为“装着函数和代码的文件”。")
    add_code(doc, "def add(a, b):\n    return a + b")
    add_body(doc, "如果上面的代码写在 tools.py 中，那么 tools.py 就是一个模块，里面定义了 add() 函数。")
    add_key(doc, "核心理解：模块就是用来组织和复用代码的 Python 文件。")

    add_heading(doc, "二、为什么要使用模块")
    add_body(doc, "当程序变大时，不适合把所有内容都写在一个文件里。可以把不同功能拆到不同模块中。")
    add_code(doc, "# 成绩统计函数放一个文件\n# 字符串处理函数放一个文件\n# 主程序再放一个文件")
    add_body(doc, "这样做以后，代码会更清楚，也更容易维护和复用。")
    add_key(doc, "这就是前面学过的“模块化思维”在文件层面的应用。")

    add_heading(doc, "三、import 是什么")
    add_body(doc, "import 的作用是：把别的模块里的内容拿过来使用。")
    add_code(doc, "import tools\n\nprint(tools.add(10, 20))")
    add_body(doc, "这里 import tools 表示导入 tools 模块，tools.add(10, 20) 表示调用模块中的 add() 函数。")

    add_heading(doc, "四、为什么调用时要写模块名")
    add_body(doc, "写成 tools.add() 的好处是来源清楚。看到代码时，你能知道这个函数来自哪个模块。")
    add_code(doc, "import tools\n\nresult = tools.add(3, 5)\nprint(result)")
    add_body(doc, "如果多个模块里都有同名函数，带上模块名就不容易混乱。")
    add_key(doc, "对初学者来说，import 模块名 再用 模块名.函数名 的写法更容易理解。")

    add_heading(doc, "五、from ... import ...")
    add_body(doc, "除了 import tools 这种写法，还可以只导入模块里的某个函数。")
    add_code(doc, "from tools import add\n\nprint(add(10, 20))")
    add_body(doc, "这样调用时可以直接写 add()，不需要再写 tools.。")
    add_key(doc, "这种写法更短，但初学阶段更推荐先熟悉带模块名前缀的写法。")

    add_heading(doc, "六、两种导入方式对比")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, text in enumerate(["写法", "示例", "特点"]):
        paragraph = table.rows[0].cells[i].paragraphs[0]
        run = paragraph.add_run(text)
        set_font(run, color=(255, 255, 255), bold=True)
        shade(paragraph, "4472C4")

    rows = [
        ("import 模块", "import tools", "来源清楚，调用时写 tools.add()"),
        ("from 模块 import 函数", "from tools import add", "写法更短，可以直接写 add()"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            run = cells[i].paragraphs[0].add_run(text)
            set_font(run)

    add_heading(doc, "七、标准库模块")
    add_body(doc, "Python 自带很多现成模块，叫标准库。它们不需要自己写，直接 import 就可以使用。")
    add_code(doc, "import math\n\nprint(math.sqrt(25))\nprint(math.pi)")
    add_body(doc, "math.sqrt(25) 用来求平方根，math.pi 表示圆周率。")
    add_code(doc, "import random\n\nprint(random.randint(1, 10))")
    add_body(doc, "random.randint(1, 10) 表示随机生成 1 到 10 之间的整数。")
    add_key(doc, "标准库就是 Python 已经帮你准备好的工具箱。")

    add_heading(doc, "八、自己写模块")
    add_body(doc, "你也可以自己创建模块，把常用函数放进去，然后在别的文件中导入使用。")
    add_code(doc, 'def add(a, b):\n    return a + b\n\ndef subtract(a, b):\n    return a - b\n\ndef say_hello(name):\n    print(f"你好，{name}")')
    add_body(doc, "比如把这些函数写在 my_tools.py 中，再在主程序文件里 import my_tools，就能重复使用这些函数。")
    add_key(doc, "自己写模块，是把“写函数”进一步升级成“写可复用的工具文件”。")

    add_heading(doc, "九、if __name__ == '__main__'")
    add_body(doc, "以后你会经常看到 if __name__ == '__main__'。它常用来区分：当前文件是被直接运行，还是被别的文件导入。")
    add_code(doc, 'def add(a, b):\n    return a + b\n\nif __name__ == "__main__":\n    print("主程序正在运行")\n    print(add(10, 20))')
    add_body(doc, "直接运行这个文件时，这段代码会执行；如果这个文件只是被 import，这里的代码通常不会自动执行。")
    add_key(doc, "现阶段先记住它的作用：放测试代码、主程序入口代码。")

    add_heading(doc, "十、常见错误")
    add_body(doc, "1. 模块名写错。")
    add_code(doc, "import my_tool   # 实际文件名是 my_tools.py")
    add_body(doc, "2. 使用 import tools 后，调用时忘了写模块名。")
    add_code(doc, "import tools\n\nadd(10, 20)        # 错误\ntools.add(10, 20)  # 正确")
    add_body(doc, "3. 对没有 return 的函数再套一层 print，导致输出 None。")
    add_code(doc, 'print(my_tools.say_hello("张三"))   # 会多输出一个 None')

    add_heading(doc, "十一、本节小结")
    for item in [
        "一个 .py 文件就可以看成一个模块。",
        "模块可以帮助我们拆分、组织和复用代码。",
        "import 的作用是导入模块。",
        "import tools 后，调用函数要写 tools.add() 这样的形式。",
        "from tools import add 可以直接使用 add()。",
        "Python 自带很多标准库模块，例如 math 和 random。",
        "自己也可以编写模块，比如 my_tools.py。",
        "if __name__ == '__main__' 常用来放主程序入口或测试代码。",
    ]:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_font(run)

    add_heading(doc, "十二、复习题")
    for question in [
        "什么是模块？",
        "为什么程序变大后要把代码拆到不同模块里？",
        "import tools 和 from tools import add 有什么区别？",
        "为什么 import tools 后，通常要写 tools.add()？",
        "math.sqrt(25) 的作用是什么？",
        "为什么 print(my_tools.say_hello('张三')) 会多输出一个 None？",
        "if __name__ == '__main__' 一般用来做什么？",
    ]:
        paragraph = doc.add_paragraph(style="List Number")
        run = paragraph.add_run(question)
        set_font(run)

    doc.add_page_break()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Python 学习笔记：第四阶段 · 第一节")
    set_font(run, size=20, color=(31, 78, 121), bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("主题：文件读写基础")
    set_font(run, size=11, color=(89, 89, 89))

    add_heading(doc, "一、什么是文件")
    add_body(doc, "文件就是保存在电脑里的数据。文本文件、Python 文件、Word 文件都属于文件。")
    add_code(doc, "study_note.txt\nhello.py\npython_learn.docx")
    add_body(doc, "程序运行结束后，变量里的内容通常会消失；写进文件里的内容可以保留下来，下次还能继续读取。")
    add_key(doc, "核心理解：文件让程序的数据可以长期保存，而不是只存在于运行过程里。")

    add_heading(doc, "二、为什么要学文件读写")
    add_body(doc, "很多程序都要把内容保存到磁盘中，或者从磁盘读取已有数据。")
    add_code(doc, "# 保存学习记录\n# 保存学生成绩\n# 读取配置文件\n# 把程序结果写到文本里")
    add_body(doc, "所以文件读写是 Python 非常基础也非常常用的能力。")

    add_heading(doc, "三、open() 是什么")
    add_body(doc, "open() 用来打开文件。打开后，程序就可以对文件进行读取、写入或追加。")
    add_code(doc, 'file = open("test.txt", "r", encoding="utf-8")')
    add_body(doc, "这里 test.txt 是文件名，r 是读取模式，encoding='utf-8' 表示按 UTF-8 编码处理文件。")
    add_key(doc, "处理中文文件时，通常都要写 encoding='utf-8'。")

    add_heading(doc, "四、推荐写法：with open(...)")
    add_body(doc, "更推荐使用 with open(...) 这种写法。它会在用完文件后自动关闭文件。")
    add_code(doc, 'with open("test.txt", "r", encoding="utf-8") as file:\n    content = file.read()\n    print(content)')
    add_body(doc, "这种写法更安全，也更符合日常开发习惯。以后优先使用它。")
    add_key(doc, "你可以先把 with open(...) 理解成“打开文件，用完自动收尾”。")

    add_heading(doc, "五、三种最常用模式")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, text in enumerate(["模式", "作用", "特点"]):
        paragraph = table.rows[0].cells[i].paragraphs[0]
        run = paragraph.add_run(text)
        set_font(run, color=(255, 255, 255), bold=True)
        shade(paragraph, "4472C4")

    rows = [
        ("r", "读取", "文件必须已经存在"),
        ("w", "写入", "不存在就创建，已存在会覆盖原内容"),
        ("a", "追加", "不存在就创建，已存在就在末尾继续写"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            run = cells[i].paragraphs[0].add_run(text)
            set_font(run)

    add_heading(doc, "六、读取整个文件：read()")
    add_body(doc, "read() 会把整个文件内容一次性读出来。")
    add_code(doc, 'with open("study_note.txt", "r", encoding="utf-8") as file:\n    content = file.read()\n    print(content)')
    add_body(doc, "适合文件内容不大、想一次性查看完整文本时使用。")

    add_heading(doc, "七、写入文件：write()")
    add_body(doc, "write() 用来向文件写入内容。常和 w 或 a 模式一起使用。")
    add_code(doc, 'with open("study_note.txt", "w", encoding="utf-8") as file:\n    file.write("Python 学习开始了\\n")\n    file.write("我正在学习文件读写\\n")')
    add_body(doc, "如果使用 w 模式，原文件内容会被覆盖；如果文件不存在，则会自动创建。")
    add_key(doc, "w 是重写，不是接着写。想保留原内容时，不要误用 w。")

    add_heading(doc, "八、追加文件内容：a")
    add_body(doc, "如果想在原文件末尾继续写，而不是覆盖原内容，就用 a 模式。")
    add_code(doc, 'with open("study_note.txt", "a", encoding="utf-8") as file:\n    file.write("今天学习了 open 和 with\\n")')
    add_body(doc, "a 可以理解为“append”，也就是追加。")

    add_heading(doc, "九、按行读取：readlines()")
    add_body(doc, "readlines() 会按行读取文件，并返回一个列表。列表中的每个元素就是文件的一行。")
    add_code(doc, 'with open("scores.txt", "r", encoding="utf-8") as file:\n    lines = file.readlines()\n    print(lines)\n    print(len(lines))')
    add_body(doc, "如果文件有 3 行，那么 lines 就会是一个长度为 3 的列表。")
    add_key(doc, "readlines() 常用于“统计行数”或“逐行处理文本”。")

    add_heading(doc, "十、处理换行符")
    add_body(doc, "文本文件中的每一行末尾通常带有换行符 \\n。直接打印时，常会多出空行。")
    add_code(doc, 'for line in lines:\n    print(line.strip())')
    add_body(doc, "strip() 可以去掉首尾空白字符，包括行尾的 \\n。")
    add_code(doc, 'file.write("第一行\\n")\nfile.write("第二行\\n")')
    add_body(doc, "\\n 表示换行。写文件时，如果想让内容分成多行，通常要手动加上它。")

    add_heading(doc, "十一、路径变量的作用")
    add_body(doc, "写文件时，先把路径保存到变量里，会让代码更清楚，也更容易统一修改。")
    add_code(doc, 'file_path = r"F:\\pythonprojects\\py_try\\study_note.txt"\n\nwith open(file_path, "w", encoding="utf-8") as file:\n    file.write("Python 学习开始了\\n")')
    add_body(doc, "这样后面再次读取、写入、追加时，都可以重复使用同一个 file_path 变量。")
    add_key(doc, "先定义路径变量，再统一使用，是一个很好的习惯。")

    add_heading(doc, "十二、常见错误")
    add_body(doc, "1. 用 r 读取一个不存在的文件。")
    add_code(doc, 'with open("abc.txt", "r", encoding="utf-8") as file:\n    print(file.read())')
    add_body(doc, "2. 把 w 当成追加，结果把原内容覆盖了。")
    add_code(doc, 'with open("study_note.txt", "w", encoding="utf-8") as file:\n    file.write("新内容")')
    add_body(doc, "3. 直接 print(line) 导致每行之间多空一行。")
    add_code(doc, 'for line in lines:\n    print(line)         # 常会多空行\n    print(line.strip()) # 更常用')

    add_heading(doc, "十三、本节小结")
    for item in [
        "文件用于把数据保存到磁盘中。",
        "open() 用来打开文件。",
        "更推荐使用 with open(...) 的写法。",
        "r 表示读取，w 表示写入，a 表示追加。",
        "w 会覆盖原内容，a 会在末尾继续写。",
        "read() 可以一次性读取整个文件。",
        "readlines() 会按行读取，并返回列表。",
        "写多行内容时，通常要手动写 \\n。",
        "逐行输出时，常用 strip() 去掉换行符。",
        "先定义文件路径变量，再统一使用，会让代码更清楚。",
    ]:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_font(run)

    add_heading(doc, "十四、复习题")
    for question in [
        "为什么程序需要文件，而不能只用变量？",
        "open('test.txt', 'r', encoding='utf-8') 中 r 和 encoding 分别表示什么？",
        "为什么更推荐使用 with open(...)？",
        "w 和 a 的区别是什么？",
        "read() 和 readlines() 的区别是什么？",
        "为什么逐行 print(line) 时常会多出空行？",
        "strip() 在这里的作用是什么？",
        "为什么把文件路径先保存到变量里是个好习惯？",
    ]:
        paragraph = doc.add_paragraph(style="List Number")
        run = paragraph.add_run(question)
        set_font(run)

    section = doc.sections[0]
    section.top_margin = Pt(54)
    section.bottom_margin = Pt(54)
    section.left_margin = Pt(60)
    section.right_margin = Pt(60)

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
