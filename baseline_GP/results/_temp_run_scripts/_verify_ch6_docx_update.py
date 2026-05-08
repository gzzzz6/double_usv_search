from pathlib import Path
from docx import Document
import zipfile

path = Path(r'F:\pythonprojects\my_report.docx')
doc = Document(path)
texts = [p.text for p in doc.paragraphs]

print('doc:', path)
print('paragraph_count:', len(texts))
print('table_count:', len(doc.tables))
print('\nCH6_PARAGRAPHS')
for i, t in enumerate(texts):
    s = t.strip()
    if s.startswith(('6 ', '6.', '表6-')):
        print(f'{i}: {s[:160]}')

with zipfile.ZipFile(path) as z:
    document_xml = z.read('word/document.xml').decode('utf-8', errors='replace')

print('\nTOKEN_CHECK')
for token in ['表6-1', '表6-2', '表6-15', 'simple_ring_v1', 'paper_simple_ring_mainline_20260505', '后续工作需要', '\ufffd']:
    print(token, token in document_xml)

print('\nBAD_MARKERS')
bad_markers = ['锛', '鈥', '乱码', '????', '\ufffd']
found = False
joined = '\n'.join(texts)
for marker in bad_markers:
    if marker in joined or marker in document_xml:
        print('BAD_MARKER_FOUND:', marker)
        found = True
if not found:
    print('none')

print('\nTABLE_SAMPLES')
for idx in [1, 2, 5, 10, 15]:
    if idx < len(doc.tables):
        tbl = doc.tables[idx]
        print(f'TABLE {idx}: rows={len(tbl.rows)} cols={len(tbl.columns)}')
        for r in range(min(2, len(tbl.rows))):
            print(' | '.join(c.text.replace('\n','/') for c in tbl.rows[r].cells[:min(6, len(tbl.columns))]))
