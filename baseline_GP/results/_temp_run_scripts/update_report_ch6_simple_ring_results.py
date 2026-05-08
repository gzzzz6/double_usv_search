from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table


PROJECT_ROOT = Path(r"F:\pythonprojects")
REPORT_PATH = PROJECT_ROOT / "my_report.docx"
TABLE_ROOT = (
    PROJECT_ROOT
    / "baseline_GP"
    / "results"
    / "paper_simple_ring_mainline_20260505"
    / "paper_tables"
)

MAP_LABELS = {
    "ALL": "总体",
    "open_water": "open_water",
    "harbor_cove": "harbor_cove",
    "peninsula_passage": "peninsula_passage",
}
MAP_ORDER = ("ALL", "open_water", "harbor_cove", "peninsula_passage")


def read_csv(name: str) -> list[dict[str, str]]:
    path = TABLE_ROOT / name
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def as_float(value: object) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.upper() == "N/A":
        return None
    return float(text)


def fmt_num(value: object, digits: int = 1) -> str:
    v = as_float(value)
    if v is None:
        return "N/A"
    return f"{v:.{digits}f}"


def fmt_pct(value: object, digits: int = 1) -> str:
    v = as_float(value)
    if v is None:
        return "N/A"
    return f"{100.0 * v:.{digits}f}"


def fmt_pp(value: object, digits: int = 1) -> str:
    v = as_float(value)
    if v is None:
        return "N/A"
    return f"{100.0 * v:+.{digits}f}"


def fmt_delta(value: object, digits: int = 1) -> str:
    v = as_float(value)
    if v is None:
        return "N/A"
    return f"{v:+.{digits}f}"


def fmt_delta_found(value: object) -> str:
    v = as_float(value)
    if v is None:
        return "N/A"
    return f"{v:+.2f}"


def row_for(
    rows: list[dict[str, str]],
    map_kind: str,
    *,
    assignment_mode: str | None = None,
    clue_mode: str | None = None,
) -> dict[str, str]:
    for row in rows:
        if row.get("map_kind") != map_kind:
            continue
        if assignment_mode is not None and row.get("assignment_mode") != assignment_mode:
            continue
        if clue_mode is not None and row.get("clue_acquisition_mode") != clue_mode:
            continue
        return row
    raise KeyError((map_kind, assignment_mode, clue_mode))


def set_run_font(run, size: float = 10.5, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def set_paragraph_text(paragraph, text: str, *, size: float = 10.5, bold: bool = False) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)


def new_paragraph_before(doc: Document, before_elm, text: str = "", style: str | None = None):
    p = OxmlElement("w:p")
    before_elm.addprevious(p)
    paragraph = doc.paragraphs[-1].__class__(p, doc)
    if style:
        paragraph.style = style
    if text:
        set_paragraph_text(paragraph, text)
    return paragraph


def add_paragraph_before(
    doc: Document,
    before_elm,
    text: str,
    *,
    style: str | None = "Normal",
    align: WD_ALIGN_PARAGRAPH | None = None,
):
    paragraph = new_paragraph_before(doc, before_elm, text, style)
    if align is not None:
        paragraph.alignment = align
    return paragraph


def add_heading_before(doc: Document, before_elm, text: str, style: str):
    paragraph = add_paragraph_before(doc, before_elm, text, style=style)
    return paragraph


def add_caption_before(doc: Document, before_elm, text: str):
    paragraph = add_paragraph_before(doc, before_elm, text, style="表目录项")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return paragraph


def set_cell_text(cell, text: str, *, bold: bool = False, size: float = 9.0) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in kwargs[edge].items():
            element.set(qn(f"w:{key}"), str(value))


def format_three_line_table(table: Table) -> None:
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                top={"val": "nil"},
                left={"val": "nil"},
                bottom={"val": "nil"},
                right={"val": "nil"},
            )
    if not table.rows:
        return
    top = {"val": "single", "sz": "12", "color": "000000"}
    thin = {"val": "single", "sz": "6", "color": "000000"}
    for cell in table.rows[0].cells:
        set_cell_border(cell, top=top, bottom=thin)
    for cell in table.rows[-1].cells:
        set_cell_border(cell, bottom=top)


def add_table_before(
    doc: Document,
    before_elm,
    headers: list[str],
    rows: list[list[str]],
    *,
    font_size: float = 8.5,
) -> Table:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = True
    for j, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], header, bold=True, size=font_size)
    for row_values in rows:
        row = table.add_row()
        for j, value in enumerate(row_values):
            set_cell_text(row.cells[j], value, size=font_size)
    format_three_line_table(table)
    table._tbl.getparent().remove(table._tbl)
    before_elm.addprevious(table._tbl)
    return table


def delete_chapter_6(doc: Document):
    body = doc.element.body
    children = list(body.iterchildren())
    start = None
    end = None
    for i, child in enumerate(children):
        if child.tag == qn("w:p"):
            text = "".join(node.text or "" for node in child.iter(qn("w:t"))).strip()
            if text == "6 实验设计与结果分析":
                start = i
            elif start is not None and text == "第七章 总结与展望":
                end = i
                break
    if start is None or end is None:
        raise RuntimeError("Could not locate chapter 6 range.")
    before_elm = children[end]
    for child in children[start:end]:
        body.remove(child)
    return before_elm


def single_table_rows(rows: list[dict[str, str]]) -> list[list[str]]:
    out: list[list[str]] = []
    for mk in MAP_ORDER:
        r = row_for(rows, mk)
        out.append(
            [
                MAP_LABELS[mk],
                r["n"],
                r["n_first_detected"],
                r["n_all_found"],
                fmt_num(r["time_to_first_detection_mean"]),
                fmt_num(r["time_to_all_found_mean"]),
                fmt_pct(r["success_all_found_mean"]),
                fmt_pct(r["detection_rate_mean"]),
                fmt_num(r["found_count_mean"], 2),
                fmt_pct(r["known_free_observation_ratio_final_mean"]),
                fmt_num(r["path_length_mean"]),
            ]
        )
    return out


def assignment_rows(rows: list[dict[str, str]]) -> list[list[str]]:
    out: list[list[str]] = []
    for mode in ("independent", "coordinated"):
        for mk in MAP_ORDER:
            r = row_for(rows, mk, assignment_mode=mode)
            out.append(
                [
                    MAP_LABELS[mk],
                    "independent" if mode == "independent" else "coordinated",
                    r["n"],
                    r["n_first_detected"],
                    r["n_all_found"],
                    fmt_num(r["time_to_first_detection_mean"]),
                    fmt_num(r["time_to_all_found_mean"]),
                    fmt_pct(r["success_all_found_mean"]),
                    fmt_pct(r["detection_rate_mean"]),
                    fmt_num(r["found_count_mean"], 2),
                    fmt_pct(r["known_free_observation_ratio_final_mean"]),
                    fmt_pct(r["duplicate_viewpoint_ratio_mean"]),
                    fmt_pct(r["cross_region_assignment_ratio_mean"]),
                    fmt_num(r["wait_count_total_mean"]),
                ]
            )
    return out


def coord_delta_rows(rows: list[dict[str, str]]) -> list[list[str]]:
    out: list[list[str]] = []
    for mk in MAP_ORDER:
        r = row_for(rows, mk)
        out.append(
            [
                MAP_LABELS[mk],
                fmt_delta(r["delta_time_to_first_detection_mean"]),
                fmt_delta(r["delta_time_to_all_found_mean"]),
                fmt_pp(r["delta_success_all_found_mean"]),
                fmt_pp(r["delta_detection_rate_mean"]),
                fmt_delta_found(r["delta_found_count_mean"]),
                fmt_pp(r["delta_duplicate_viewpoint_ratio_mean"]),
                fmt_pp(r["delta_cross_region_assignment_ratio_mean"]),
                fmt_delta(r["delta_wait_count_total_mean"]),
            ]
        )
    return out


def anomaly_main_rows(rows: list[dict[str, str]]) -> list[list[str]]:
    out: list[list[str]] = []
    for clue in ("ucb", "anomaly_upper_tail"):
        for mk in MAP_ORDER:
            r = row_for(rows, mk, clue_mode=clue)
            out.append(
                [
                    MAP_LABELS[mk],
                    "UCB" if clue == "ucb" else "anomaly_upper_tail",
                    r["n"],
                    r["n_first_detected"],
                    r["n_all_found"],
                    fmt_num(r["time_to_first_detection_mean"]),
                    fmt_num(r["time_to_all_found_mean"]),
                    fmt_pct(r["success_all_found_mean"]),
                    fmt_pct(r["detection_rate_mean"]),
                    fmt_num(r["found_count_mean"], 2),
                    fmt_pct(r["known_free_observation_ratio_final_mean"]),
                ]
            )
    return out


def anomaly_delta_rows(rows: list[dict[str, str]]) -> list[list[str]]:
    out: list[list[str]] = []
    for mk in MAP_ORDER:
        r = row_for(rows, mk)
        out.append(
            [
                MAP_LABELS[mk],
                fmt_delta(r["delta_time_to_first_detection_mean"]),
                fmt_delta(r["delta_time_to_all_found_mean"]),
                fmt_pp(r["delta_success_all_found_mean"]),
                fmt_pp(r["delta_detection_rate_mean"]),
                fmt_delta_found(r["delta_found_count_mean"]),
                fmt_pp(r["delta_known_free_observation_ratio_final_mean"]),
            ]
        )
    return out


def main() -> None:
    doc = Document(REPORT_PATH)
    before_elm = delete_chapter_6(doc)

    t62_static = read_csv("table_6_2_static_single_ucb.csv")
    t62_random = read_csv("table_6_2_random_walk_single_ucb.csv")
    t63_static = read_csv("table_6_3_static_two_usv_ucb_assignment.csv")
    t63_random = read_csv("table_6_3_random_walk_two_usv_ucb_assignment.csv")
    t63_delta_static = read_csv("table_6_3_static_two_usv_coord_minus_independent.csv")
    t63_delta_random = read_csv("table_6_3_random_walk_two_usv_coord_minus_independent.csv")
    t64_single_static = read_csv("table_6_4_single_static_ucb_vs_anomaly.csv")
    t64_single_random = read_csv("table_6_4_single_random_walk_ucb_vs_anomaly.csv")
    t64_two_static = read_csv("table_6_4_two_coordinated_static_ucb_vs_anomaly.csv")
    t64_two_random = read_csv("table_6_4_two_coordinated_random_walk_ucb_vs_anomaly.csv")
    t64_d_single_static = read_csv("table_6_4_single_static_delta_anomaly_minus_ucb.csv")
    t64_d_single_random = read_csv("table_6_4_single_random_walk_delta_anomaly_minus_ucb.csv")
    t64_d_two_static = read_csv("table_6_4_two_coordinated_static_delta_anomaly_minus_ucb.csv")
    t64_d_two_random = read_csv("table_6_4_two_coordinated_random_walk_delta_anomaly_minus_ucb.csv")

    single_headers = [
        "地图",
        "n",
        "n_first",
        "n_all",
        "T_first",
        "T_all",
        "success（%）",
        "detection（%）",
        "found",
        "obs（%）",
        "path",
    ]
    assignment_headers = [
        "地图",
        "方法",
        "n",
        "n_first",
        "n_all",
        "T_first",
        "T_all",
        "success（%）",
        "detection（%）",
        "found",
        "obs（%）",
        "duplicate（%）",
        "cross（%）",
        "wait",
    ]
    coord_delta_headers = [
        "地图",
        "ΔT_first",
        "ΔT_all",
        "Δsuccess（百分点）",
        "Δdetection（百分点）",
        "Δfound",
        "Δduplicate（百分点）",
        "Δcross（百分点）",
        "Δwait",
    ]
    anomaly_headers = [
        "地图",
        "采集方式",
        "n",
        "n_first",
        "n_all",
        "T_first",
        "T_all",
        "success（%）",
        "detection（%）",
        "found",
        "obs（%）",
    ]
    anomaly_delta_headers = [
        "地图",
        "ΔT_first",
        "ΔT_all",
        "Δsuccess（百分点）",
        "Δdetection（百分点）",
        "Δfound",
        "Δobs（百分点）",
    ]

    # Chapter 6
    add_heading_before(doc, before_elm, "6 实验设计与结果分析", "一级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "本章围绕已知静态障碍地图下的海上可疑目标搜索任务，设计用于验证单 USV 信息驱动搜索基线、双 USV 协同搜索机制以及 anomaly-aware acquisition 作用的实验方案。实验分析以第 3 章建立的场模型、第 4 章单艇路径规划方法和第 5 章双艇协同决策方法为基础，所有实验均采用当前主线 simple_ring_v1 候选视点生成方式和 soft_clearance_astar_v1 路径安全生成层。",
    )
    add_paragraph_before(
        doc,
        before_elm,
        "本章结果均来自 F:\\pythonprojects\\baseline_GP\\results\\paper_simple_ring_mainline_20260505 目录下的真实批量运行数据。单艇实验与双艇实验采用不同地图尺寸，因此第 6.2 节与第 6.3 节不进行单艇和双艇的直接同表数值比较；第 6.4 节则在相同系统结构内比较 UCB 与 anomaly_upper_tail 采集函数。",
    )

    # 6.1
    add_heading_before(doc, before_elm, "6.1 实验设置", "二级标题")
    add_heading_before(doc, before_elm, "6.1.1 仿真环境与地图设置", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "实验在二维栅格化海域仿真环境中进行。地图由自由水域与静态障碍物组成，USV 只能在自由栅格中运动，障碍物在整个实验过程中保持不变。本文研究对象为 known static map 条件，因此规划器在路径生成阶段可以使用完整静态障碍地图，但目标位置和线索状态对规划器仍然未知。",
    )
    add_paragraph_before(
        doc,
        before_elm,
        "实验地图选取 open_water、harbor_cove 和 peninsula_passage 三类典型场景。open_water 用于观察少障碍条件下的信息驱动搜索行为；harbor_cove 具有凹形港湾结构，能够检验避障路径生成与热点观测之间的耦合；peninsula_passage 包含半岛和通道结构，适合分析狭窄通路、绕行代价和多艇分工对搜索效率的影响。",
    )
    add_heading_before(doc, before_elm, "6.1.2 目标与线索场设置", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "实验目标为海域中的可疑目标。目标真实位置不直接暴露给搜索策略，USV 只能通过传感器探测结果和线索观测逐步更新搜索状态。目标数量设置为 3 个，并采用上界式目标数量设定，以保证算法面对的是未知目标状态而不是已知目标列表。",
    )
    add_paragraph_before(
        doc,
        before_elm,
        "目标运动模式包括 static 和 random_walk 两类。线索场由目标诱导的弱观测生成，并通过高斯过程模型形成 GP clue 场。主线基线实验使用 UCB 采集函数；anomaly_upper_tail 仅在第 6.4 节作为采集函数消融变体进行比较，不改变 target intensity、recency / staleness 或路径安全层语义。",
    )
    add_heading_before(doc, before_elm, "6.1.3 对比方法设置", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "单艇实验以 marine_knownmap_path_v2_infosampled 为主线策略，并启用 simple_ring_v1 候选视点生成和 soft_clearance_astar_v1 路径安全生成层。该设置对应第 4 章所述的“综合信息价值场驱动热点选择、环形候选视点生成、短时域路径段评分、避障路径生成和滚动执行”主链。",
    )
    add_paragraph_before(
        doc,
        before_elm,
        "双艇实验以 marine_knownmap_path_v2_infosampled_2usv 的 coordinated 模式为本文协同方法，并以 independent 模式作为对比。independent 表示两艇分别使用局部规划逻辑，不使用 residual_map 顺序分配和团队联合评分；coordinated 表示第 5 章提出的中心化共享状态、软责任区先验、顺序分配、重复覆盖惩罚和艇间冲突安全执行机制。",
    )
    add_paragraph_before(
        doc,
        before_elm,
        "anomaly-aware acquisition 消融实验比较 ucb 与 anomaly_upper_tail 两种 clue acquisition mode。其中 anomaly_upper_tail 固定 anomaly_tail_quantile=0.90，并采用经敏感性实验选定的 anomaly_weight_lambda=1.25。除采集函数模式及其对应上尾加权参数外，地图、目标设置、传感器参数、路径安全模式和随机种子保持一致。",
    )
    add_heading_before(doc, before_elm, "6.1.4 评价指标", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "实验评价首先关注搜索任务完成效率，包括 time_to_first_detection 和 time_to_all_found。前者表示首次发现任一目标所需步数，后者表示发现全部目标所需步数。由于部分 episode 在最大步数内未发现全部目标，表中同时给出 n、n_first 和 n_all，其中 n 表示总 episode 数，n_first 表示存在首次发现时间的 episode 数，n_all 表示在 240 步内发现全部目标的 episode 数。",
    )
    add_paragraph_before(
        doc,
        before_elm,
        "搜索完成程度由 success_all_found、detection_rate 和 found_count 共同刻画。success_all_found 定义为 n_all / n，detection_rate 表示最终发现目标数量占目标数量的比例，found_count 表示仿真结束时发现的目标个数。对于双艇实验，还统计 duplicate_viewpoint_ratio、cross_region_assignment_ratio 和 wait_count_total，以解释 coordinated 是否减少重复观测、是否形成合理分工以及 reservation_v1 是否带来额外等待成本。",
    )

    # 6.2
    add_heading_before(doc, before_elm, "6.2 单 USV 已知地图搜索基线及目标运动鲁棒性实验", "二级标题")
    add_heading_before(doc, before_elm, "6.2.1 实验目的", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "本节实验用于验证第 4 章所述单 USV 信息驱动安全路径规划方法在已知静态障碍地图下的基础搜索能力，并考察该方法在 static 与 random_walk 两类目标运动模式下的稳定性。与第 6.4 节的 anomaly-aware acquisition 消融不同，本节只取 clue_acquisition_mode=ucb 的结果作为单艇基线。",
    )
    add_heading_before(doc, before_elm, "6.2.2 实验设置", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "两组实验均使用三类 60×40 已知静态地图，每类地图使用 0–9 共 10 个随机种子，最大仿真步数为 240。为保持单艇基线实验语义一致，结果表只统计 UCB 采集模式；同一批运行中生成的 anomaly_upper_tail 结果将在第 6.4 节用于采集函数消融分析。",
    )
    add_caption_before(doc, before_elm, "表6-1 单艇实验参数设置")
    add_table_before(
        doc,
        before_elm,
        ["参数项", "设置"],
        [
            ["实验对象", "单 USV 已知静态地图搜索"],
            ["地图尺寸", "60×40 栅格"],
            ["地图类型", "open_water、harbor_cove、peninsula_passage"],
            ["目标运动模式", "static、random_walk"],
            ["随机种子", "0–9，共 10 个 episode"],
            ["最大仿真步数", "240"],
            ["基础策略", "marine_knownmap_path_v2_infosampled"],
            ["候选视点生成模式", "simple_ring_v1"],
            ["线索采集模式", "UCB"],
            ["路径安全模式", "soft_clearance_astar_v1"],
            ["传感器与分辨率", "sensor_range_m=25.0，resolution_m=5.0"],
            ["anomaly 说明", "本节不使用 anomaly_upper_tail，该结果留作第 6.4 节消融"],
        ],
        font_size=9.5,
    )
    add_heading_before(doc, before_elm, "6.2.3 评价指标与结果组织方式", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "本节结果围绕搜索效率、搜索完成程度和搜索行为辅助指标组织。time_to_first_detection 与 time_to_all_found 用于衡量搜索效率；success_all_found、detection_rate 和 found_count 用于反映搜索完成程度；known_free_observation_ratio_final 与 path_length 用于说明搜索覆盖范围和实际运动代价。",
    )
    add_heading_before(doc, before_elm, "6.2.4 静态目标模式下的单艇搜索结果", "三级标题")
    add_caption_before(doc, before_elm, "表6-2 静态目标模式下单艇 UCB 搜索结果")
    add_table_before(doc, before_elm, single_headers, single_table_rows(t62_static), font_size=7.5)
    add_paragraph_before(
        doc,
        before_elm,
        "由表6-2可见，在 static 目标模式下，单艇 UCB 基线的总体首次发现时间为 45.3 步，成功完成全部目标搜索的样本中 T_all 均值为 158.3 步；n_all=14，对应 success_all_found 为 46.7%。最终 detection_rate 为 81.1%，平均 found_count 为 2.43，说明在 240 步预算内单艇能够稳定发现大部分目标，但仍存在部分随机种子未能完成全部目标搜索。",
    )
    add_paragraph_before(
        doc,
        before_elm,
        "分地图看，harbor_cove 的首次发现时间最短，为 33.3 步，说明港湾结构中的高价值线索较容易在早期引导 USV 接近目标区域；peninsula_passage 的 T_all 最高，为 192.3 步，且 path_length 达到 220.9，反映出半岛和通道结构会显著增加绕行与覆盖代价。",
    )
    add_heading_before(doc, before_elm, "6.2.5 随机游走目标模式下的单艇搜索结果", "三级标题")
    add_caption_before(doc, before_elm, "表6-3 随机游走目标模式下单艇 UCB 搜索结果")
    add_table_before(doc, before_elm, single_headers, single_table_rows(t62_random), font_size=7.5)
    add_paragraph_before(
        doc,
        before_elm,
        "由表6-3可见，在 random_walk 目标模式下，单艇 UCB 基线的总体首次发现时间为 53.2 步，T_all 均值为 153.1 步，n_all=17，对应 success_all_found 为 56.7%。与 static 模式相比，首次发现时间有所增加，但完成全部目标搜索的 episode 数量和最终 detection_rate 均略有提高，说明轻微目标运动并未破坏单艇信息驱动搜索主链。",
    )
    add_paragraph_before(
        doc,
        before_elm,
        "在 open_water 中，random_walk 模式下 n_all 达到 7，success_all_found 为 70.0%，但 T_first 也升至 71.1 步，表明开阔地图中的早期搜索方向更依赖当前线索排序。harbor_cove 的 T_first 仍较低，为 32.6 步；peninsula_passage 的 T_all 为 176.7 步，继续体现通道型拓扑对搜索完成时间的影响。",
    )
    add_heading_before(doc, before_elm, "6.2.6 不同地图结构下的搜索行为分析", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "综合 static 与 random_walk 两类结果可以看出，地图结构对单艇搜索行为的影响较为稳定。open_water 障碍少，路径更直接，但首次发现时间对目标运动更敏感；harbor_cove 在两类目标模式下都具有较短的首次发现时间，说明信息价值场能够较快将 USV 引向港湾结构中的高价值区域。",
    )
    add_paragraph_before(
        doc,
        before_elm,
        "peninsula_passage 的共同特征是 path_length 和 T_all 偏高，说明半岛与狭窄通路会使单艇必须执行更长路径、覆盖更多自由区域后才能完成搜索。该结果为后续双艇协同实验提供了动机：在复杂通道地图中，多艇协同若能形成互补分工，理论上更有可能降低单艇长距离往返搜索带来的时间代价。",
    )
    add_heading_before(doc, before_elm, "6.2.7 本节小结", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "本节基于 simple_ring_v1 候选视点生成、UCB 线索采集和 soft_clearance_astar_v1 安全路径生成的真实运行结果验证了单 USV 已知地图搜索基线。结果表明，单艇方法在 static 与 random_walk 两类目标模式下均能够发现大部分目标，但在 240 步预算内并不能保证所有 episode 都完成全部目标搜索。",
    )

    # 6.3
    add_heading_before(doc, before_elm, "6.3 双 USV 协同搜索实验", "二级标题")
    add_heading_before(doc, before_elm, "6.3.1 实验目的", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "本节实验用于验证第 5 章提出的双 USV 协同搜索决策机制是否能够在相同双艇平台与相同已知静态地图设置下提高团队搜索效果。由于单艇实验采用 60×40 地图，而双艇实验采用 60×80 地图，二者搜索面积和目标分布空间不同，因此本节不将单 USV 与双 USV 进行直接数值比较。",
    )
    add_heading_before(doc, before_elm, "6.3.2 实验设置", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "本节使用 open_water、harbor_cove 和 peninsula_passage 三类 60×80 已知静态地图，每类地图使用 0–9 共 10 个随机种子，最大仿真步数为 240。所有结果均固定 clue_acquisition_mode=ucb，以便将实验变量限定为双艇团队分配机制本身。",
    )
    add_caption_before(doc, before_elm, "表6-4 双艇协同实验参数设置")
    add_table_before(
        doc,
        before_elm,
        ["参数项", "设置"],
        [
            ["实验对象", "双 USV 已知静态地图搜索"],
            ["比较方法", "two_usv_independent / two_usv_coordinated"],
            ["地图尺寸", "60×80 栅格"],
            ["地图类型", "open_water、harbor_cove、peninsula_passage"],
            ["目标模式", "static、random_walk"],
            ["随机种子", "0–9，共 10 个 episode"],
            ["最大仿真步数", "240"],
            ["policy", "marine_knownmap_path_v2_infosampled_2usv"],
            ["候选视点生成模式", "simple_ring_v1"],
            ["clue_acquisition_mode", "ucb"],
            ["path_safety_mode", "soft_clearance_astar_v1"],
            ["team_path_avoidance_mode", "reservation_v1"],
            ["协同差异", "independent 不使用 residual_map 顺序分配；coordinated 使用 sequential allocation-lite 与重叠抑制"],
        ],
        font_size=9.0,
    )
    add_heading_before(doc, before_elm, "6.3.3 评价指标与结果组织方式", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "本节分别给出 static 与 random_walk 两类目标模式下 independent 与 coordinated 的 UCB 搜索结果，并进一步给出 coordinated 减 independent 的差值表。对于时间类指标，差值小于 0 表示 coordinated 更快；对于 success_all_found、detection_rate 和 found_count，差值大于 0 表示 coordinated 完成程度更高。",
    )
    add_heading_before(doc, before_elm, "6.3.4 静态目标模式下的双艇协同搜索结果", "三级标题")
    add_caption_before(doc, before_elm, "表6-5 静态目标模式下双艇 UCB independent 与 coordinated 对比结果")
    add_table_before(doc, before_elm, assignment_headers, assignment_rows(t63_static), font_size=6.5)
    add_caption_before(doc, before_elm, "表6-6 静态目标模式下 coordinated - independent 差值结果")
    add_table_before(doc, before_elm, coord_delta_headers, coord_delta_rows(t63_delta_static), font_size=7.5)
    add_paragraph_before(
        doc,
        before_elm,
        "由表6-5和表6-6可见，在 static 目标模式下，coordinated 相比 independent 的总体 T_first 降低 24.8 步，success_all_found 提高 13.3 个百分点，detection_rate 提高 13.3 个百分点，found_count 提高 0.40。这说明顺序分配与重叠抑制机制能够更早形成有效目标发现，并提高静态目标条件下的整体搜索完成程度。",
    )
    add_paragraph_before(
        doc,
        before_elm,
        "协同行为指标也显示 coordinated 的实际作用：duplicate_viewpoint_ratio 平均降低 3.2 个百分点，cross_region_assignment_ratio 平均降低 28.5 个百分点，wait_count_total 平均减少 0.3。peninsula_passage 中 ΔT_all 为 N/A，这是因为成对差值统计中有效完成全部目标的样本不足，不应强行解释为具体时间收益。",
    )
    add_heading_before(doc, before_elm, "6.3.5 随机游走目标模式下的双艇协同搜索结果", "三级标题")
    add_caption_before(doc, before_elm, "表6-7 随机游走目标模式下双艇 UCB independent 与 coordinated 对比结果")
    add_table_before(doc, before_elm, assignment_headers, assignment_rows(t63_random), font_size=6.5)
    add_caption_before(doc, before_elm, "表6-8 随机游走目标模式下 coordinated - independent 差值结果")
    add_table_before(doc, before_elm, coord_delta_headers, coord_delta_rows(t63_delta_random), font_size=7.5)
    add_paragraph_before(
        doc,
        before_elm,
        "由表6-7和表6-8可见，在 random_walk 目标模式下，coordinated 的总体 T_first 比 independent 降低 15.3 步，T_all 降低 8.1 步，说明协同机制仍有助于提高早期发现效率和成功样本中的完成效率。但 success_all_found 总体下降 6.7 个百分点，detection_rate 仅提高 1.1 个百分点，found_count 仅提高 0.03，说明动态目标下 coordinated 的完成可靠性收益并非单调成立。",
    )
    add_paragraph_before(
        doc,
        before_elm,
        "从协同行为看，random_walk 模式下 coordinated 仍显著降低 duplicate_viewpoint_ratio 和 cross_region_assignment_ratio，二者分别下降 2.9 和 24.0 个百分点。该结果说明 residual_map 顺序分配仍然有效抑制重复搜索，但动态目标带来的线索变化会削弱协同分配对最终全部发现成功率的稳定提升。",
    )
    add_heading_before(doc, before_elm, "6.3.6 不同地图结构下的协同效果分析", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "open_water 中障碍较少，coordinated 在两类目标模式下均缩短 T_first，并降低重复视点选择。在 static 模式下，coordinated 的 detection_rate 相比 independent 提高 10.0 个百分点；在 random_walk 模式下，success_all_found 提高 10.0 个百分点，说明开阔地图中协同分配更容易转化为互补覆盖收益。",
    )
    add_paragraph_before(
        doc,
        before_elm,
        "harbor_cove 中凹形岸线和港湾入口会使两艇容易被相近热点吸引。coordinated 在 static 模式下降低 T_first，但 T_all 差值为 +36.0 步；在 random_walk 模式下，success_all_found 下降 30.0 个百分点。这表明港湾拓扑下 coordinated 能抑制重复视点，却不一定稳定提高全部目标完成率。",
    )
    add_paragraph_before(
        doc,
        before_elm,
        "peninsula_passage 中通道与半岛结构对双艇协同影响最明显。static 模式下，coordinated 将 success_all_found 从 10.0% 提升至 50.0%，detection_rate 从 53.3% 提升至 83.3%；random_walk 模式下则保持相同 success_all_found，但提高 detection_rate 和 found_count。该现象说明在通道型地图中，显式团队分配更有利于避免两艇在相同通道或相同热点附近重复搜索。",
    )
    add_heading_before(doc, before_elm, "6.3.7 本节小结", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "本节在相同 60×80 双 USV 已知地图设置下，比较了 UCB 线索采集模式下的 independent 与 coordinated 两类双艇搜索方法。实验表明，coordinated 的稳定收益主要体现在降低重复视点选择、减少跨责任区偏移并改善早期发现效率；但在 random_walk 目标模式和部分复杂港湾地图中，全部发现成功率并不总是高于 independent。",
    )

    # 6.4
    add_heading_before(doc, before_elm, "6.4 anomaly-aware acquisition 消融实验", "二级标题")
    add_heading_before(doc, before_elm, "6.4.1 实验目的与变量控制", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "本节用于检验 anomaly-aware acquisition 中 anomaly_upper_tail 采集方式相对于标准 UCB 的实际作用。根据前置参数敏感性实验，本文正式 anomaly_upper_tail 设置为 anomaly_tail_quantile=0.90、anomaly_weight_lambda=1.25。该变体只改变 GP clue acquisition 的排序，不改变 target intensity、recency、综合信息价值场融合方式或路径安全约束。",
    )
    add_heading_before(doc, before_elm, "6.4.2 评价指标与差值定义", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "本节采用 time_to_first_detection、time_to_all_found、success_all_found、detection_rate 和 found_count 作为核心指标。差值统一定义为 anomaly_upper_tail 减去 UCB。对于时间类指标，差值小于 0 表示 anomaly_upper_tail 更快；对于 success_all_found、detection_rate 和 found_count，差值大于 0 表示 anomaly_upper_tail 的完成程度更高。",
    )
    add_caption_before(doc, before_elm, "表6-9 anomaly 消融实验评价指标与差值定义")
    add_table_before(
        doc,
        before_elm,
        ["指标", "含义", "解释重点"],
        [
            ["time_to_first_detection", "首次发现任一目标所需步数", "越小表示越早获得首个目标发现"],
            ["time_to_all_found", "发现全部目标所需步数", "仅对成功完成全部目标的 episode 统计均值"],
            ["n_all", "在 240 步内发现全部目标的 episode 数", "用于解释 time_to_all_found 的有效样本数量"],
            ["success_all_found", "n_all / n", "越高表示整体完成可靠性越强"],
            ["detection_rate", "最终发现目标数占目标数比例", "用于反映未全部完成时的搜索完成程度"],
            ["ΔT_all", "T_all(anomaly_upper_tail) - T_all(UCB)", "小于 0 表示 anomaly 更快，大于 0 表示 anomaly 更慢"],
        ],
        font_size=8.5,
    )
    add_heading_before(doc, before_elm, "6.4.3 单 USV 下 anomaly acquisition 的影响", "三级标题")
    add_caption_before(doc, before_elm, "表6-10 单艇 static 模式下 UCB 与 anomaly_upper_tail 对比")
    add_table_before(doc, before_elm, anomaly_headers, anomaly_main_rows(t64_single_static), font_size=7.0)
    add_caption_before(doc, before_elm, "表6-11 单艇 random_walk 模式下 UCB 与 anomaly_upper_tail 对比")
    add_table_before(doc, before_elm, anomaly_headers, anomaly_main_rows(t64_single_random), font_size=7.0)
    add_caption_before(doc, before_elm, "表6-12 单艇 anomaly_upper_tail - UCB 差值结果")
    rows_612 = [["static"] + row for row in anomaly_delta_rows(t64_d_single_static)] + [
        ["random_walk"] + row for row in anomaly_delta_rows(t64_d_single_random)
    ]
    add_table_before(
        doc,
        before_elm,
        ["目标模式"] + anomaly_delta_headers,
        rows_612,
        font_size=7.0,
    )
    add_paragraph_before(
        doc,
        before_elm,
        "由表6-10至表6-12可见，单艇下 anomaly_upper_tail 的效果具有条件性。static 模式中，anomaly_upper_tail 将 success_all_found 从 46.7% 提升至 73.3%，detection_rate 从 81.1% 提升至 90.0%，但 T_all 差值为 -0.8 步，时间收益并不显著。random_walk 模式中，anomaly_upper_tail 将 success_all_found 提高 10.0 个百分点，但 T_all 差值为 +14.2 步，说明其提高了完成可靠性，却未必缩短成功样本中的全部发现时间。",
    )
    add_heading_before(doc, before_elm, "6.4.4 双 USV coordinated 下 anomaly acquisition 的影响", "三级标题")
    add_caption_before(doc, before_elm, "表6-13 双艇 coordinated static 模式下 UCB 与 anomaly_upper_tail 对比")
    add_table_before(doc, before_elm, anomaly_headers, anomaly_main_rows(t64_two_static), font_size=7.0)
    add_caption_before(doc, before_elm, "表6-14 双艇 coordinated random_walk 模式下 UCB 与 anomaly_upper_tail 对比")
    add_table_before(doc, before_elm, anomaly_headers, anomaly_main_rows(t64_two_random), font_size=7.0)
    add_caption_before(doc, before_elm, "表6-15 双艇 coordinated anomaly_upper_tail - UCB 差值结果")
    rows_615 = [["static"] + row for row in anomaly_delta_rows(t64_d_two_static)] + [
        ["random_walk"] + row for row in anomaly_delta_rows(t64_d_two_random)
    ]
    add_table_before(
        doc,
        before_elm,
        ["目标模式"] + anomaly_delta_headers,
        rows_615,
        font_size=7.0,
    )
    add_paragraph_before(
        doc,
        before_elm,
        "由表6-13至表6-15可见，在双艇 coordinated 系统中，static 模式下 anomaly_upper_tail 使 T_first 降低 7.8 步、T_all 降低 25.4 步，但 success_all_found 下降 3.3 个百分点，detection_rate 与 found_count 基本不变。random_walk 模式下，anomaly_upper_tail 将 success_all_found 从 46.7% 提升至 66.7%，detection_rate 从 81.1% 提升至 85.6%，并使 T_all 降低 9.4 步，但 T_first 略有上升。",
    )
    add_heading_before(doc, before_elm, "6.4.5 地图结构对 anomaly 效果的影响", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "从地图维度看，anomaly_upper_tail 的收益不是单调一致的。open_water 障碍少，上尾异常线索更容易转化为可执行路径收益；harbor_cove 中 anomaly 对完成率和探测率的影响随单艇/双艇结构发生变化；peninsula_passage 受通道拓扑限制，即使异常线索改变热点排序，也不一定能显著缩短 T_all。",
    )
    add_paragraph_before(
        doc,
        before_elm,
        "因此，anomaly_upper_tail 的作用应被解释为线索采集偏置对候选区域排序的影响，而不是对目标真实存在概率或避障可达性的改变。在复杂地图中，路径代价、候选视点可达性和双艇 residual 分配都会共同决定异常线索能否转化为搜索收益。",
    )
    add_heading_before(doc, before_elm, "6.4.6 机制解释与边界讨论", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "从机制上看，anomaly_upper_tail 提高当前 GP clue 后验中上尾异常区域的排序权重，使路径规划器更倾向于观察可能具有异常高线索价值的区域。该机制适合发现相对于当前整体线索场偏高的局部区域，因此在部分地图和目标运动模式下能够提高 success_all_found 或缩短 T_all。",
    )
    add_paragraph_before(
        doc,
        before_elm,
        "但 anomaly_upper_tail 也存在边界。第一，它只改变线索采集排序，不改变 target intensity、recency 或避障可达性；第二，在双艇 coordinated 中，路径选择还受到 residual_map、overlap penalty、same-viewpoint penalty 和 reservation_v1 的共同影响；第三，对于 random_walk 目标，历史异常区域可能随目标移动而失效。因此，本文不将 anomaly_upper_tail 表述为在所有条件下全面优于 UCB 的采集函数。",
    )
    add_heading_before(doc, before_elm, "6.4.7 本节小结", "三级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "本节基于单 USV 与双 USV coordinated 的真实运行结果，对 UCB 与 anomaly_upper_tail（q=0.90，λ=1.25）进行了消融比较。结果表明，anomaly_upper_tail 在部分场景中能改善搜索完成可靠性或缩短成功样本中的全部发现时间，但收益受地图结构、目标运动模式和系统结构共同影响。因而，本文将其作为 paper-inspired upper-tail variant 进行分析，而不是将其包装为无条件优于 UCB 的主导机制。",
    )

    # 6.5
    add_heading_before(doc, before_elm, "6.5 本章小结", "二级标题")
    add_paragraph_before(
        doc,
        before_elm,
        "本章基于 simple_ring_v1 主线实验数据，对单 USV 基线、双 USV 协同机制和 anomaly-aware acquisition 变体进行了系统分析。单艇实验表明，simple_ring_v1 候选视点生成与 soft_clearance_astar_v1 安全路径生成能够在两类目标运动模式下稳定发现大部分目标，但受地图拓扑和 240 步预算限制，仍有部分 episode 未能发现全部目标。",
    )
    add_paragraph_before(
        doc,
        before_elm,
        "双艇实验表明，coordinated 的主要收益体现在降低重复视点选择、减少跨责任区偏移并改善早期发现效率；但其完成全部目标的成功率并非在所有地图和目标运动模式下都稳定优于 independent。anomaly 消融实验进一步说明，anomaly_upper_tail（q=0.90，λ=1.25）具有条件性收益，适合作为上尾异常线索采集变体分析，而不应被解释为无条件优于 UCB。",
    )

    doc.save(REPORT_PATH)


if __name__ == "__main__":
    main()
