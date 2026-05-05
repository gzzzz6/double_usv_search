# -*- coding: utf-8 -*-
"""Update thesis anomaly_upper_tail wording and tables to lambda=1.25."""

from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


REPORT_PATH = Path(r"F:\pythonprojects\my_report.docx")


def set_run_fonts(run) -> None:
    run.font.name = "Times New Roman"
    r_fonts = run._element.rPr.rFonts
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "宋体")


def replace_paragraph_text(paragraph, text: str) -> None:
    for run in list(paragraph.runs):
        run.text = ""
    run = paragraph.add_run(text)
    set_run_fonts(run)


def set_cell_text(cell, text: str) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            set_run_fonts(run)


def replace_first_paragraph_starting(doc: Document, prefix: str, text: str) -> None:
    for paragraph in doc.paragraphs:
        current = paragraph.text.strip()
        if current.startswith(prefix) and "\t" not in current:
            replace_paragraph_text(paragraph, text)
            return
    raise RuntimeError(f"Could not find paragraph starting with: {prefix}")


def replace_exact_caption(doc: Document, old: str, new: str) -> None:
    for paragraph in doc.paragraphs:
        current = paragraph.text.strip()
        if current == old:
            replace_paragraph_text(paragraph, new)
            return
    raise RuntimeError(f"Could not find caption: {old}")


def update_table(table, rows: list[list[str]]) -> None:
    if len(table.rows) != len(rows):
        raise RuntimeError(f"Row count mismatch: doc={len(table.rows)} target={len(rows)}")
    for row, values in zip(table.rows, rows):
        if len(row.cells) != len(values):
            raise RuntimeError(
                f"Column count mismatch: doc={len(row.cells)} target={len(values)}"
            )
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value)


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = REPORT_PATH.with_name(
        f"my_report_backup_before_anomaly_lambda125_{stamp}.docx"
    )
    shutil.copy2(REPORT_PATH, backup)

    doc = Document(REPORT_PATH)

    replace_first_paragraph_starting(
        doc,
        "用此方法构造的anomaly权重保留了UCB",
        (
            "用此方法构造的anomaly权重保留了UCB在“高均值利用”与“不确定性探索”之间的基本平衡，"
            "同时对更可能属于高值上尾的区域给予额外强调，从而提升搜索策略对潜在异常热点的敏感性。"
            "本文实验中上尾分位参数取 q=0.90；异常强调系数 λ 通过后续敏感性实验在 "
            "{0.50, 0.75, 1.00, 1.25} 中选取。综合单艇、双艇以及 static、random_walk "
            "目标模式下的平均结果，λ=1.25 的整体表现优于原 λ=1.00 设置，因此本文正式 "
            "anomaly_upper_tail 结果采用 anomaly_weight_lambda=1.25。该参数只影响 GP clue "
            "acquisition 层，不改变 target intensity、recency 或路径安全层。"
        ),
    )

    replace_first_paragraph_starting(
        doc,
        "anomaly-aware acquisition 消融实验比较 ucb",
        (
            "anomaly-aware acquisition 消融实验比较 ucb 与 anomaly_upper_tail 两种 clue acquisition mode。"
            "其中 anomaly_upper_tail 固定 anomaly_tail_quantile=0.90，并采用经敏感性实验选定的 "
            "anomaly_weight_lambda=1.25。除采集函数模式及其对应上尾加权参数外，地图、目标设置、"
            "传感器参数、路径安全模式和随机种子保持一致，以便将性能差异主要归因于线索采集策略的变化。"
        ),
    )

    replace_first_paragraph_starting(
        doc,
        "本节用于检验 anomaly-aware acquisition",
        (
            "本节用于检验 anomaly-aware acquisition 中 anomaly_upper_tail 采集方式相对于标准 UCB 的实际作用。"
            "与第 6.2 节不同，本节不再只把 UCB 作为单艇基线，而是直接比较 clue_acquisition_mode=ucb "
            "与 clue_acquisition_mode=anomaly_upper_tail 两种线索采集方式在相同系统结构下的搜索表现。"
            "根据前置参数敏感性实验，本文正式 anomaly_upper_tail 设置为 anomaly_tail_quantile=0.90、"
            "anomaly_weight_lambda=1.25。"
        ),
    )

    replace_exact_caption(
        doc,
        "表6-8 单 USV 下 UCB 与 anomaly_upper_tail 总体结果对比",
        "表6-8 单 USV 下 UCB 与 anomaly_upper_tail（λ=1.25）总体结果对比",
    )
    replace_exact_caption(
        doc,
        "表6-9 双 USV coordinated 下 UCB 与 anomaly_upper_tail 总体结果对比",
        "表6-9 双 USV coordinated 下 UCB 与 anomaly_upper_tail（λ=1.25）总体结果对比",
    )
    replace_exact_caption(
        doc,
        "表6-10 各地图 anomaly 效果差值汇总",
        "表6-10 各地图 anomaly_upper_tail（λ=1.25）效果差值汇总",
    )

    replace_first_paragraph_starting(
        doc,
        "由表6-8可见",
        (
            "由表6-8可见，单 USV 下 anomaly_upper_tail 的效果具有明显条件性。在 static 目标模式中，"
            "anomaly_upper_tail 的总体 time_to_all_found 为 131.5 步，UCB 为 158.7 步，"
            "ΔT_all=-27.2，说明成功完成全部目标搜索的样本中完成时间明显缩短；但 success_all_found "
            "从 53.3% 降至 46.7%，detection_rate 从 80.0% 小幅提高到 81.1%。因此，静态目标下 "
            "λ=1.25 的 anomaly_upper_tail 更主要体现为成功样本完成时间缩短，而不是全面提高完成可靠性。"
        ),
    )
    replace_first_paragraph_starting(
        doc,
        "在 random_walk 目标模式中，anomaly_upper_tail 的总体 time_to_all_found 从 UCB 的 165.7",
        (
            "在 random_walk 目标模式中，anomaly_upper_tail 的总体 time_to_all_found 从 UCB 的 165.7 步"
            "下降到 152.6 步，ΔT_all=-13.1；同时 success_all_found 从 63.3% 提升到 70.0%，"
            "detection_rate 从 85.6% 提升到 88.9%，found_count 从 2.57 提升到 2.67。"
            "这说明在单艇动态目标场景中，λ=1.25 的上尾异常加权不仅能缩短成功样本的完成时间，"
            "也能提高整体搜索完成程度。"
        ),
    )
    replace_first_paragraph_starting(
        doc,
        "由表6-9可见",
        (
            "由表6-9可见，在双 USV coordinated 系统中，static 目标模式下 anomaly_upper_tail 带来轻度整体收益。"
            "其总体 time_to_all_found 为 152.6 步，低于 UCB 的 157.5 步，ΔT_all=-4.9；"
            "success_all_found 从 56.7% 提升到 60.0%，detection_rate 从 80.0% 提升到 81.1%。"
            "这说明在 λ=1.25 设置下，anomaly 偏置没有改变双艇协同主链，但可以在部分静态目标 episode 中"
            "帮助团队更早聚焦高值线索区域。"
        ),
    )
    replace_first_paragraph_starting(
        doc,
        "在 random_walk 目标模式下，双 USV coordinated 的 anomaly_upper_tail 呈现另一种特征",
        (
            "在 random_walk 目标模式下，双 USV coordinated 的 anomaly_upper_tail 同样表现出稳定改善："
            "time_to_all_found 从 159.9 步下降到 155.3 步，ΔT_all=-4.6；success_all_found "
            "从 60.0% 提升到 73.3%，detection_rate 从 83.3% 提升到 88.9%，found_count 从 2.50 "
            "提升到 2.67。也就是说，双艇动态目标场景中 λ=1.25 的 anomaly_upper_tail 更明显地改善了"
            "搜索完成可靠性，同时略微缩短了成功样本中的全部发现时间。"
        ),
    )
    replace_first_paragraph_starting(
        doc,
        "表6-10显示",
        (
            "表6-10显示，λ=1.25 时 anomaly_upper_tail 的效果仍不是单调一致的，而是受到地图拓扑、"
            "目标运动模式和系统结构共同影响。open_water 障碍少，异常上尾加权更容易转化为路径收益："
            "单艇 random_walk 下 ΔT_all=-8.3，双艇 random_walk 下 ΔT_all=-38.1，且两者的 "
            "success_all_found 和 detection_rate 均有所提升。"
        ),
    )
    replace_first_paragraph_starting(
        doc,
        "harbor_cove 中 anomaly 的作用",
        (
            "harbor_cove 中 anomaly 的收益主要体现为完成可靠性改善，但时间收益随系统结构和目标模式变化。"
            "单艇 static 下 anomaly_upper_tail 使 ΔT_all=-16.9，success_all_found 提高 30.0 个百分点，"
            "说明港湾结构中的异常线索有时能帮助单艇更早聚焦有效区域；双艇 static 下 ΔT_all=-6.8，"
            "success_all_found 保持不变，detection_rate 提高 10.0 个百分点。random_walk 下，单艇仍表现为"
            "完成时间缩短，而双艇则主要表现为完成率和探测率提升。"
        ),
    )
    replace_first_paragraph_starting(
        doc,
        "peninsula_passage 中通道拓扑限制了候选路径可达性",
        (
            "peninsula_passage 中通道拓扑限制了候选路径可达性。即使 anomaly_upper_tail 改变了热点排序，"
            "USV 仍可能需要经过固定通道或绕行路径才能到达有效观测位置。因此，该地图下 anomaly 效果仍然不稳定："
            "单艇 static 的 ΔT_all=-3.8，但 success_all_found 下降 20.0 个百分点；单艇 random_walk 的 "
            "ΔT_all=-10.8，完成率保持不变；双艇 random_walk 下 ΔT_all=+27.7，说明动态目标和双艇分配叠加后，"
            "异常线索偏置不一定能在通道受限地图中转化为更快完成。"
        ),
    )
    replace_first_paragraph_starting(
        doc,
        "因此，本文不将 anomaly_upper_tail 表述为",
        (
            "因此，本文不将 anomaly_upper_tail 表述为在所有条件下全面优于 UCB 的采集函数。更准确的结论是："
            "在 q=0.90、λ=1.25 设置下，anomaly_upper_tail 的总体平均表现优于 λ=1.00 的原始设置，"
            "并能在部分场景中改善 time_to_all_found、success_all_found 或 detection_rate；但其收益仍依赖地图结构、"
            "目标运动模式和单艇/双艇系统结构，适合在消融实验中作为 paper-inspired upper-tail variant 进行分析。"
        ),
    )
    replace_first_paragraph_starting(
        doc,
        "本节基于单 USV 与双 USV coordinated 的真实运行结果",
        (
            "本节基于单 USV 与双 USV coordinated 的真实运行结果，对 UCB 与 anomaly_upper_tail（q=0.90，λ=1.25）"
            "进行了消融比较。结果表明，λ=1.25 相比原 λ=1.00 具有更强的总体平均表现；在单艇 random_walk "
            "和双艇 random_walk 中，anomaly_upper_tail 能同时提高搜索完成可靠性和探测完成程度；在 static "
            "目标模式下，其收益更依赖地图结构和系统结构。总体而言，anomaly_upper_tail 可以作为本文上尾异常线索"
            "偏置的有效变体，但不应被解释为在所有场景中稳定优于 UCB。"
        ),
    )

    update_table(
        doc.tables[8],
        [
            ["目标模式", "采集方式", "time_to_first_detection", "time_to_all_found", "success_all_found（%）", "detection_rate（%）", "found_count"],
            ["static", "UCB", "45.6", "158.7", "53.3", "80.0", "2.40"],
            ["static", "anomaly_upper_tail", "45.6", "131.5", "46.7", "81.1", "2.43"],
            ["static", "Δ(anomaly-UCB)", "+0.0", "-27.2", "-6.7", "+1.1", "+0.03"],
            ["random_walk", "UCB", "46.0", "165.7", "63.3", "85.6", "2.57"],
            ["random_walk", "anomaly_upper_tail", "46.9", "152.6", "70.0", "88.9", "2.67"],
            ["random_walk", "Δ(anomaly-UCB)", "+0.9", "-13.1", "+6.7", "+3.3", "+0.10"],
        ],
    )
    update_table(
        doc.tables[9],
        [
            ["目标模式", "采集方式", "time_to_first_detection", "time_to_all_found", "success_all_found（%）", "detection_rate（%）", "found_count"],
            ["static", "UCB", "61.4", "157.5", "56.7", "80.0", "2.40"],
            ["static", "anomaly_upper_tail", "63.3", "152.6", "60.0", "81.1", "2.43"],
            ["static", "Δ(anomaly-UCB)", "+1.9", "-4.9", "+3.3", "+1.1", "+0.03"],
            ["random_walk", "UCB", "61.1", "159.9", "60.0", "83.3", "2.50"],
            ["random_walk", "anomaly_upper_tail", "64.8", "155.3", "73.3", "88.9", "2.67"],
            ["random_walk", "Δ(anomaly-UCB)", "+3.7", "-4.6", "+13.3", "+5.6", "+0.17"],
        ],
    )
    update_table(
        doc.tables[10],
        [
            ["系统", "目标模式", "地图", "ΔT_all", "Δsuccess_all_found（百分点）", "Δdetection_rate（百分点）"],
            ["单 USV", "static", "open_water", "-32.6", "-30.0", "-6.7"],
            ["单 USV", "static", "harbor_cove", "-16.9", "+30.0", "+13.3"],
            ["单 USV", "static", "peninsula_passage", "-3.8", "-20.0", "-3.3"],
            ["单 USV", "random_walk", "open_water", "-8.3", "+20.0", "+13.3"],
            ["单 USV", "random_walk", "harbor_cove", "-18.3", "0.0", "-3.3"],
            ["单 USV", "random_walk", "peninsula_passage", "-10.8", "0.0", "0.0"],
            ["双 USV coordinated", "static", "open_water", "-6.8", "+20.0", "+6.7"],
            ["双 USV coordinated", "static", "harbor_cove", "-6.8", "0.0", "+10.0"],
            ["双 USV coordinated", "static", "peninsula_passage", "+8.4", "-10.0", "-13.3"],
            ["双 USV coordinated", "random_walk", "open_water", "-38.1", "+20.0", "+10.0"],
            ["双 USV coordinated", "random_walk", "harbor_cove", "+0.8", "+10.0", "+3.3"],
            ["双 USV coordinated", "random_walk", "peninsula_passage", "+27.7", "+10.0", "+3.3"],
        ],
    )

    doc.save(REPORT_PATH)
    print(f"backup={backup}")
    print(f"updated={REPORT_PATH}")


if __name__ == "__main__":
    main()
