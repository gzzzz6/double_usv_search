from pathlib import Path
from docx import Document
import zipfile

path = Path(r'F:\pythonprojects\my_report.docx')
doc = Document(path)
texts = [p.text for p in doc.paragraphs]
with zipfile.ZipFile(path) as z:
    document_xml = z.read('word/document.xml').decode('utf-8', errors='replace')

print('BAD_MARKERS')
bad_markers = ['锛', '鈥', '乱码', '????', '\ufffd']
found = False
joined = '\n'.join(texts)
for marker in bad_markers:
    if marker in joined or marker in document_xml:
        print('BAD_MARKER_FOUND:', repr(marker))
        found = True
if not found:
    print('none')

print('\nACTUAL_CH6_CAPTIONS')
for i, t in enumerate(texts):
    s = t.strip()
    if i >= 660 and s.startswith('表6-'):
        print(f'{i}: {s}')

print('\nTABLE_SAMPLES')
for idx in [1, 2, 5, 10, 15]:
    if idx < len(doc.tables):
        tbl = doc.tables[idx]
        print(f'TABLE {idx}: rows={len(tbl.rows)} cols={len(tbl.columns)}')
        for r in range(min(3, len(tbl.rows))):
            print(' | '.join(c.text.replace('\n','/') for c in tbl.rows[r].cells[:min(8, len(tbl.columns))]))

print('\nCH6_TEXT_TOKENS')
for token in ['simple_ring_v1', 'soft_clearance_astar_v1', 'reservation_v1', 'anomaly_tail_quantile=0.90', 'anomaly_weight_lambda=1.25', 'lambda=1.25', 'λ=1.25']:
    print(token, token in joined or token in document_xml)
