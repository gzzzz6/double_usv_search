from pathlib import Path
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
import re

p=Path(r'F:\pythonprojects\my_report.docx')
doc=Document(p)
print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables), 'inline_shapes', len(doc.inline_shapes))
print('\nCH45_RANGE')
for i,para in enumerate(doc.paragraphs):
    t=para.text.strip()
    if 488 <= i <= 540 or t.startswith(('4.5','4.6')):
        has_pic=bool(para._p.xpath('.//pic:pic'))
        has_math=bool(para._p.xpath('.//m:oMath'))
        print(f'{i}: style={para.style.name!r} pic={has_pic} math={has_math} text={t}')

print('\nFIG_CAPTIONS_BODY')
for i,para in enumerate(doc.paragraphs):
    t=para.text.strip()
    if t.startswith('图') and para.style.name == '图目录项':
        print(i, t)

print('\nCITATIONS')
for i,para in enumerate(doc.paragraphs):
    t=para.text.strip()
    if '[7,35' in t or '[7,34' in t or '[32-34' in t or '[34]' in t:
        print(i, t)

print('\nFORMULA_LABELS_45_46')
for i,para in enumerate(doc.paragraphs):
    t=para.text.strip()
    if re.fullmatch(r'（\d+）', t) and 490 <= i <= 540:
        print(i, t)

# Body block caption ordering
body=doc.element.body
blocks=[]
for child in body.iterchildren():
    if isinstance(child, CT_P):
        para=Paragraph(child, doc)
        blocks.append(('P', para.text.strip(), para.style.name, bool(child.xpath('.//pic:pic'))))
    elif isinstance(child, CT_Tbl):
        blocks.append(('T','<TABLE>','',False))

print('\nCAPTION_POSITION_CHECK')
for idx,(kind,text,style,pic) in enumerate(blocks):
    if style in {'图目录项','表目录项'} and text.startswith(('图','表')):
        prev=blocks[idx-1] if idx>0 else None
        nxt=blocks[idx+1] if idx+1<len(blocks) else None
        print(idx, text, 'prev=', prev[:3] if prev else None, 'next=', nxt[:3] if nxt else None)

print('\nBAD_MARKERS')
joined='\n'.join(p.text for p in doc.paragraphs)
for m in ['锛','鈥','????','�']:
    print(m, m in joined)
