from __future__ import annotations

from pathlib import Path

from docx import Document


REPORT = Path(r"F:\pythonprojects\my_report.docx")


def set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def main() -> None:
    doc = Document(REPORT)
    # Front table-directory result and appendix table captions. These entries belong
    # to the experiment chapter, which has moved from Chapter 6 to Chapter 5.
    for idx in list(range(262, 277)) + list(range(757, 776)):
        text = doc.paragraphs[idx].text
        if "表6-" in text:
            set_paragraph_text(doc.paragraphs[idx], text.replace("表6-", "表5-"))
    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()
