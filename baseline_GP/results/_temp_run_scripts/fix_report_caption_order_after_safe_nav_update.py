from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.text.paragraph import Paragraph


ROOT = Path(r"F:\pythonprojects")
DOCX_PATH = ROOT / "my_report.docx"
BACKUP_PATH = ROOT / f"my_report_before_caption_order_fix_{datetime.now():%Y%m%d_%H%M%S}.docx"


def find_heading(doc: Document, prefix: str, *, start: int = 0, style_name: str | None = None) -> int:
    for idx, paragraph in enumerate(doc.paragraphs[start:], start=start):
        if paragraph.text.strip().startswith(prefix) and (
            style_name is None or paragraph.style.name == style_name
        ):
            return idx
    raise RuntimeError(f"Cannot find heading starting with {prefix!r}")


def is_table_caption_element(element, doc: Document) -> bool:
    if not isinstance(element, CT_P):
        return False
    paragraph = Paragraph(element, doc)
    return paragraph.style.name == "表目录项" and paragraph.text.strip().startswith("表")


def fix_consecutive_table_caption_runs(doc: Document) -> int:
    body = doc.element.body
    changed_total = 0
    changed = True
    while changed:
        changed = False
        children = list(body.iterchildren())
        idx = 0
        while idx < len(children):
            if not isinstance(children[idx], CT_Tbl):
                idx += 1
                continue
            table_start = idx
            while idx < len(children) and isinstance(children[idx], CT_Tbl):
                idx += 1
            tables = children[table_start:idx]
            if len(tables) <= 1:
                idx += 1
                continue
            cap_start = idx
            while idx < len(children) and is_table_caption_element(children[idx], doc):
                idx += 1
            captions = children[cap_start:idx]
            if len(captions) < len(tables):
                continue
            for cap in captions[: len(tables)]:
                cap.getparent().remove(cap)
            for offset in reversed(range(len(tables))):
                tables[offset].addnext(captions[offset])
            changed_total += len(tables)
            changed = True
            break
    return changed_total


def normalize_formula_styles_in_ch45_ch46(doc: Document) -> int:
    start = find_heading(doc, "4.5", style_name="二级标题")
    end = find_heading(doc, "4.7", start=start + 1, style_name="二级标题")
    changed = 0
    for paragraph in doc.paragraphs[start:end]:
        text = paragraph.text.strip()
        if paragraph.style.name != "公式目录项":
            continue
        if not text:
            continue
        if re.fullmatch(r"（\d+）", text):
            continue
        paragraph.style = doc.styles["Normal"]
        changed += 1
    return changed


def main() -> None:
    shutil.copy2(DOCX_PATH, BACKUP_PATH)
    doc = Document(DOCX_PATH)
    fixed_runs = fix_consecutive_table_caption_runs(doc)
    fixed_styles = normalize_formula_styles_in_ch45_ch46(doc)
    doc.save(DOCX_PATH)
    print(f"updated={DOCX_PATH}")
    print(f"backup={BACKUP_PATH}")
    print(f"fixed_consecutive_table_caption_runs={fixed_runs}")
    print(f"fixed_formula_text_styles={fixed_styles}")


if __name__ == "__main__":
    main()
