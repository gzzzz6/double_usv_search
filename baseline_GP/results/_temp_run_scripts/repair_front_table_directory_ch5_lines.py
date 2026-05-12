from __future__ import annotations

from pathlib import Path

from docx import Document


REPORT = Path(r"F:\pythonprojects\my_report.docx")


def hard_set_paragraph_text(paragraph, text: str) -> None:
    # TOC/list entries can contain hyperlink/field runs that paragraph.runs does not
    # fully expose. Clear XML content to avoid duplicate old and new entries.
    paragraph._p.clear_content()
    paragraph.add_run(text)


def main() -> None:
    doc = Document(REPORT)
    for idx in range(262, 277):
        text = doc.paragraphs[idx].text.strip()
        if "表5-" in text:
            text = "表5-" + text.split("表5-", 1)[1]
        elif "表6-" in text:
            text = text.replace("表6-", "表5-")
        hard_set_paragraph_text(doc.paragraphs[idx], text)
    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()
