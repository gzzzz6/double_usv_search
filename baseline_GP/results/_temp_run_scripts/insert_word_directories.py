# -*- coding: utf-8 -*-
"""Insert Word TOC/list fields for headings, formulas, figures, and tables."""

from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


REPORT_PATH = Path(r"F:\pythonprojects\my_report.docx")


def set_run_font(run, *, name: str = "宋体", size_pt: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def clear_paragraph(paragraph) -> None:
    for run in paragraph.runs:
        run.text = ""


def set_paragraph_text(paragraph, text: str, *, bold: bool = False, size_pt: float = 16.0) -> None:
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size_pt=size_pt, bold=bold)


def add_complex_field(paragraph, instruction: str, placeholder: str) -> None:
    run_begin = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    fld_begin.set(qn("w:dirty"), "true")
    run_begin._r.append(fld_begin)

    run_instr = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " " + instruction + " "
    run_instr._r.append(instr)

    run_sep = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run_sep._r.append(fld_sep)

    run_text = paragraph.add_run(placeholder)
    set_run_font(run_text, size_pt=12.0)

    run_end = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run_end._r.append(fld_end)


def add_tc_field(paragraph, entry: str, identifier: str) -> None:
    # Avoid adding duplicate TC fields if the script is rerun.
    xml_text = paragraph._p.xml
    if f' TC "{entry}" \\f {identifier} ' in xml_text:
        return
    run_begin = paragraph.add_run()
    run_begin.font.hidden = True
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run_begin._r.append(fld_begin)

    run_instr = paragraph.add_run()
    run_instr.font.hidden = True
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f' TC "{entry}" \\f {identifier} \\l 1 '
    run_instr._r.append(instr)

    run_end = paragraph.add_run()
    run_end.font.hidden = True
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run_end._r.append(fld_end)


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    parent.remove(element)


def is_actual_figure_caption(text: str) -> bool:
    if "\t" in text:
        return False
    if not re.match(r"^图\d+[-.]\d+\s+", text):
        return False
    if "给出" in text or "展示" in text or "可以看到" in text:
        return False
    return len(text) <= 80


def is_actual_table_caption(text: str) -> bool:
    if "\t" in text:
        return False
    if not re.match(r"^表\d+[-.]\d+\s+", text):
        return False
    return len(text) <= 90


def is_formula_number(text: str) -> bool:
    return bool(re.match(r"^（\d+）$", text))


def mark_entries(doc: Document) -> tuple[int, int, int]:
    figure_count = 0
    table_count = 0
    formula_count = 0
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if is_actual_figure_caption(text):
            add_tc_field(paragraph, text, "F")
            figure_count += 1
        elif is_actual_table_caption(text):
            add_tc_field(paragraph, text, "T")
            table_count += 1
        elif is_formula_number(text):
            add_tc_field(paragraph, f"公式{text}", "E")
            formula_count += 1
    return formula_count, figure_count, table_count


def insert_directory_block(doc: Document) -> None:
    first_heading = None
    first_heading_idx = None
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text.startswith("3 场建模"):
            first_heading = paragraph
            first_heading_idx = idx
            break
    if first_heading is None or first_heading_idx is None:
        raise RuntimeError("Could not find first body heading '3 场建模'.")

    # Remove the old manual directory block before the first body heading.
    for paragraph in list(doc.paragraphs[23:first_heading_idx]):
        delete_paragraph(paragraph)

    blocks = [
        (
            "目 录",
            r'TOC \h \z \t "一级标题,1,二级标题,2,三级标题,3"',
            "请在 Word 中按 Ctrl+A 后按 F9 更新目录",
        ),
        ("公式目录", r"TOC \h \z \f E", "请更新域以生成公式目录"),
        ("图目录", r"TOC \h \z \f F", "请更新域以生成图目录"),
        ("表目录", r"TOC \h \z \f T", "请更新域以生成表目录"),
    ]

    for title, field_instruction, placeholder in blocks:
        title_para = first_heading.insert_paragraph_before()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_text(title_para, title, bold=True, size_pt=16.0)

        field_para = first_heading.insert_paragraph_before()
        add_complex_field(field_para, field_instruction, placeholder)

        spacer = first_heading.insert_paragraph_before()
        spacer.add_run().add_break()


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = REPORT_PATH.with_name(
        f"my_report_backup_before_directories_{stamp}.docx"
    )
    shutil.copy2(REPORT_PATH, backup)

    doc = Document(REPORT_PATH)
    formula_count, figure_count, table_count = mark_entries(doc)
    insert_directory_block(doc)
    doc.save(REPORT_PATH)

    print(f"backup={backup}")
    print(f"updated={REPORT_PATH}")
    print(f"formula_entries={formula_count}")
    print(f"figure_entries={figure_count}")
    print(f"table_entries={table_count}")


if __name__ == "__main__":
    main()
