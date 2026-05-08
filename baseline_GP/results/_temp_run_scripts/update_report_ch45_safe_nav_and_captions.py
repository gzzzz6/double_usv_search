from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(r"F:\pythonprojects")
DOCX_PATH = ROOT / "my_report.docx"
BACKUP_PATH = ROOT / f"my_report_before_ch45_safe_nav_update_{datetime.now():%Y%m%d_%H%M%S}.docx"


def set_run_font(run, *, east_asia: str = "宋体", ascii_font: str = "Times New Roman", size_pt: float = 12.0) -> None:
    run.font.name = ascii_font
    run.font.size = Pt(size_pt)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    r_fonts.set(ns + "eastAsia", east_asia)
    r_fonts.set(ns + "ascii", ascii_font)
    r_fonts.set(ns + "hAnsi", ascii_font)


def format_paragraph(paragraph, *, size_pt: float = 12.0, align=None) -> None:
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        set_run_font(run, size_pt=size_pt)


def m_run(text: str):
    r = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    r.append(t)
    return r


def m_sub(base: str, sub: str):
    ssub = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    e.append(m_run(base))
    sub_el = OxmlElement("m:sub")
    sub_el.append(m_run(sub))
    ssub.append(e)
    ssub.append(sub_el)
    return ssub


def add_formula_before(anchor_paragraph, parts: list[object], number: str):
    paragraph = anchor_paragraph.insert_paragraph_before("", style="公式目录项")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath = OxmlElement("m:oMath")
    for part in parts:
        if isinstance(part, tuple) and len(part) == 2:
            omath.append(m_sub(part[0], part[1]))
        else:
            omath.append(m_run(str(part)))
    paragraph._p.append(omath)
    run = paragraph.add_run(f"    {number}")
    set_run_font(run, size_pt=10.5)
    return paragraph


def insert_paragraph_before(anchor_paragraph, text: str, style: str = "Normal", *, size_pt: float = 12.0):
    paragraph = anchor_paragraph.insert_paragraph_before(text, style=style)
    format_paragraph(paragraph, size_pt=size_pt)
    return paragraph


def find_heading(doc: Document, prefix: str, *, start: int = 0, style_name: str | None = None) -> int:
    for idx, paragraph in enumerate(doc.paragraphs[start:], start=start):
        if paragraph.text.strip().startswith(prefix) and (
            style_name is None or paragraph.style.name == style_name
        ):
            return idx
    raise RuntimeError(f"Cannot find heading starting with {prefix!r}")


def replace_exact_text(paragraph, text: str) -> None:
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
        set_run_font(paragraph.runs[0], size_pt=10.5 if paragraph.style.name in {"图目录项", "表目录项"} else 12.0)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size_pt=10.5 if paragraph.style.name in {"图目录项", "表目录项"} else 12.0)


def shift_formula_numbers(doc: Document, *, start_after_idx: int, start_num: int = 48) -> None:
    pattern = re.compile(r"^（(\d+)）$")
    for idx, paragraph in enumerate(doc.paragraphs):
        if idx <= start_after_idx:
            continue
        text = paragraph.text.strip()
        match = pattern.match(text)
        if not match:
            continue
        value = int(match.group(1))
        if value >= start_num:
            replace_exact_text(paragraph, f"（{value + 1}）")


def move_table_captions_below_tables(doc: Document) -> int:
    body = doc.element.body
    moved = 0
    changed = True
    while changed:
        changed = False
        children = list(body.iterchildren())
        for idx, child in enumerate(children[:-1]):
            if not isinstance(child, CT_P):
                continue
            para = Paragraph(child, doc)
            if para.style.name != "表目录项" or not para.text.strip().startswith("表"):
                continue
            next_child = children[idx + 1]
            if not isinstance(next_child, CT_Tbl):
                continue
            child.getparent().remove(child)
            next_child.addnext(child)
            moved += 1
            changed = True
            break
    return moved


def add_a_star_caption(doc: Document, *, section_start: int, section_end: int) -> None:
    pic_idx = None
    for idx in range(section_start, section_end):
        paragraph = doc.paragraphs[idx]
        if paragraph._p.xpath(".//pic:pic"):
            pic_idx = idx
            paragraph.style = doc.styles["Normal"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            break
    if pic_idx is None:
        raise RuntimeError("Cannot find A* image paragraph in 4.5.2.")

    next_text = doc.paragraphs[pic_idx + 1].text.strip() if pic_idx + 1 < len(doc.paragraphs) else ""
    if next_text.startswith("图4-4 A*算法"):
        return
    caption = doc.paragraphs[pic_idx + 1].insert_paragraph_before(
        "图4-4 A*算法在栅格地图中的搜索过程示意图",
        style="图目录项",
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(caption, size_pt=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)


def main() -> None:
    if not DOCX_PATH.exists():
        raise FileNotFoundError(DOCX_PATH)
    shutil.copy2(DOCX_PATH, BACKUP_PATH)

    doc = Document(DOCX_PATH)

    chapter4_idx = find_heading(doc, "4 信息型路径规划", style_name="一级标题")
    section45_idx = find_heading(doc, "4.5", start=chapter4_idx + 1, style_name="二级标题")
    h452_idx = find_heading(doc, "4.5.2", start=section45_idx + 1, style_name="三级标题")
    h453_idx = find_heading(doc, "4.5.3", start=h452_idx + 1, style_name="三级标题")
    h46_idx = find_heading(doc, "4.6", start=h453_idx + 1, style_name="二级标题")

    # Keep 4.5.2 body paragraphs in normal text style except true formula/image paragraphs.
    for idx in range(h452_idx + 1, h453_idx):
        paragraph = doc.paragraphs[idx]
        text = paragraph.text.strip()
        has_pic = bool(paragraph._p.xpath(".//pic:pic"))
        if has_pic:
            paragraph.style = doc.styles["Normal"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if text and not re.fullmatch(r"（\d+）", text):
            # The two displayed formulas in this area have empty text because they are OMML objects.
            if not paragraph._p.xpath(".//m:oMath"):
                paragraph.style = doc.styles["Normal"]
                format_paragraph(paragraph)

    add_a_star_caption(doc, section_start=h452_idx, section_end=h453_idx)

    # Re-resolve indices after inserting the A* caption.
    chapter4_idx = find_heading(doc, "4 信息型路径规划", style_name="一级标题")
    section45_idx = find_heading(doc, "4.5", start=chapter4_idx + 1, style_name="二级标题")
    h452_idx = find_heading(doc, "4.5.2", start=section45_idx + 1, style_name="三级标题")
    h453_idx = find_heading(doc, "4.5.3", start=h452_idx + 1, style_name="三级标题")

    # Shift formula labels from the original short-horizon formula onward.
    shift_formula_numbers(doc, start_after_idx=h453_idx - 1, start_num=48)

    # Insert safe-nav as the new 4.5.3.
    old_h453 = doc.paragraphs[h453_idx]
    insert_paragraph_before(old_h453, "4.5.3 Safe-nav A*安全路径生成", style="三级标题")
    insert_paragraph_before(
        old_h453,
        "在本文当前主线中，路径生成并非只使用普通 A* 的单位步长代价，而是在 A* 搜索框架中加入安全导航代价，形成 soft_clearance_astar_v1 路径生成方式。其基本思想是：障碍物仍作为不可通行区域参与可达性判断，而靠近障碍物但尚可通行的自由栅格则通过软净空代价进行惩罚，使规划器在存在多条可达路径时更倾向于选择远离障碍物的路径。",
    )
    insert_paragraph_before(
        old_h453,
        "设从当前节点 n 扩展到相邻节点 n'，普通 A* 中的基础移动代价为 1。加入 soft clearance 后，单步代价和 A* 评价函数可写为：",
    )
    add_formula_before(
        old_h453,
        [
            "c(n,n') = 1 + ",
            ("λ", "clear"),
            " ",
            ("C", "clear"),
            "(n'),    f(n') = [g(n) + c(n,n')] + h(n')",
        ],
        "（48）",
    )
    insert_paragraph_before(
        old_h453,
        "其中，C_clear(n') 表示相邻节点 n' 的软净空代价，λ_clear 表示净空代价权重。当某一自由栅格距离障碍物较近时，C_clear 较大，经过该格点的路径累计代价也会增大；当自由栅格远离障碍物时，该项较小，路径代价更接近普通 A* 的单位步长代价。因此，safe-nav 并没有改变 A* 的启发式搜索框架，而是改变了路径代价的定义，使路径生成在路径长度和障碍安全距离之间进行权衡。",
    )
    insert_paragraph_before(
        old_h453,
        "结合本文代码实现，当 path_safety_mode=\"soft_clearance_astar_v1\" 时，系统调用安全路径生成模块进行候选 viewpoint 的可达路径搜索。当前主线中 safe_nav_lambda_clearance=1.0，并采用软净空半径约束对贴近障碍的自由格点施加额外代价。因此，若多条路径都能到达同一候选 viewpoint，系统更倾向于选择障碍距离更安全的路径，而不是单纯选择几何步数最短的路径。",
    )
    insert_paragraph_before(
        old_h453,
        "需要强调的是，safe-nav A* 属于路径生成与执行安全层。它只改变候选 viewpoint 的可达路径形态和路径执行代价，不改变 GP clue、target intensity、recency 或综合信息价值场的语义。换言之，信息场回答“哪里值得搜索”，safe-nav A* 回答“如何更安全地到达候选观测位置”。最终执行哪一段路径，仍需在后续路径段评分中结合搜索信息收益、观测刷新收益、执行代价和掉头惩罚统一决定。",
    )

    # Original 4.5.3 becomes 4.5.4.
    old_h453 = find_heading(doc, "4.5.3 短时域路径段截断", start=h452_idx + 1, style_name="三级标题")
    replace_exact_text(doc.paragraphs[old_h453], "4.5.4 短时域路径段截断")
    doc.paragraphs[old_h453].style = doc.styles["三级标题"]

    # Existing Figure 4-4 becomes Figure 4-5 after adding the A* caption.
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "图4-4 候选路径段评分与最优路径段选择流程图":
            replace_exact_text(paragraph, "图4-5 候选路径段评分与最优路径段选择流程图")
            paragraph.style = doc.styles["图目录项"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif "图4-4 候选路径段评分与最优路径段选择流程图" in text:
            replace_exact_text(
                paragraph,
                paragraph.text.replace(
                    "图4-4 候选路径段评分与最优路径段选择流程图",
                    "图4-5 候选路径段评分与最优路径段选择流程图",
                ),
            )

    # Reference numbering affected by inserting LaValle as [34].
    for paragraph in doc.paragraphs:
        text = paragraph.text
        new_text = text.replace("[7,34-35]", "[7,35-36]").replace("[7,34]", "[7,35]")
        if new_text != text:
            replace_exact_text(paragraph, new_text)

    moved_tables = move_table_captions_below_tables(doc)

    doc.save(DOCX_PATH)
    print(f"updated={DOCX_PATH}")
    print(f"backup={BACKUP_PATH}")
    print(f"moved_table_captions={moved_tables}")


if __name__ == "__main__":
    main()
