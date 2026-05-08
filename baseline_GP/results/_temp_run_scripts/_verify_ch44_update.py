from pathlib import Path
from docx import Document
p=Path(r'F:\pythonprojects\my_report.docx')
doc=Document(p)
print('paragraph_count', len(doc.paragraphs))
for i, para in enumerate(doc.paragraphs):
    t=para.text.strip()
    if 470 <= i <= 510 or (t.startswith('4.4') or t.startswith('4.5')):
        print(f'{i}: style={para.style.name!r} text={t}')

# Body-only checks between real 4.4 and 4.5.
ch4 = next(i for i,p in enumerate(doc.paragraphs) if p.text.strip().startswith('4 信息型路径规划') and p.style.name=='一级标题')
sec44 = next(i for i,p in enumerate(doc.paragraphs[ch4+1:], start=ch4+1) if p.text.strip().startswith('4.4') and p.style.name=='二级标题')
sec45 = next(i for i,p in enumerate(doc.paragraphs[sec44+1:], start=sec44+1) if p.text.strip().startswith('4.5') and p.style.name=='二级标题')
body='\n'.join(p.text for p in doc.paragraphs[sec44:sec45])
print('\nBODY_RANGE', sec44, sec45)
print(body)
print('\nBODY_TOKEN_CHECK')
for token in ['simple_ring_v1', 'info_perturb', 'continuity_biased', 'current_segment_endpoint', 'simple_nearest_free', 'r_pref', 'I_bar(v)', '（45）', '（46）']:
    print(token, token in body)
print('\nWHOLE_DOC_BAD_MARKERS')
joined='\n'.join(p.text for p in doc.paragraphs)
for m in ['锛','鈥','????','�']:
    print(m, m in joined)
