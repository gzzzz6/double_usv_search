from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document


REPORT = Path(r"F:\pythonprojects\my_report.docx")
OUT = Path(r"F:\pythonprojects\baseline_GP\results\_temp_run_scripts\report_numbering_audit.json")


def main() -> None:
    doc = Document(REPORT)
    patterns = {
        "chapter_ref": re.compile(r"第\s*[一二三四五六123456]\s*章"),
        "section_ref_6": re.compile(r"(?:第\s*)?6\.\d+(?:\.\d+)?\s*节?"),
        "section_ref_5": re.compile(r"(?:第\s*)?5\.\d+(?:\.\d+)?\s*节?"),
        "fig_table_ref": re.compile(r"[图表][23456]-\d+"),
        "old_fig_table_ref_6": re.compile(r"[图表]6-\d+"),
    }
    findings: dict[str, list[dict[str, object]]] = {key: [] for key in patterns}
    headings: list[dict[str, object]] = []
    captions: list[dict[str, object]] = []

    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        style = p.style.name
        if style in {"一级标题", "二级标题", "三级标题"}:
            headings.append({"idx": i, "style": style, "text": text})
        if style in {"图目录项", "表目录项", "公式目录项"}:
            captions.append({"idx": i, "style": style, "text": text})
        if not text:
            continue
        for key, rx in patterns.items():
            hits = rx.findall(text)
            if hits:
                findings[key].append(
                    {"idx": i, "style": style, "hits": hits, "text": text}
                )

    payload = {
        "report": str(REPORT),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "headings": headings,
        "captions": captions,
        "findings": findings,
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUT)
    print(f"headings={len(headings)} captions={len(captions)}")
    for key, rows in findings.items():
        print(f"{key}={len(rows)}")
        for row in rows[:20]:
            print(f"  P{row['idx']:04d} {row['hits']} {row['text'][:180]}")


if __name__ == "__main__":
    main()
