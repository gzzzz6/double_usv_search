"""
Fix formatting of my_report.docx to match 论文写作规范.docx requirements.
"""
from docx import Document
from docx.shared import Pt, Cm, Emu
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree
import re
import shutil
import os

REPORT_PATH = r'F:\pythonprojects\my_report.docx'
BACKUP_PATH = r'F:\pythonprojects\my_report_format_backup_20260507.docx'

# Create backup
shutil.copy2(REPORT_PATH, BACKUP_PATH)
print(f'Backup saved to: {BACKUP_PATH}')

doc = Document(REPORT_PATH)

# ============================================================
# Helper: set East Asian font on a run
# ============================================================
def set_run_fonts(run, ascii_font='Times New Roman', ea_font=None, size=None, bold=None):
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), ascii_font)
    rFonts.set(qn('w:hAnsi'), ascii_font)
    rFonts.set(qn('w:cs'), ascii_font)
    if ea_font:
        rFonts.set(qn('w:eastAsia'), ea_font)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold

def set_run_size_bold(run, size=None, bold=None):
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold

# ============================================================
# 1. Fix page margins for sections 2 and 3
# ============================================================
print('\n=== 1. Fixing page margins ===')
for i, sec in enumerate(doc.sections):
    if i == 0:
        continue  # Skip cover section
    print(f'  Section {i+1}: before top={sec.top_margin}, bottom={sec.bottom_margin}, left={sec.left_margin}, right={sec.right_margin}')
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.0)
    sec.footer_distance = Cm(1.5)
    print(f'  Section {i+1}: after top={sec.top_margin}, bottom={sec.bottom_margin}, left={sec.left_margin}, right={sec.right_margin}')

# ============================================================
# 2. Fix Normal style
# ============================================================
print('\n=== 2. Fixing Normal style ===')
normal_style = doc.styles['Normal']
normal_style.font.name = 'Times New Roman'
normal_style.font.size = Pt(12)

# Set East Asian font in Normal style XML
style_elem = normal_style._element
rPr = style_elem.find(qn('w:rPr'))
if rPr is None:
    rPr = OxmlElement('w:rPr')
    style_elem.insert(0, rPr)
rFonts = rPr.find(qn('w:rFonts'))
if rFonts is None:
    rFonts = OxmlElement('w:rFonts')
    rPr.insert(0, rFonts)
rFonts.set(qn('w:eastAsia'), '宋体')

# Line spacing 1.25
normal_style.paragraph_format.line_spacing = 1.25
# First line indent
normal_style.paragraph_format.first_line_indent = Cm(0.74)

print(f'  Normal: font={normal_style.font.name} size={normal_style.font.size} line_spacing={normal_style.paragraph_format.line_spacing} first_line_indent={normal_style.paragraph_format.first_line_indent}')

# ============================================================
# Helper: check if paragraph is in TOC
# ============================================================
def is_toc_para(p):
    sname = p.style.name if p.style else ''
    if 'toc' in sname.lower():
        return True
    if sname in ('目录标题', '图目录标题', '公式目录标', '一级标题'):
        return True
    if re.search(r'\t\d+$', p.text.strip()):
        return True
    return False

# ============================================================
# 3. Fix chapter/section headings
# ============================================================
print('\n=== 3. Fixing headings ===')

heading_patterns = [
    (re.compile(r'^第[一二三四五六七八九十]+章'), 'chapter', Pt(16), True, True, Pt(24), Pt(18)),
    (re.compile(r'^\d+\.\d+\.\d+'), 'subsection (1.1.1)', Pt(12), True, False, Pt(6), Pt(6)),
    (re.compile(r'^\d+\.\d+'), 'section (1.1)', Pt(14), True, False, Pt(12), Pt(6)),
    (re.compile(r'^\d+ '), 'heading-1', Pt(16), True, True, Pt(24), Pt(18)),
]

heading_count = 0
for i, p in enumerate(doc.paragraphs):
    if is_toc_para(p):
        continue
    text = p.text.strip()
    if not text:
        continue

    for pat, label, size, bold, center, space_before, space_after in heading_patterns:
        if pat.match(text):
            heading_count += 1
            pf = p.paragraph_format

            if center:
                pf.alignment = 1  # CENTER
            else:
                pf.alignment = 0  # LEFT

            if space_before is not None:
                pf.space_before = space_before
            if space_after is not None:
                pf.space_after = space_after
            pf.first_line_indent = Emu(0)  # Remove indent for headings

            for run in p.runs:
                set_run_fonts(run, ascii_font='Times New Roman', ea_font='黑体', size=size, bold=bold)

            if label == 'chapter':
                print(f'  P{i}: [{label}] {text[:80]}')
            break

print(f'  Fixed {heading_count} heading paragraphs')

# ============================================================
# 4. Fix table/figure captions
# ============================================================
print('\n=== 4. Fixing table/figure captions ===')

fig_count = 0
tab_count = 0
for i, p in enumerate(doc.paragraphs):
    if is_toc_para(p):
        continue
    text = p.text.strip()
    if not text:
        continue

    is_fig = re.match(r'^图\s*\d+', text)
    is_tab = re.match(r'^表\s*\d+', text)

    if is_fig or is_tab:
        pf = p.paragraph_format
        pf.alignment = 1  # CENTER

        if is_fig:
            pf.space_before = Pt(6)
            pf.space_after = Pt(12)
            fig_count += 1
            if fig_count <= 3:
                print(f'  P{i} Fig: {text[:80]}')
        else:
            pf.space_before = Pt(6)
            pf.space_after = Pt(6)
            tab_count += 1
            if tab_count <= 3:
                print(f'  P{i} Tab: {text[:80]}')

        pf.first_line_indent = Emu(0)

        for run in p.runs:
            set_run_fonts(run, ascii_font='Times New Roman', ea_font='宋体', size=Pt(10.5), bold=None)

print(f'  Fixed {fig_count} figure captions, {tab_count} table captions')

# ============================================================
# 5. Fix references
# ============================================================
print('\n=== 5. Fixing references ===')

# Find references heading
ref_start = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == '参考文献':
        ref_start = i
        break

if ref_start:
    # Fix reference heading
    p = doc.paragraphs[ref_start]
    pf = p.paragraph_format
    pf.alignment = 1  # CENTER
    pf.space_before = Pt(24)
    pf.space_after = Pt(18)
    pf.first_line_indent = Emu(0)
    for run in p.runs:
        set_run_fonts(run, ascii_font='Times New Roman', ea_font='黑体', size=Pt(16), bold=True)
    print(f'  P{ref_start}: 参考文献 heading fixed')

    # Fix reference entries
    ref_count = 0
    for i in range(ref_start + 1, len(doc.paragraphs)):
        p = doc.paragraphs[i]
        text = p.text.strip()
        if not text:
            continue
        if text.startswith('附录') or text.startswith('致谢'):
            break

        if re.match(r'^\[\d+\]', text):
            ref_count += 1
            pf = p.paragraph_format
            pf.line_spacing = Pt(16)
            pf.first_line_indent = Emu(0)

            for run in p.runs:
                set_run_fonts(run, ascii_font='Times New Roman', ea_font='宋体', size=Pt(10.5), bold=None)

    print(f'  Fixed {ref_count} reference entries')
else:
    print('  WARNING: 参考文献 heading not found!')

# ============================================================
# 6. Fix abstract/dedication etc body paragraphs - ensure they have proper formatting
# ============================================================
print('\n=== 6. Fixing remaining body paragraphs ===')

# Find body text paragraphs that should have first-line indent and proper formatting
body_fixed = 0
for i, p in enumerate(doc.paragraphs):
    if is_toc_para(p):
        continue
    text = p.text.strip()
    if not text:
        continue

    # Skip already-handled headings and captions
    if re.match(r'^(第[一二三四五六七八九十]+章|\d+\.\d+\.\d+|\d+\.\d+|\d+ |图\s*\d+|表\s*\d+)', text):
        continue
    # Skip reference entries
    if re.match(r'^\[\d+\]', text):
        continue
    # Skip very short items (titles, signatures, etc.)
    if len(text) < 20:
        continue

    # Check if this is likely body text (Normal style with substantial content)
    pf = p.paragraph_format

    # Set first-line indent if not already set
    if pf.first_line_indent is None or pf.first_line_indent == 0:
        pf.first_line_indent = Cm(0.74)
        body_fixed += 1

    # Ensure runs have proper font
    for run in p.runs:
        if run.font.name is None or run.font.name == 'Times New Roman':
            set_run_fonts(run, ascii_font='Times New Roman', ea_font='宋体', size=Pt(12), bold=None)
        if run.font.size is None:
            run.font.size = Pt(12)

print(f'  Adjusted {body_fixed} body paragraphs (added first-line indent)')

# ============================================================
# Save
# ============================================================
print('\n=== Saving ===')
doc.save(REPORT_PATH)
print(f'Saved: {REPORT_PATH}')
print('Done!')
