from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


REPORT = Path(r"F:\pythonprojects\my_report.docx")


def insert_after(paragraph, text: str = "", style: str | None = None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    if style is not None:
        inserted.style = style
    if text:
        inserted.add_run(text)
    return inserted


def set_caption_format(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = REPORT.with_name(f"my_report_backup_before_ch5_fig_captions_{stamp}.docx")
    shutil.copy2(REPORT, backup)

    doc = Document(REPORT)

    existing = "\n".join(p.text for p in doc.paragraphs)
    if "图5-4 实验采用的三类已知静态障碍地图" not in existing:
        map_img_para = doc.paragraphs[626]
        p = insert_after(map_img_para, "图5-4 实验采用的三类已知静态障碍地图", "图目录项")
        set_caption_format(p)
        p = insert_after(
            p,
            (
                "图5-4给出了本文实验采用的三类已知静态障碍地图。三类地图均由自由水域与静态障碍物组成，"
                "其中白色区域表示 USV 可通行的 free 区域，黑色区域表示静态障碍物或地图边界。"
                "open_water 仅包含边界约束，用于表示障碍较少、路径可达性较高的开阔水域场景；"
                "obstacle_field 在水域中布置多个离散矩形障碍物，用于检验搜索策略在局部绕行和安全路径生成约束下的执行能力；"
                "peninsula_passage 包含半岛式障碍和局部通道结构，用于构造更明显的拓扑约束和绕行代价。"
            ),
            "Normal",
        )
        insert_after(
            p,
            (
                "通过设置这三类地图，本文能够在由简单到复杂的空间结构中比较搜索策略的表现："
                "open_water 主要考察信息驱动搜索本身的有效性，obstacle_field 重点考察障碍分布对路径生成和热点观测的影响，"
                "peninsula_passage 则用于分析通道结构对搜索完成时间和协同行为的限制。"
                "后续单艇基线实验、双艇协同实验和 anomaly-aware acquisition 消融实验均在上述地图条件下展开。"
            ),
            "Normal",
        )

    # Re-open after insertion because paragraph indices after 626 may have shifted.
    doc.save(REPORT)
    doc = Document(REPORT)
    existing = "\n".join(p.text for p in doc.paragraphs)
    if "图5-5 单 USV 在不同地图与目标运动模式下的 UCB 搜索轨迹" not in existing:
        six_img_idx = None
        for idx, paragraph in enumerate(doc.paragraphs):
            if 640 <= idx <= 660 and paragraph._p.xpath(".//w:drawing"):
                six_img_idx = idx
                break
        if six_img_idx is None:
            raise RuntimeError("Could not locate the 5.2 six-panel trajectory image.")
        six_img_para = doc.paragraphs[six_img_idx]
        p = insert_after(
            six_img_para,
            "图5-5 单 USV 在不同地图与目标运动模式下的 UCB 搜索轨迹",
            "图目录项",
        )
        set_caption_format(p)
        p = insert_after(
            p,
            (
                "图5-5给出了单 USV 在 UCB 采集函数下的六组典型运行轨迹，其中三列分别对应 open_water、"
                "obstacle_field 和 peninsula_passage 三类已知静态地图，上下两行分别对应 static 和 random_walk 两类目标运动模式。"
                "图中青色折线表示 USV 已执行轨迹，绿色线段表示当前提交执行的短时域路径段，蓝色圆点表示 USV 当前所在位置，"
                "黄色叉号表示当前候选 viewpoint，紫色菱形表示对应 anchor，星形符号表示目标状态。"
            ),
            "Normal",
        )
        p = insert_after(
            p,
            (
                "从图中可以看出，单艇搜索并不是沿固定覆盖路径机械扫描，而是随着 GP clue、target intensity 与 recency 状态的更新，"
                "在不同阶段不断调整搜索方向。在 open_water 场景中，USV 轨迹整体较为自由，主要体现为对高价值线索区域的逐步覆盖；"
                "在 obstacle_field 场景中，轨迹会围绕离散障碍产生绕行，说明 soft_clearance_astar_v1 安全路径生成机制能够在避开障碍的同时维持搜索推进；"
                "在 peninsula_passage 场景中，通道和半岛结构限制了可达路径，USV 需要沿可通行区域进行更长距离搜索。"
            ),
            "Normal",
        )
        insert_after(
            p,
            (
                "对比 static 与 random_walk 两类目标模式可以发现，目标运动会改变局部线索分布和后续搜索方向，"
                "但系统仍能通过滚动重规划持续更新候选观测位置并推进搜索过程。因此，该图从轨迹层面补充说明了表5-2和表5-3中的实验结果："
                "本文单艇 UCB 主线在不同地图结构和目标运动模式下均能够形成可执行的连续搜索过程，但其搜索效率会受到障碍分布和地图拓扑结构影响。"
            ),
            "Normal",
        )

    doc.save(REPORT)
    print(f"backup={backup}")
    print(f"updated={REPORT}")


if __name__ == "__main__":
    main()
