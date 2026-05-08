from pathlib import Path
from docx import Document
p=Path(r'F:\pythonprojects\my_report.docx')
doc=Document(p)
for start,end,title in [(414,438,'4.1'),(447,489,'4.3-4.4'),(490,524,'4.5-4.6')]:
    print('\n====', title, '====')
    for i in range(start,end+1):
        if i < len(doc.paragraphs):
            t=doc.paragraphs[i].text.strip()
            if t:
                print(f'{i}: {t}')
