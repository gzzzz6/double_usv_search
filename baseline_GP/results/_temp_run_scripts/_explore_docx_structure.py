"""Explore my_report2.docx structure: TOC, TOF, TOT, captions."""
from docx import Document
from docx.oxml.ns import qn
import re

DOC = r"F:\pythonprojects\my_report2.docx"
doc = Document(DOC)

print("=== SECTION / PARAGRAPH OVERVIEW ===\n")
for i, para in enumerate(doc.paragraphs):
    style = para.style.name if para.style else "None"
    text = para.text.strip()
    if not text:
        continue
    # Print only informative paragraphs
    if any(kw in text for kw in ["目录", "图", "表", "Figure", "Table", "TOC", "HYPERLINK", "PAGEREF", "图表", "插图"]):
        # Check for hyperlinks
        hyperlinks = para._element.findall(qn('w:hyperlink'))
        fld_simple = para._element.findall(qn('w:fldSimple'))
        fld_char = para._element.findall(qn('w:fldChar'))
        print(f"[P{i:04d}] style={style:30s} | hyperlinks={len(hyperlinks)} fldSimple={len(fld_simple)} fldChar={len(fld_char)}")
        print(f"         text: {text[:120]}")
        print()

print("\n=== HYPERLINKS IN DOCUMENT ===\n")
hyperlink_count = 0
for para in doc.paragraphs:
    for hl in para._element.findall(qn('w:hyperlink')):
        anchor = hl.get(qn('w:anchor'))
        target = hl.get(qn('w:target'))
        if hyperlink_count < 30:
            print(f"  hyperlink anchor={anchor} target={target}")
        hyperlink_count += 1
print(f"  Total hyperlinks: {hyperlink_count}")

print("\n=== PARAGRAPHS NEAR 图目录 / 表目录 ===")
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if "图目录" in text or "表目录" in text:
        print(f"\n  Found: P{i:04d} '{text}'")
        # Print next 30 paragraphs
        for j in range(i, min(i + 30, len(doc.paragraphs))):
            t = doc.paragraphs[j].text.strip()
            if t:
                print(f"    P{j:04d}: {t[:150]}")
