from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm


PROJECT_ROOT = Path("F:/pythonprojects")
DOCX_PATH = PROJECT_ROOT / "my_report.docx"
FIG4_1_PATH = PROJECT_ROOT / "images" / "fig4_1_static_responsibility_prior_current_maps.png"
FIG4_2_PATH = PROJECT_ROOT / "images" / "fig4_2_residual_map_sequential_allocation_obstacle_field.png"


def _clear_paragraph(paragraph) -> None:
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def _replace_paragraph_with_picture(paragraph, image_path: Path, width_cm: float) -> None:
    _clear_paragraph(paragraph)
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))


def _replace_paragraph_text(paragraph, text: str) -> None:
    _clear_paragraph(paragraph)
    paragraph.style = "Normal"
    paragraph.alignment = None
    paragraph.add_run(text)


def _find_caption_index(doc: Document, caption_text: str) -> int:
    for idx, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == caption_text:
            return idx
    raise RuntimeError(f"Caption not found: {caption_text}")


def _replace_figure_before_caption(
    doc: Document,
    *,
    caption_text: str,
    image_path: Path,
    width_cm: float,
    following_text: str,
) -> dict[str, object]:
    caption_idx = _find_caption_index(doc, caption_text)
    image_idx = caption_idx - 1
    if image_idx < 0:
        raise RuntimeError(f"No paragraph before caption: {caption_text}")

    image_paragraph = doc.paragraphs[image_idx]
    had_drawing = bool(image_paragraph._p.xpath(".//w:drawing"))
    _replace_paragraph_with_picture(image_paragraph, image_path, width_cm=width_cm)

    text_idx = caption_idx + 1
    if text_idx >= len(doc.paragraphs):
        raise RuntimeError(f"No following explanation paragraph after caption: {caption_text}")
    _replace_paragraph_text(doc.paragraphs[text_idx], following_text)

    return {
        "caption": caption_text,
        "caption_paragraph_index": caption_idx,
        "image_paragraph_index": image_idx,
        "image_paragraph_had_drawing": had_drawing,
        "explanation_paragraph_index": text_idx,
        "image_path": str(image_path),
    }


def main() -> None:
    if not DOCX_PATH.exists():
        raise FileNotFoundError(DOCX_PATH)
    if not FIG4_1_PATH.exists():
        raise FileNotFoundError(FIG4_1_PATH)
    if not FIG4_2_PATH.exists():
        raise FileNotFoundError(FIG4_2_PATH)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = DOCX_PATH.with_name(f"my_report_backup_before_ch4_figures_{timestamp}.docx")
    shutil.copy2(DOCX_PATH, backup_path)

    doc = Document(DOCX_PATH)
    replacements = []
    replacements.append(
        _replace_figure_before_caption(
            doc,
            caption_text="图4-1 静态软责任区先验可视化",
            image_path=FIG4_1_PATH,
            width_cm=15.8,
            following_text=(
                "图4-1给出了 open_water、obstacle_field 和 peninsula_passage 三张正式实验地图上的静态软责任区先验。"
                "蓝色区域表示第0艘USV责任分数为正的区域，橙色区域表示第1艘USV责任分数占优的区域，"
                "灰色区域表示距离差较小的 buffer band。可以看到，责任边界由两艇固定起点到自由格的已知地图可达距离决定；"
                "在障碍或通道结构存在时，责任区域随可航空间发生变形，而不是简单的欧氏直线分割。"
            ),
        )
    )
    replacements.append(
        _replace_figure_before_caption(
            doc,
            caption_text="图4-2 residual_map 顺序分配过程",
            image_path=FIG4_2_PATH,
            width_cm=15.8,
            following_text=(
                "图4-2给出了 obstacle_field 场景下一次 coordinated 分配中的 residual_map 变化。"
                "第一艇在共享综合信息价值场上规划后，其路径段可见区域被从临时残差信息图中置零；"
                "第二艇随后在残差信息图上规划，从而降低对已分配观测区域的重复关注。"
                "该过程只改变团队分配阶段的临时排序图，不改变共享综合信息价值场本身。"
            ),
        )
    )
    doc.save(DOCX_PATH)

    report_path = PROJECT_ROOT / "baseline_GP" / "results" / "_temp_run_scripts" / "replace_ch4_figures_4_1_4_2_report.txt"
    report_path.write_text(
        "\n".join(
            [
                f"backup={backup_path}",
                *(str(item) for item in replacements),
            ]
        ),
        encoding="utf-8",
    )
    print(f"backup={backup_path}")
    print(f"report={report_path}")
    for item in replacements:
        print(item)


if __name__ == "__main__":
    main()
