from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


DOC_PATH = Path(r"F:\pythonprojects\my_report.docx")
ASCII_RE = re.compile(r"([A-Za-z0-9_./+\-=<>%]+)")

HEADERS = [
    "地图",
    "方法",
    "success_all_found（%）",
    "T_first",
    "T_all",
    "detection_rate（%）",
    "found_count",
    "obs_ratio（%）",
    "duplicate（%）",
    "cross_region（%）",
]

PARAM_ROWS = [
    ["实验对象", "双 USV 已知静态地图搜索"],
    ["比较方法", "two_usv_independent / two_usv_coordinated"],
    ["地图尺寸", "60×80"],
    ["地图类型", "open_water / harbor_cove / peninsula_passage"],
    ["目标模式", "static / random_walk"],
    ["随机种子", "0–9"],
    ["最大仿真步数", "240"],
    ["policy", "marine_knownmap_path_v2_infosampled_2usv"],
    ["clue_acquisition_mode", "ucb"],
    ["path_safety_mode", "soft_clearance_astar_v1"],
    ["team_path_avoidance_mode", "reservation_v1"],
    [
        "协同差异",
        "independent 不使用 residual_map 顺序分配；coordinated 使用 sequential allocation-lite 与重叠抑制",
    ],
]

STATIC_ROWS = [
    ["总体", "independent", "26.7±44.2", "83.5±62.4", "128.8±67.4", "64.4±25.7", "1.93±0.77", "62.0±15.0", "2.07±3.02", "29.61±16.61"],
    ["总体", "coordinated", "56.7±49.6", "61.4±47.8", "157.5±61.0", "80.0±28.0", "2.40±0.84", "67.6±16.7", "0.01±0.07", "1.47±1.46"],
    ["open_water", "independent", "40.0±49.0", "81.6±70.3", "112.2±41.2", "76.7±21.3", "2.30±0.64", "53.8±12.9", "1.47±2.61", "23.52±12.14"],
    ["open_water", "coordinated", "70.0±45.8", "65.2±46.9", "161.4±62.6", "90.0±15.3", "2.70±0.46", "63.9±16.6", "0.00±0.00", "1.54±1.54"],
    ["harbor_cove", "independent", "30.0±45.8", "68.0±55.0", "114.7±73.3", "63.3±27.7", "1.90±0.83", "70.0±18.5", "1.25±2.59", "20.28±15.11"],
    ["harbor_cove", "coordinated", "50.0±50.0", "52.0±49.7", "117.6±59.6", "66.7±39.4", "2.00±1.18", "69.9±21.7", "0.04±0.12", "1.36±1.70"],
    ["peninsula_passage", "independent", "10.0±30.0", "100.9±56.3", "237.0±0.0", "53.3±22.1", "1.60±0.66", "62.3±5.6", "3.50±3.27", "45.02±9.50"],
    ["peninsula_passage", "coordinated", "50.0±50.0", "65.2±46.1", "191.8±29.0", "83.3±16.7", "2.50±0.50", "68.9±8.3", "0.00±0.00", "1.52±1.04"],
]

RANDOM_ROWS = [
    ["总体", "independent", "36.7±48.2", "76.1±58.1", "153.6±48.6", "73.3±23.4", "2.20±0.70", "71.3±14.2", "3.83±5.30", "26.21±14.16"],
    ["总体", "coordinated", "60.0±49.0", "61.1±48.9", "159.9±47.7", "83.3±22.4", "2.50±0.67", "73.2±15.3", "0.00±0.00", "2.27±2.41"],
    ["open_water", "independent", "40.0±49.0", "59.7±47.9", "158.0±42.8", "73.3±24.9", "2.20±0.75", "67.3±10.7", "4.86±4.53", "31.43±12.51"],
    ["open_water", "coordinated", "60.0±49.0", "64.8±54.1", "164.3±46.9", "83.3±22.4", "2.50±0.67", "72.3±14.2", "0.00±0.00", "1.75±2.15"],
    ["harbor_cove", "independent", "50.0±50.0", "68.9±50.3", "144.4±58.7", "80.0±22.1", "2.40±0.66", "73.6±18.5", "0.96±1.67", "14.02±11.66"],
    ["harbor_cove", "coordinated", "60.0±49.0", "53.1±41.7", "159.2±55.4", "83.3±22.4", "2.50±0.67", "77.4±17.5", "0.00±0.00", "1.42±1.71"],
    ["peninsula_passage", "independent", "20.0±40.0", "99.6±66.7", "168.0±16.0", "66.7±21.1", "2.00±0.63", "73.1±11.2", "5.66±6.94", "33.19±9.17"],
    ["peninsula_passage", "coordinated", "60.0±49.0", "65.5±49.2", "156.2±39.1", "83.3±22.4", "2.50±0.67", "69.9±12.7", "0.00±0.00", "3.62±2.66"],
]


def clear_content(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_run_font(run, english=False, size=10.5, bold=None):
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.name = "Times New Roman" if english else "宋体"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def set_mixed_text(paragraph, text, size=10.5, bold=None):
    clear_content(paragraph)
    for part in ASCII_RE.split(text):
        if not part:
            continue
        run = paragraph.add_run(part)
        set_run_font(run, english=bool(ASCII_RE.fullmatch(part)), size=size, bold=bold)


def insert_paragraph(anchor, text, style=None, align=None, size=10.5):
    paragraph = anchor.insert_paragraph_before()
    if style is not None:
        paragraph.style = style
    if align is not None:
        paragraph.alignment = align
    set_mixed_text(paragraph, text, size=size)
    return paragraph


def insert_table(anchor, headers, rows, font_size=8.0):
    doc = anchor.part.document
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Normal Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_mixed_text(paragraph, header, size=font_size, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_mixed_text(paragraph, str(value), size=font_size)
    anchor._p.addprevious(table._tbl)
    return table


def locate_and_clear_old_63(doc):
    body = doc.element.body
    children = list(body.iterchildren())
    start = end = None
    anchor = None
    style_h2 = style_h3 = None
    for idx, child in enumerate(children):
        if not child.tag.endswith("}p"):
            continue
        paragraph = Paragraph(child, doc)
        text = paragraph.text.strip()
        if idx > 500 and text.startswith("6.3 "):
            start = idx
            style_h2 = paragraph.style
        elif idx > 500 and start is not None and text.startswith("6.3.1 "):
            style_h3 = paragraph.style
        elif idx > 500 and start is not None and text.startswith("6.4 "):
            end = idx
            anchor = paragraph
            break
    if start is None or end is None or anchor is None:
        raise RuntimeError("Could not locate body section 6.3/6.4")
    if style_h3 is None:
        raise RuntimeError("Could not capture level-3 heading style")
    for child in children[start + 1 : end]:
        body.remove(child)
    return anchor, style_h2, style_h3


def replace_64_table_numbers(doc):
    replacements = [
        ("表6-7", "表6-10"),
        ("表6-6", "表6-9"),
        ("表6-5", "表6-8"),
        ("表6-4", "表6-7"),
    ]
    in_64 = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("6.4 "):
            in_64 = True
        elif text.startswith("6.5 "):
            in_64 = False
        if not in_64:
            continue
        updated = paragraph.text
        for old, new in replacements:
            updated = updated.replace(old, new)
        if updated != paragraph.text:
            size = 10.5
            if paragraph.style.name in {"二级标题", "三级标题"}:
                size = 12.0
            set_mixed_text(paragraph, updated, size=size)


def add_63(anchor, style_h3):
    insert_paragraph(anchor, "6.3.1 实验目的", style_h3)
    insert_paragraph(anchor, "本节实验用于验证第 5 章提出的双 USV 协同搜索决策机制是否能够在相同双艇平台与相同已知静态地图设置下提高团队搜索效果。与第 6.2 节不同，本节不将单 USV 与双 USV 进行直接数值比较，原因是单艇实验采用 60×40 地图，而双艇实验采用 60×80 地图；二者搜索面积和目标分布空间不同，不能作为严格公平的同表对比对象。")
    insert_paragraph(anchor, "因此，本节只比较 two_usv_independent 与 two_usv_coordinated 两类双艇方法。前者表示两艇在共享地图、共享 GP clue、target intensity、recency 和综合信息价值场的条件下分别独立规划路径段；后者在相同共享团队状态上进一步引入 sequential allocation-lite、residual_map、overlap penalty 与 same-viewpoint penalty，用于降低两艇对同一高价值区域的重复选择。")

    insert_paragraph(anchor, "6.3.2 实验设置", style_h3)
    insert_paragraph(anchor, "本节使用 open_water、harbor_cove 和 peninsula_passage 三类 60×80 已知静态地图，每类地图使用 0–9 共 10 个随机种子，最大仿真步数为 240。目标运动模式分别设置为 static 和 random_walk。为保证本节只考察协同机制本身，所有结果均固定 clue_acquisition_mode=ucb，不引入 anomaly_upper_tail 或 ucb_anomaly_conditional 采集变体。")
    insert_paragraph(anchor, "两类双艇方法均使用 marine_knownmap_path_v2_infosampled_2usv 策略、soft_clearance_astar_v1 路径安全模式以及 reservation_v1 艇间冲突安全执行层。也就是说，independent 与 coordinated 的差别不在地图、目标、线索采集、避障或执行安全设置，而在于团队分配层是否使用 residual_map 与重叠抑制机制进行协同路径段选择。表6-4给出了本节主要实验参数。")
    insert_paragraph(anchor, "表6-4 双艇协同实验参数设置", align=WD_ALIGN_PARAGRAPH.CENTER)
    insert_table(anchor, ["参数项", "设置"], PARAM_ROWS, font_size=9.0)

    insert_paragraph(anchor, "6.3.3 评价指标与结果组织方式", style_h3)
    insert_paragraph(anchor, "本节指标分为三类。第一类为搜索效率指标，包括 time_to_first_detection 与 time_to_all_found，分别表示首次发现任一目标和发现全部目标所需步数。由于部分 episode 未能在 240 步预算内发现全部目标，time_to_all_found 只对成功完成全部目标搜索的 episode 统计均值与标准差。")
    insert_paragraph(anchor, "第二类为搜索完成程度指标，包括 success_all_found、detection_rate 和 found_count。success_all_found 表示在最大步数内发现全部目标的比例，detection_rate 和 found_count 则刻画仿真结束时的最终发现程度。第三类为协同行为指标，包括 known_free_observation_ratio_final、duplicate_viewpoint_ratio、cross_region_assignment_ratio、path_length_total 和 reservation_wait_fallback_count 等。其中 duplicate_viewpoint_ratio 反映两艇是否选择重复视点，cross_region_assignment_ratio 反映路径段选择是否频繁跨越软责任区先验。")
    insert_paragraph(anchor, "结果组织上，本文分别给出 static 与 random_walk 两类目标模式下的 independent 与 coordinated 对比表。表中 T_first 表示 time_to_first_detection，T_all 表示 time_to_all_found，obs_ratio 表示 known_free_observation_ratio_final。")

    insert_paragraph(anchor, "6.3.4 静态目标模式下的双艇协同搜索结果", style_h3)
    insert_paragraph(anchor, "表6-5 静态目标模式下双艇 independent 与 coordinated 对比结果", align=WD_ALIGN_PARAGRAPH.CENTER)
    insert_table(anchor, HEADERS, STATIC_ROWS, font_size=7.2)
    insert_paragraph(anchor, "由表6-5可见，在 static 目标模式下，coordinated 的总体 success_all_found 为 56.7%，高于 independent 的 26.7%；detection_rate 由 64.4% 提升到 80.0%，found_count 由 1.93 提升到 2.40。同时，coordinated 的总体首次发现时间为 61.4 步，低于 independent 的 83.5 步，说明协同分配机制能够更早形成有效目标发现。")
    insert_paragraph(anchor, "需要注意的是，coordinated 的总体 T_all 为 157.5 步，高于 independent 的 128.8 步。该现象不能简单解释为 coordinated 更慢，因为 independent 的 success_all_found 只有 26.7%，T_all 只来自少量成功完成的 episode；coordinated 使更多困难 episode 进入成功集合，因此完成全部目标的平均时间可能上升。结合 success_all_found、detection_rate 和 found_count，可以更准确地认为 coordinated 提高了静态目标下的整体搜索可靠性。")
    insert_paragraph(anchor, "从协同行为看，coordinated 的 duplicate_viewpoint_ratio 由 independent 的 2.07% 降至 0.01%，cross_region_assignment_ratio 由 29.61% 降至 1.47%。这说明 sequential allocation-lite 与 residual_map 的作用确实进入了团队路径段选择过程，使两艇更少选择重复视点，也更少偏离由固定起点形成的软责任区先验。reservation_wait_fallback_count 在两类方法中均为 0，说明本组实验中 reservation_v1 没有引入明显等待代价。")

    insert_paragraph(anchor, "6.3.5 随机游走目标模式下的双艇协同搜索结果", style_h3)
    insert_paragraph(anchor, "表6-6 随机游走目标模式下双艇 independent 与 coordinated 对比结果", align=WD_ALIGN_PARAGRAPH.CENTER)
    insert_table(anchor, HEADERS, RANDOM_ROWS, font_size=7.2)
    insert_paragraph(anchor, "由表6-6可见，在 random_walk 目标模式下，coordinated 仍保持了较好的协同收益。其总体 success_all_found 为 60.0%，高于 independent 的 36.7%；detection_rate 从 73.3% 提升到 83.3%，found_count 从 2.20 提升到 2.50。首次发现时间方面，coordinated 的总体 T_first 为 61.1 步，低于 independent 的 76.1 步。")
    insert_paragraph(anchor, "与 static 模式类似，random_walk 下 coordinated 的总体 T_all 为 159.9 步，略高于 independent 的 153.6 步。但该差异需要结合完成率理解：coordinated 完成全部目标搜索的 episode 比例更高，因而其 T_all 覆盖了更多复杂样本。若只看成功样本完成时间，容易低估 independent 在未完成 episode 上的失败代价。")
    insert_paragraph(anchor, "协同行为指标进一步说明了 coordinated 的实际作用。random_walk 下 independent 的 duplicate_viewpoint_ratio 为 3.83%，cross_region_assignment_ratio 为 26.21%；coordinated 则分别降至 0.00% 和 2.27%。这表明即使目标位置发生轻微随机游走，基于 residual_map 的顺序分配仍能抑制两艇对同一热点区域的重复追踪，并维持较稳定的空间分工。")

    insert_paragraph(anchor, "6.3.6 不同地图结构下的协同效果分析", style_h3)
    insert_paragraph(anchor, "在 open_water 场景中，障碍较少，两艇路径可达性较好，coordinated 的主要作用表现为降低重复选择并提高搜索完成程度。static 模式下，open_water 的 success_all_found 由 independent 的 40.0% 提升至 coordinated 的 70.0%，detection_rate 由 76.7% 提升至 90.0%；random_walk 模式下，success_all_found 由 40.0% 提升至 60.0%，detection_rate 由 73.3% 提升至 83.3%。这说明开阔地图中协同分配更容易转化为互补覆盖收益。")
    insert_paragraph(anchor, "在 harbor_cove 场景中，凹形岸线和港湾入口会使两艇容易被相近热点吸引。coordinated 在 static 模式下将 success_all_found 从 30.0% 提升到 50.0%，在 random_walk 模式下从 50.0% 提升到 60.0%；同时 duplicate_viewpoint_ratio 几乎降为 0。该结果说明 overlap penalty 与 same-viewpoint penalty 对港湾结构中的重复观测具有抑制作用，但由于港湾拓扑本身限制可达路径，完成时间收益并不总是同步扩大。")
    insert_paragraph(anchor, "在 peninsula_passage 场景中，通道与半岛结构对双艇协同的影响最明显。static 模式下，independent 的 success_all_found 仅为 10.0%，coordinated 提升至 50.0%，detection_rate 也由 53.3% 提升至 83.3%；random_walk 模式下，success_all_found 由 20.0% 提升至 60.0%。同时，coordinated 在该地图下的 T_all 低于 independent，说明在通道型地图中，显式团队分配能够减少两艇在同一通道或同一热点附近重复搜索的代价。")
    insert_paragraph(anchor, "从路径与执行代价看，coordinated 的总体 path_length_total 在 static 模式下为 386.5，低于 independent 的 420.7；在 random_walk 模式下为 383.9，同样低于 independent 的 416.6。两类目标模式下 reservation_wait_fallback_count 均为 0，说明本节观察到的协同收益主要来自分配层的路径段选择差异，而不是通过增加等待或冲突规避代价换取的结果。")

    insert_paragraph(anchor, "6.3.7 本节小结", style_h3)
    insert_paragraph(anchor, "本节在相同 60×80 双 USV 已知地图设置下，比较了 UCB 线索采集模式下的 independent 与 coordinated 两类双艇搜索方法。由于单艇实验采用 60×40 地图，本节没有将 single USV 与双艇结果进行直接数值对比，而是将实验变量限定为双艇团队分配机制本身。")
    insert_paragraph(anchor, "实验结果表明，coordinated 在 static 与 random_walk 两类目标模式下均提高了 success_all_found、detection_rate 和 found_count，并显著降低 duplicate_viewpoint_ratio 与 cross_region_assignment_ratio。这说明 residual_map 顺序分配、重叠惩罚和相同视点惩罚确实改变了两艇路径段选择，使团队搜索更趋向互补分工。与此同时，coordinated 的 time_to_all_found 并非在所有总体统计中都低于 independent，原因是其成功完成全部目标搜索的 episode 更多，T_all 统计集合也包含了更多困难样本。因此，本节结论应表述为 coordinated 提高了双艇搜索完成可靠性和协同覆盖质量，而不是无条件缩短所有完成时间指标。")


def main():
    doc = Document(DOC_PATH)
    anchor, _style_h2, style_h3 = locate_and_clear_old_63(doc)
    add_63(anchor, style_h3)
    replace_64_table_numbers(doc)
    doc.save(DOC_PATH)
    print(DOC_PATH)


if __name__ == "__main__":
    main()
