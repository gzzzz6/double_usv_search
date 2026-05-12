from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document


REPORT = Path(r"F:\pythonprojects\my_report.docx")
PAYLOAD = Path(
    r"F:\pythonprojects\baseline_GP\results\paper_obstacle_field_mainline_20260510"
    r"\paper_tables_merged_for_report\table_payload.json"
)


def set_paragraph_text(paragraph, text: str) -> None:
    """Replace paragraph text while keeping paragraph style and first-run formatting."""
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_cell_text(cell, text: str) -> None:
    """Replace cell text while preserving the cell's first paragraph style."""
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text)
    for extra in cell.paragraphs[1:]:
        set_paragraph_text(extra, "")


def replace_in_paragraph(paragraph, old: str, new: str) -> None:
    text = paragraph.text
    if old in text:
        set_paragraph_text(paragraph, text.replace(old, new))


def replace_table(table, rows: list[list[str]]) -> None:
    if len(table.rows) != len(rows):
        raise AssertionError(f"table row mismatch: doc={len(table.rows)}, payload={len(rows)}")
    for r_idx, row_values in enumerate(rows):
        if len(table.rows[r_idx].cells) != len(row_values):
            raise AssertionError(
                f"table col mismatch at row {r_idx}: doc={len(table.rows[r_idx].cells)}, payload={len(row_values)}"
            )
        for c_idx, value in enumerate(row_values):
            set_cell_text(table.rows[r_idx].cells[c_idx], value)


def main() -> None:
    if not REPORT.exists():
        raise FileNotFoundError(REPORT)
    payload = json.loads(PAYLOAD.read_text(encoding="utf-8"))

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = REPORT.with_name(f"my_report_backup_before_obstacle_field_update_{stamp}.docx")
    shutil.copy2(REPORT, backup)

    doc = Document(REPORT)

    # The experiment chapter has been moved to Chapter 5. Only adjust table references
    # inside the experiment chapter range.
    for idx in range(642, 714):
        replace_in_paragraph(doc.paragraphs[idx], "表6-", "表5-")

    paragraph_updates = {
        661: (
            "由表5-2可见，在 static 目标模式下，单艇 UCB 基线的总体首次发现时间为 47.0 步，"
            "成功完成全部目标搜索的样本中 T_all 均值为 153.7 步；n_all=15，对应 success_all_found 为 50.0%。"
            "最终 detection_rate 为 83.3%，说明在 240 步预算内单艇能够稳定发现大部分目标，但仍存在部分随机种子未能完成全部目标搜索。"
        ),
        662: (
            "分地图看，obstacle_field 的首次发现时间为 38.5 步，T_all 为 117.3 步，"
            "success_all_found 为 60.0%，是三类地图中完成效率较高的场景。该结果说明离散障碍会带来局部绕行，"
            "但不会像狭窄通道那样形成持续瓶颈；peninsula_passage 的 T_all 最高，为 192.2 步，且 path_length 达到 220.9，"
            "反映出半岛和通道结构会显著增加绕行与覆盖代价。"
        ),
        665: (
            "由表5-3可见，在 random_walk 目标模式下，单艇 UCB 基线的总体首次发现时间为 60.5 步，"
            "T_all 均值为 145.6 步，n_all=19，对应 success_all_found 为 63.3%。与 static 模式相比，"
            "首次发现时间有所增加，但完成全部目标搜索的 episode 数量和最终 detection_rate 均略有提高，"
            "说明轻微目标运动并未破坏单艇信息驱动搜索主链。"
        ),
        666: (
            "在 open_water 中，random_walk 模式下 n_all 达到 7，success_all_found 为 70.0%，但 T_first 也升至 71.1 步，"
            "表明开阔地图中的早期搜索方向更依赖当前线索排序。obstacle_field 的 T_first 为 54.7 步，T_all 为 138.0 步，"
            "说明离散障碍场景下安全路径生成仍能保持较高完成效率；peninsula_passage 的 T_all 为 176.7 步，"
            "继续体现通道型拓扑对搜索完成时间的影响。"
        ),
        668: (
            "综合 static 与 random_walk 两类结果可以看出，地图结构对单艇搜索行为的影响较为稳定。"
            "open_water 障碍少，路径更直接，但首次发现时间对目标运动更敏感；obstacle_field 包含多个离散障碍，"
            "会引入局部绕行和安全净空代价，但整体搜索完成效率仍优于狭窄通道场景。"
        ),
        676: (
            "本节使用 open_water、obstacle_field 和 peninsula_passage 三类 60×80 已知静态地图，每类地图使用 0–9 共 10 个随机种子，"
            "最大仿真步数为 240。所有结果均固定 clue_acquisition_mode=ucb，以便将实验变量限定为双艇团队分配机制本身。"
        ),
        681: (
            "由表5-5和表5-6可见，在 static 目标模式下，coordinated 相比 independent 的总体 T_first 降低 22.6 步，"
            "T_all 成对差值降低 4.8 步，success_all_found 提高 20.0 个百分点，detection_rate 提高 16.7 个百分点。"
            "这说明顺序分配与重叠抑制机制能够更早形成有效目标发现，并提高静态目标条件下的整体搜索完成程度。"
        ),
        682: (
            "协同行为指标也显示 coordinated 的实际作用：duplicate_viewpoint_ratio 平均降低 4.9 个百分点，"
            "cross_region_assignment_ratio 平均降低 30.5 个百分点，wait_count_total 平均减少 0.7。"
            "peninsula_passage 中 ΔT_all 为 N/A，这是因为成对差值统计中有效完成全部目标的样本不足，不应强行解释为具体时间收益。"
        ),
        685: (
            "由表5-7和表5-8可见，在 random_walk 目标模式下，coordinated 的总体 T_first 比 independent 降低 9.2 步，"
            "T_all 降低 5.4 步，success_all_found 提高 10.0 个百分点，detection_rate 提高 7.8 个百分点，"
            "说明协同机制在动态目标条件下仍能改善总体完成程度。但 obstacle_field 中 T_first 和 T_all 均有所增加，"
            "表明离散障碍环境下的分工收益会受到局部绕行和路径安全约束影响。"
        ),
        686: (
            "从协同行为看，random_walk 模式下 coordinated 仍显著降低 duplicate_viewpoint_ratio 和 cross_region_assignment_ratio，"
            "二者分别下降 3.8 和 33.5 个百分点。该结果说明 residual_map 顺序分配仍然有效抑制重复搜索，"
            "但动态目标带来的线索变化和地图拓扑约束会共同影响协同分配对最终全部发现成功率的提升幅度。"
        ),
        689: (
            "obstacle_field 中多个离散障碍会改变两艇从起点到候选观测区域的可达路径。coordinated 在 static 模式下使 T_first 降低 4.6 步，"
            "T_all 降低 7.5 步，并将 success_all_found 提高 20.0 个百分点；在 random_walk 模式下，虽然 T_first 和 T_all 分别增加 4.6 步和 27.5 步，"
            "但 success_all_found 与 detection_rate 仍分别提高 20.0 和 10.0 个百分点。该结果表明离散障碍场景中 coordinated 能减少重复与跨区偏移，"
            "但其时间收益会受到目标运动和局部绕行代价影响。"
        ),
        692: (
            "本节在相同 60×80 双 USV 已知地图设置下，比较了 UCB 线索采集模式下的 independent 与 coordinated 两类双艇搜索方法。"
            "实验表明，coordinated 的稳定收益主要体现在降低重复观测点选择、减少跨责任区偏移并改善早期发现效率；"
            "但在 random_walk 目标模式和部分障碍场景中，全部发现时间并不总是低于 independent。"
        ),
        700: (
            "由表5-10至表5-12可见，单艇下 anomaly_upper_tail 的效果具有条件性。static 模式中，anomaly_upper_tail 将 success_all_found 从 50.0% 提升至 73.3%，"
            "detection_rate 从 83.3% 提升至 88.9%，但 T_all 成对差值为 -4.7 步，时间收益并不显著。random_walk 模式中，"
            "anomaly_upper_tail 将 success_all_found 提高 6.7 个百分点，并使 T_first 降低 13.6 步，但 T_all 差值为 +20.5 步，"
            "说明其能够改善早期发现和完成可靠性，却未必缩短成功样本中的全部发现时间。"
        ),
        702: (
            "由表5-13至表5-15可见，在双艇 coordinated 系统中，static 模式下 anomaly_upper_tail 使 T_first 降低 5.2 步，"
            "但 T_all 成对差值为 +1.8 步，success_all_found 下降 3.3 个百分点，detection_rate 基本不变。random_walk 模式下，"
            "anomaly_upper_tail 将 success_all_found 从 53.3% 提升至 66.7%，detection_rate 从 83.3% 提升至 85.6%，并使 T_all 降低 29.4 步，"
            "但 T_first 略有上升。"
        ),
        704: (
            "从地图维度看，anomaly_upper_tail 的收益不是单调一致的。open_water 障碍少，上尾异常线索更容易转化为可执行路径收益；"
            "obstacle_field 中 anomaly 在 random_walk 条件下更容易改善早期发现或全部发现时间，但在 static 条件下对成功率和时间指标的影响并不一致；"
            "peninsula_passage 受通道拓扑限制，即使异常线索改变热点排序，也不一定能显著缩短 T_all。"
        ),
    }

    for idx, text in paragraph_updates.items():
        set_paragraph_text(doc.paragraphs[idx], text)

    # Parameter table in 5.3.
    set_cell_text(doc.tables[1].rows[4].cells[1], "open_water、obstacle_field、peninsula_passage")

    table_mapping = {
        2: "two_random_delta",
        4: "single_static_metrics",
        5: "single_static_counts",
        6: "single_random_metrics",
        7: "single_random_counts",
        8: "two_static_assignment",
        9: "two_static_counts",
        10: "two_static_delta",
        11: "two_random_assignment",
        12: "two_random_counts",
        13: "single_static_anomaly",
        14: "single_static_anomaly_counts",
        15: "single_random_anomaly",
        16: "single_random_anomaly_counts",
        17: "single_anomaly_delta",
        18: "two_static_anomaly",
        19: "two_static_anomaly_counts",
        20: "two_random_anomaly",
        21: "two_random_anomaly_counts",
        22: "two_anomaly_delta",
    }

    for table_idx, payload_key in table_mapping.items():
        replace_table(doc.tables[table_idx], payload[payload_key])

    doc.save(REPORT)
    print(f"backup={backup}")
    print(f"updated={REPORT}")


if __name__ == "__main__":
    main()
