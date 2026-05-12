from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


ROOT = Path(r"F:\pythonprojects")
DOCX_PATH = ROOT / "my_report.docx"
REPORT_PATH = (
    ROOT
    / "baseline_GP"
    / "results"
    / "_temp_run_scripts"
    / "front_figure_directory_repair_report.json"
)


def clear_and_set_text(paragraph: Paragraph, text: str, style_name: str | None = None) -> None:
    """Clear paragraph XML content, preserving paragraph properties, then set plain text."""
    if style_name:
        paragraph.style = style_name
    paragraph._p.clear_content()
    paragraph.add_run(text)


def main() -> None:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = ROOT / f"my_report_backup_before_front_figure_directory_repair_{timestamp}.docx"
    shutil.copy2(DOCX_PATH, backup_path)

    doc = Document(DOCX_PATH)

    desired_entries = [
        "图2-1 已知静态地图下多源搜索场可视化\t12",
        "图3-1 单艇信息驱动安全路径规划总体流程图\t15",
        "图3-2 Anchor 与 Viewpoint 的空间语义区分示意图\t19",
        "图3-3 A*算法在栅格地图中的搜索过程示意图\t21",
        "图3-4 候选路径段评分与最优路径段选择流程图\t24",
        "图4-1 静态软责任区先验可视化\t28",
        "图4-2 residual_map 顺序分配过程\t30",
        "图5-1 实验采用的三类已知静态障碍地图\t33",
        "图5-2 单 USV 在不同地图与目标运动模式下的 UCB 搜索轨迹\t35",
    ]

    # Locate the manually maintained figure directory block.
    fig_title_idx = None
    table_title_idx = None
    for idx, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == "图目录":
            fig_title_idx = idx
        elif fig_title_idx is not None and paragraph.text.strip() == "表目录":
            table_title_idx = idx
            break

    if fig_title_idx is None or table_title_idx is None:
        raise RuntimeError("Could not locate front figure directory block.")

    style_name = None
    for idx in range(fig_title_idx + 1, table_title_idx):
        if doc.paragraphs[idx].style.name.lower().startswith("toc"):
            style_name = doc.paragraphs[idx].style.name
            break
    style_name = style_name or "toc 1"

    block_indices = list(range(fig_title_idx + 1, table_title_idx))
    if len(block_indices) < len(desired_entries):
        raise RuntimeError(
            f"Figure directory block has {len(block_indices)} paragraphs, "
            f"needs at least {len(desired_entries)}."
        )

    for idx, text in zip(block_indices, desired_entries):
        clear_and_set_text(doc.paragraphs[idx], text, style_name=style_name)

    for idx in block_indices[len(desired_entries) :]:
        clear_and_set_text(doc.paragraphs[idx], "", style_name="Normal")

    doc.save(DOCX_PATH)

    report = {
        "docx": str(DOCX_PATH),
        "backup": str(backup_path),
        "figure_directory_range": [fig_title_idx + 1, table_title_idx - 1],
        "entries": desired_entries,
    }
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
