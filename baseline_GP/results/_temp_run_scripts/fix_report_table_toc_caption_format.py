from __future__ import annotations

import json
import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(r"F:\pythonprojects")
DOCX_PATH = ROOT / "my_report.docx"
OUT_JSON = (
    ROOT
    / "baseline_GP"
    / "results"
    / "_temp_run_scripts"
    / "fix_report_table_toc_caption_format_report.json"
)


FIG_STYLE = "图目录项"
TABLE_STYLE = "表目录项"
SONGTI = "宋体"


def ensure_ppr(element):
    ppr = element.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        element.insert(0, ppr)
    return ppr


def ensure_ind(element):
    ppr = ensure_ppr(element)
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    return ind


def set_no_indent(element) -> None:
    ind = ensure_ind(element)
    for attr in (
        "left",
        "leftChars",
        "right",
        "rightChars",
        "firstLine",
        "firstLineChars",
        "hanging",
        "hangingChars",
    ):
        qattr = qn(f"w:{attr}")
        if qattr in ind.attrib:
            del ind.attrib[qattr]
    ind.set(qn("w:left"), "0")
    ind.set(qn("w:leftChars"), "0")
    ind.set(qn("w:firstLine"), "0")
    ind.set(qn("w:firstLineChars"), "0")


def clear_direct_indent(element) -> None:
    ppr = element.find(qn("w:pPr"))
    if ppr is None:
        return
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        return
    for attr in (
        "left",
        "leftChars",
        "right",
        "rightChars",
        "firstLine",
        "firstLineChars",
        "hanging",
        "hangingChars",
    ):
        qattr = qn(f"w:{attr}")
        if qattr in ind.attrib:
            del ind.attrib[qattr]


def set_toc_style_indent(style, left_twips: int | None = None, left_chars: int | None = None) -> None:
    ind = ensure_ind(style.element)
    for attr in (
        "left",
        "leftChars",
        "right",
        "rightChars",
        "firstLine",
        "firstLineChars",
        "hanging",
        "hangingChars",
    ):
        qattr = qn(f"w:{attr}")
        if qattr in ind.attrib:
            del ind.attrib[qattr]
    if left_twips is not None:
        ind.set(qn("w:left"), str(left_twips))
    if left_chars is not None:
        ind.set(qn("w:leftChars"), str(left_chars))
    ind.set(qn("w:firstLine"), "0")
    ind.set(qn("w:firstLineChars"), "0")


def set_run_font(run, font_name: str = SONGTI, size_pt: float = 10.5) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)


def set_style_font(style, font_name: str = SONGTI, size_pt: float = 10.5) -> None:
    style.font.name = font_name
    style.font.size = Pt(size_pt)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)


def configure_caption_style(doc: Document, style_name: str, *, after_pt: float) -> None:
    style = doc.styles[style_name]
    set_style_font(style, SONGTI, 10.5)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style.paragraph_format.line_spacing = 1
    style.paragraph_format.space_before = Pt(6)
    style.paragraph_format.space_after = Pt(after_pt)
    set_no_indent(style.element)


def format_caption_paragraph(paragraph: Paragraph, style_name: str) -> None:
    paragraph.style = style_name
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(12 if style_name == FIG_STYLE else 6)
    set_no_indent(paragraph._p)
    for run in paragraph.runs:
        set_run_font(run, SONGTI, 10.5)


TABLE_CAPTION_RE = re.compile(r"^表\d+-\d+(?:续)?\s+.+")
FIG_CAPTION_RE = re.compile(r"^图\d+-\d+\s+.+")


def paragraph_text_from_element(element, doc: Document) -> str:
    return Paragraph(element, doc).text.strip()


def is_empty_paragraph_element(element, doc: Document) -> bool:
    return element.tag == qn("w:p") and not paragraph_text_from_element(element, doc)


def is_table_caption_element(element, doc: Document) -> bool:
    if element.tag != qn("w:p"):
        return False
    paragraph = Paragraph(element, doc)
    text = paragraph.text.strip()
    if not TABLE_CAPTION_RE.match(text):
        return False
    if paragraph.style.name.lower().startswith("toc"):
        return False
    # Exclude narrative paragraphs such as "表5-7协同行为指标也显示".
    return True


def previous_nonempty_paragraph_or_table(children, idx: int, doc: Document):
    for j in range(idx - 1, -1, -1):
        child = children[j]
        if is_empty_paragraph_element(child, doc):
            continue
        return j, child
    return None, None


def next_nonempty_paragraph_or_table(children, idx: int, doc: Document, max_lookahead: int = 6):
    seen = 0
    for j in range(idx + 1, len(children)):
        child = children[j]
        if is_empty_paragraph_element(child, doc):
            continue
        seen += 1
        if seen > max_lookahead:
            break
        return j, child
    return None, None


def move_table_captions_above(doc: Document) -> list[dict[str, object]]:
    body = doc.element.body
    moved: list[dict[str, object]] = []
    changed = True
    while changed:
        changed = False
        children = list(body.iterchildren())
        for idx, child in enumerate(children):
            if child.tag != qn("w:tbl"):
                continue

            prev_idx, prev_child = previous_nonempty_paragraph_or_table(children, idx, doc)
            if prev_child is not None and is_table_caption_element(prev_child, doc):
                format_caption_paragraph(Paragraph(prev_child, doc), TABLE_STYLE)
                continue

            next_idx, next_child = next_nonempty_paragraph_or_table(children, idx, doc)
            if next_child is not None and is_table_caption_element(next_child, doc):
                caption_text = paragraph_text_from_element(next_child, doc)
                body.remove(next_child)
                child.addprevious(next_child)
                format_caption_paragraph(Paragraph(next_child, doc), TABLE_STYLE)
                moved.append(
                    {
                        "caption": caption_text,
                        "from_block_index": int(next_idx),
                        "to_before_table_index": int(idx),
                    }
                )
                changed = True
                break
    return moved


def clean_toc_indents(doc: Document) -> int:
    touched = 0
    toc_specs = {
        "toc 1": (0, 0),
        "TOC 1": (0, 0),
        "toc 2": (420, 200),
        "TOC 2": (420, 200),
        "toc 3": (840, 400),
        "TOC 3": (840, 400),
    }
    for name, (left_twips, left_chars) in toc_specs.items():
        try:
            set_toc_style_indent(doc.styles[name], left_twips, left_chars)
        except Exception:
            pass

    for paragraph in doc.paragraphs:
        if paragraph.style.name.lower().startswith("toc"):
            clear_direct_indent(paragraph._p)
            touched += 1
    return touched


def format_existing_captions(doc: Document) -> dict[str, int]:
    fig_count = 0
    table_count = 0
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if paragraph.style.name.lower().startswith("toc"):
            continue
        if paragraph.style.name == FIG_STYLE or (
            FIG_CAPTION_RE.match(text)
            and "给出" not in text
            and "展示" not in text
            and "为本文" not in text
            and len(text) <= 80
        ):
            format_caption_paragraph(paragraph, FIG_STYLE)
            fig_count += 1
        elif paragraph.style.name == TABLE_STYLE or TABLE_CAPTION_RE.match(text):
            format_caption_paragraph(paragraph, TABLE_STYLE)
            table_count += 1
    return {"figure_caption_count": fig_count, "table_caption_count": table_count}


def audit_table_positions(doc: Document) -> list[dict[str, object]]:
    blocks: list[tuple[str, Paragraph | Table]] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            blocks.append(("p", Paragraph(child, doc)))
        elif child.tag == qn("w:tbl"):
            blocks.append(("tbl", Table(child, doc)))
    missing = []
    for idx, (kind, obj) in enumerate(blocks):
        if kind != "tbl":
            continue
        prev_text = ""
        prev_style = ""
        for j in range(idx - 1, -1, -1):
            if blocks[j][0] == "p":
                p = blocks[j][1]
                assert isinstance(p, Paragraph)
                if p.text.strip():
                    prev_text = p.text.strip()
                    prev_style = p.style.name
                    break
            else:
                break
        if not TABLE_CAPTION_RE.match(prev_text):
            table = obj
            assert isinstance(table, Table)
            missing.append(
                {
                    "block_index": idx,
                    "rows": len(table.rows),
                    "cols": len(table.columns),
                    "previous_text": prev_text[:120],
                    "previous_style": prev_style,
                }
            )
    return missing


def audit_toc_direct_indents(doc: Document) -> list[dict[str, object]]:
    bad = []
    for idx, paragraph in enumerate(doc.paragraphs):
        if not paragraph.style.name.lower().startswith("toc"):
            continue
        ppr = paragraph._p.find(qn("w:pPr"))
        ind = None if ppr is None else ppr.find(qn("w:ind"))
        attrs = {}
        if ind is not None:
            for attr in ("left", "leftChars", "firstLine", "firstLineChars", "hanging", "hangingChars"):
                value = ind.get(qn(f"w:{attr}"))
                if value is not None:
                    attrs[attr] = value
        if attrs:
            bad.append(
                {
                    "paragraph_index": idx,
                    "style": paragraph.style.name,
                    "text": paragraph.text.strip()[:120],
                    "direct_indent_attrs": attrs,
                }
            )
    return bad


def audit_caption_styles(doc: Document) -> list[dict[str, object]]:
    bad = []
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if not text or paragraph.style.name.lower().startswith("toc"):
            continue
        if paragraph.style.name in {FIG_STYLE, TABLE_STYLE}:
            ppr = paragraph._p.find(qn("w:pPr"))
            ind = None if ppr is None else ppr.find(qn("w:ind"))
            attrs = {}
            if ind is not None:
                for attr in ("left", "leftChars", "firstLine", "firstLineChars"):
                    value = ind.get(qn(f"w:{attr}"))
                    if value is not None:
                        attrs[attr] = value
            if attrs not in ({}, {"left": "0", "leftChars": "0", "firstLine": "0", "firstLineChars": "0"}):
                bad.append(
                    {
                        "paragraph_index": idx,
                        "style": paragraph.style.name,
                        "text": text[:120],
                        "indent_attrs": attrs,
                    }
                )
    return bad


def main() -> None:
    if not DOCX_PATH.exists():
        raise FileNotFoundError(DOCX_PATH)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = ROOT / f"my_report_backup_before_table_toc_caption_format_{timestamp}.docx"
    shutil.copy2(DOCX_PATH, backup_path)

    doc = Document(DOCX_PATH)
    configure_caption_style(doc, FIG_STYLE, after_pt=12)
    configure_caption_style(doc, TABLE_STYLE, after_pt=6)

    moved = move_table_captions_above(doc)
    toc_touched = clean_toc_indents(doc)
    caption_counts = format_existing_captions(doc)

    doc.save(DOCX_PATH)

    check = Document(DOCX_PATH)
    report = {
        "docx": str(DOCX_PATH),
        "backup": str(backup_path),
        "moved_table_captions": moved,
        "moved_table_caption_count": len(moved),
        "toc_paragraphs_touched": toc_touched,
        **caption_counts,
        "table_position_missing_after": audit_table_positions(check),
        "toc_direct_indent_bad_after": audit_toc_direct_indents(check)[:20],
        "caption_style_bad_after": audit_caption_styles(check)[:20],
    }
    OUT_JSON.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "docx": str(DOCX_PATH),
                "backup": str(backup_path),
                "moved_table_caption_count": len(moved),
                "toc_paragraphs_touched": toc_touched,
                "figure_caption_count": caption_counts["figure_caption_count"],
                "table_caption_count": caption_counts["table_caption_count"],
                "table_position_missing_after_count": len(report["table_position_missing_after"]),
                "toc_direct_indent_bad_after_count": len(report["toc_direct_indent_bad_after"]),
                "caption_style_bad_after_count": len(report["caption_style_bad_after"]),
                "report": str(OUT_JSON),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
