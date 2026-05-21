"""Verify hyperlinks in my_report2.docx 图目录 and 表目录."""
import sys
sys.stdout.reconfigure(encoding="utf-8")
import re
from docx import Document
from docx.oxml.ns import qn

DOC = r"F:\pythonprojects\my_report2.docx"
doc = Document(DOC)

print("=== 图目录 hyperlinks ===\n")
in_fig = in_tbl = False
fig_ok = fig_bad = tbl_ok = tbl_bad = 0

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text == "图目录":
        in_fig = True
        continue
    if text == "表目录":
        in_fig = False
        in_tbl = True
        continue
    if (in_fig or in_tbl) and text and re.match(r'第[一二三四五六七八九十\d]+章', text):
        break
    if not text or not (in_fig or in_tbl):
        continue
    if not re.match(r'(图|表)\d+-', text):
        continue

    hls = para._element.findall(qn('w:hyperlink'))
    anchor = hls[0].get(qn('w:anchor')) if hls else "NONE"
    label = "图" if in_fig else "表"
    status = "OK" if hls else "NO LINK"

    if hls:
        if in_fig: fig_ok += 1
        else: tbl_ok += 1
    else:
        if in_fig: fig_bad += 1
        else: tbl_bad += 1

    print(f"  {label} {text[:80]:80s} -> #{anchor}")

print(f"\n图目录: {fig_ok} linked, {fig_bad} unlinked")
print(f"表目录: {tbl_ok} linked, {tbl_bad} unlinked")
print(f"Total: {fig_ok + tbl_ok} OK")

# Also verify the file opens correctly
print(f"\nDocument paragraphs: {len(doc.paragraphs)}")
print("Verification complete.")
