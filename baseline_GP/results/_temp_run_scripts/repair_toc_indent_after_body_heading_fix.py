from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"F:\pythonprojects")
DOCX_PATH = ROOT / "my_report.docx"
OUT_JSON = (
    ROOT
    / "baseline_GP"
    / "results"
    / "_temp_run_scripts"
    / "repair_toc_indent_after_body_heading_fix_report.json"
)


def ensure_ppr(element):
    ppr = element.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        element.insert(0, ppr)
    return ppr


def ensure_ind(element):
    ppr = ensure_ppr(element)
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    return ind


def remove_first_line_attrs(ind):
    for attr in ("firstLine", "firstLineChars", "hanging", "hangingChars"):
        qattr = qn(f"w:{attr}")
        if qattr in ind.attrib:
            del ind.attrib[qattr]


def set_toc_style_indent(style, left_twips: int | None = None, left_chars: int | None = None) -> None:
    ind = ensure_ind(style.element)
    remove_first_line_attrs(ind)
    ind.set(qn("w:firstLine"), "0")
    if left_twips is not None:
        ind.set(qn("w:left"), str(left_twips))
    if left_chars is not None:
        ind.set(qn("w:leftChars"), str(left_chars))


def clear_toc_paragraph_first_line(paragraph) -> dict[str, str]:
    ind = ensure_ind(paragraph._p)
    before = dict(ind.attrib)
    remove_first_line_attrs(ind)
    ind.set(qn("w:firstLine"), "0")
    return {k.split("}")[-1]: v for k, v in before.items()}


def ind_xml(element) -> dict[str, str]:
    ppr = element.find(qn("w:pPr"))
    if ppr is None:
        return {}
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        return {}
    out = {}
    for attr in ("firstLine", "firstLineChars", "left", "leftChars", "hanging", "hangingChars"):
        value = ind.get(qn(f"w:{attr}"))
        if value is not None:
            out[attr] = value
    return out


def main() -> None:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = ROOT / f"my_report_backup_before_toc_indent_repair_{timestamp}.docx"
    shutil.copy2(DOCX_PATH, backup_path)

    doc = Document(DOCX_PATH)

    # Keep TOC hierarchy left indents, but remove all first-line indents.
    toc_style_specs = {
        "toc 1": (None, None),
        "TOC 1": (None, None),
        "toc 2": (420, 200),
        "TOC 2": (420, 200),
        "toc 3": (840, 400),
        "TOC 3": (840, 400),
    }
    touched_styles = []
    for name, (left_twips, left_chars) in toc_style_specs.items():
        try:
            style = doc.styles[name]
        except Exception:
            continue
        set_toc_style_indent(style, left_twips, left_chars)
        touched_styles.append(name)

    touched_paras = []
    for idx, paragraph in enumerate(doc.paragraphs):
        if paragraph.style.name.lower().startswith("toc"):
            before = clear_toc_paragraph_first_line(paragraph)
            touched_paras.append(
                {
                    "index": idx,
                    "style": paragraph.style.name,
                    "text": paragraph.text.strip()[:120],
                    "before": before,
                    "after": ind_xml(paragraph._p),
                }
            )

    doc.save(DOCX_PATH)

    reopened = Document(DOCX_PATH)
    samples = []
    for idx, paragraph in enumerate(reopened.paragraphs[:180]):
        if paragraph.style.name.lower().startswith("toc"):
            samples.append(
                {
                    "index": idx,
                    "style": paragraph.style.name,
                    "text": paragraph.text.strip()[:100],
                    "direct": ind_xml(paragraph._p),
                    "style_indent": ind_xml(paragraph.style.element),
                }
            )

    result = {
        "docx": str(DOCX_PATH),
        "backup": str(backup_path),
        "touched_styles": touched_styles,
        "touched_toc_paragraph_count": len(touched_paras),
        "samples": samples[:80],
    }
    OUT_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in result.items() if k != "samples"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
