"""
Verify formatting fixes in my_report.docx
"""
from docx import Document
from docx.oxml.ns import qn
import re

doc = Document(r'F:\pythonprojects\my_report.docx')

def emu_to_pt(emu):
    if emu is None: return None
    return round(emu / 12700, 1)
def emu_to_cm(emu):
    if emu is None: return None
    return round(emu / 360000, 2)
def is_toc(p):
    sname = p.style.name if p.style else ''
    if 'toc' in sname.lower(): return True
    if re.search(r'\t\d+$', p.text.strip()): return True
    return False

ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

all_checks = []

# 1. Page margins
print('1. PAGE MARGINS (req: t2.5 b2.5 l2.5 r2.0 footer1.5)')
sec_ok = True
for i, sec in enumerate(doc.sections):
    t, b, l, r = emu_to_cm(sec.top_margin), emu_to_cm(sec.bottom_margin), emu_to_cm(sec.left_margin), emu_to_cm(sec.right_margin)
    f = emu_to_cm(sec.footer_distance)
    ok = t==2.5 and b==2.5 and l==2.5 and r==2.0 and f==1.5
    print(f'  Sec{i+1}: {t}/{b}/{l}/{r} footer{f} -> OK' if ok else f'  Sec{i+1}: {t}/{b}/{l}/{r} footer{f} -> FAIL')
    if not ok: sec_ok = False
all_checks.append(('Page margins', sec_ok))

# 2. Normal style
ns_style = doc.styles['Normal']
ns_ea_ok = False
rPr = ns_style._element.find(f'{{{ns}}}rPr')
if rPr is not None:
    rf = rPr.find(f'{{{ns}}}rFonts')
    if rf is not None:
        ns_ea_ok = rf.get(f'{{{ns}}}eastAsia') == '宋体'  # 宋体
ns_ok = ns_style.font.name == 'Times New Roman' and emu_to_pt(ns_style.font.size) == 12.0 and str(ns_style.paragraph_format.line_spacing) == '1.25' and ns_ea_ok
print(f'\n2. NORMAL STYLE: font={ns_style.font.name} size={emu_to_pt(ns_style.font.size)}pt ls={ns_style.paragraph_format.line_spacing} ea_ok={ns_ea_ok} -> {"OK" if ns_ok else "FAIL"}')
all_checks.append(('Normal style', ns_ok))

# 3. Chapter headings (一级标题)
print('\n3. CHAPTER HEADINGS:')
ch_ok = True
for i, p in enumerate(doc.paragraphs):
    if is_toc(p): continue
    if (p.style.name if p.style else '') == '一级标题':
        r0 = p.runs[0] if p.runs else None
        if r0:
            rPr = r0._r.find(f'{{{ns}}}rPr')
            ea = None
            if rPr is not None:
                rf = rPr.find(f'{{{ns}}}rFonts')
                if rf is not None:
                    ea = rf.get(f'{{{ns}}}eastAsia')
            sz = emu_to_pt(r0.font.size)
            ok = sz == 16.0 and r0.font.bold == True and ea == '黑体'
            if not ok: ch_ok = False
            mark = 'OK' if ok else 'FAIL'
            print(f'  P{i}: {p.text.strip()[:50]} sz={sz}pt bold={r0.font.bold} ea={ea} -> {mark}')
all_checks.append(('Chapter headings', ch_ok))

# 4. Section headings (1.1)
print('\n4. SECTION HEADINGS (1.1):')
sec1_ok = True
found = 0
for i, p in enumerate(doc.paragraphs):
    if is_toc(p): continue
    if re.match(r'^\d+\.\d+ ', p.text.strip()):
        found += 1
        if found <= 3:
            r0 = p.runs[0] if p.runs else None
            if r0:
                rPr = r0._r.find(f'{{{ns}}}rPr')
                ea = None
                if rPr is not None:
                    rf = rPr.find(f'{{{ns}}}rFonts')
                    if rf is not None:
                        ea = rf.get(f'{{{ns}}}eastAsia')
                sz = emu_to_pt(r0.font.size)
                ok = sz == 14.0 and r0.font.bold == True and ea == '黑体'
                if not ok: sec1_ok = False
                mark = 'OK' if ok else 'FAIL'
                print(f'  P{i}: {p.text.strip()[:60]} sz={sz}pt bold={r0.font.bold} ea={ea} -> {mark}')
print(f'  (sampled 3)')
all_checks.append(('Section headings (1.1)', sec1_ok))

# 5. Subsection headings (1.1.1)
print('\n5. SUBSECTION HEADINGS (1.1.1):')
sec2_ok = True
found = 0
for i, p in enumerate(doc.paragraphs):
    if is_toc(p): continue
    if re.match(r'^\d+\.\d+\.\d+', p.text.strip()):
        found += 1
        if found <= 3:
            r0 = p.runs[0] if p.runs else None
            if r0:
                rPr = r0._r.find(f'{{{ns}}}rPr')
                ea = None
                if rPr is not None:
                    rf = rPr.find(f'{{{ns}}}rFonts')
                    if rf is not None:
                        ea = rf.get(f'{{{ns}}}eastAsia')
                sz = emu_to_pt(r0.font.size)
                ok = sz == 12.0 and r0.font.bold == True and ea == '黑体'
                if not ok: sec2_ok = False
                mark = 'OK' if ok else 'FAIL'
                print(f'  P{i}: {p.text.strip()[:60]} sz={sz}pt bold={r0.font.bold} ea={ea} -> {mark}')
print(f'  (sampled 3)')
all_checks.append(('Subsection headings (1.1.1)', sec2_ok))

# 6. Captions
print('\n6. CAPTIONS:')
cap_ok = True
for label, pat in [('图', r'^图\s*\d+'), ('表', r'^表\s*\d+')]:
    found = False
    for p in doc.paragraphs:
        if is_toc(p): continue
        if re.match(pat, p.text.strip()):
            r0 = p.runs[0] if p.runs else None
            sz = emu_to_pt(r0.font.size) if r0 else None
            al = str(p.paragraph_format.alignment)
            sb = emu_to_pt(p.paragraph_format.space_before)
            sa = emu_to_pt(p.paragraph_format.space_after)
            exp_sa = 12.0 if label == '图' else 6.0
            ok = sz == 10.5 and al == 'CENTER (1)' and sb == 6.0 and sa == exp_sa
            if not ok: cap_ok = False
            mark = 'OK' if ok else 'FAIL'
            print(f'  {label}: {p.text.strip()[:60]} sz={sz}pt al={al} sb={sb}pt sa={sa}pt -> {mark}')
            found = True
            break
all_checks.append(('Captions', cap_ok))

# 7. References
print('\n7. REFERENCES:')
ref_ok = True
for p in doc.paragraphs:
    if p.text.strip() == '参考文献':
        r0 = p.runs[0] if p.runs else None
        if r0:
            rPr = r0._r.find(f'{{{ns}}}rPr')
            ea = None
            if rPr is not None:
                rf = rPr.find(f'{{{ns}}}rFonts')
                if rf is not None:
                    ea = rf.get(f'{{{ns}}}eastAsia')
            sz = emu_to_pt(r0.font.size)
            al = str(p.paragraph_format.alignment)
            ok = sz == 16.0 and r0.font.bold == True and ea == '黑体' and al == 'CENTER (1)'
            if not ok: ref_ok = False
            mark = 'OK' if ok else 'FAIL'
            print(f'  Title: sz={sz}pt bold={r0.font.bold} ea={ea} al={al} -> {mark}')
        break

found = 0
for p in doc.paragraphs:
    if re.match(r'^\[\d+\]', p.text.strip()):
        found += 1
        if found <= 2:
            r0 = p.runs[0] if p.runs else None
            sz = emu_to_pt(r0.font.size) if r0 else None
            ls = p.paragraph_format.line_spacing
            ls_pt = emu_to_pt(ls) if ls else None
            ok = sz == 10.5 and ls_pt == 16.0
            if not ok: ref_ok = False
            mark = 'OK' if ok else 'FAIL'
            print(f'  Entry{found}: sz={sz}pt ls={ls_pt}pt -> {mark}')
        if found >= 2: break
all_checks.append(('References', ref_ok))

# Summary
print()
print('=' * 40)
print('SUMMARY')
print('=' * 40)
for name, ok in all_checks:
    print(f'  {name}: {"PASS" if ok else "FAIL"}')
all_pass = all(ok for _, ok in all_checks)
print(f'\n  OVERALL: {"ALL PASS" if all_pass else "SOME FAIL"}')
