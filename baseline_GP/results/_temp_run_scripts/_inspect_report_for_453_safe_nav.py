from pathlib import Path
from docx import Document
p=Path(r'F:\pythonprojects\my_report.docx')
doc=Document(p)
print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables), 'inline_shapes', len(doc.inline_shapes))
for i,p in enumerate(doc.paragraphs):
    t=p.text.strip()
    if t.startswith('4.5') or t.startswith('4.6') or (485 <= i <= 525):
        has_pic = any(r._element.xpath('.//pic:pic') for r in p.runs)
        print(f'{i}: style={p.style.name!r} pic={has_pic} text={t}')
print('\nFIG_CAPTIONS')
for i,p in enumerate(doc.paragraphs):
    t=p.text.strip()
    if t.startswith('图'):
        print(f'{i}: style={p.style.name!r} text={t}')
print('\nCITATION_OCCURRENCES')
for i,p in enumerate(doc.paragraphs):
    t=p.text
    if '[35' in t or '[36' in t or '[37' in t or '[34' in t:
        print(f'{i}: {t.strip()}')
