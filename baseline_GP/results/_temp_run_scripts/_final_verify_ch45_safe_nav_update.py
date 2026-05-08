from pathlib import Path
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.text.paragraph import Paragraph
import re
p=Path(r'F:\pythonprojects\my_report.docx')
doc=Document(p)
print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables), 'inline_shapes', len(doc.inline_shapes))
print('\nCH45_HEADINGS_AND_CAPTIONS')
for i,para in enumerate(doc.paragraphs):
    t=para.text.strip()
    if 489 <= i <= 538:
        print(f'{i}: {para.style.name!r} pic={bool(para._p.xpath(".//pic:pic"))} math={bool(para._p.xpath(".//m:oMath"))} text={t}')
print('\nFIG_CAPTIONS')
for i,pgh in enumerate(doc.paragraphs):
    t=pgh.text.strip()
    if pgh.style.name=='图目录项' and t.startswith('图'):
        print(i,t)
print('\nCITATIONS')
for i,pgh in enumerate(doc.paragraphs):
    t=pgh.text.strip()
    if '[7,34' in t or '[7,35' in t or '[7,35-36]' in t or '[32-34]' in t or '[7,34-35]' in t:
        print(i,t)
print('\nTABLE_CAPTION_ORDER')
blocks=[]
for child in doc.element.body.iterchildren():
    if isinstance(child, CT_P):
        para=Paragraph(child, doc)
        blocks.append(('P', para.text.strip(), para.style.name, bool(child.xpath('.//pic:pic'))))
    elif isinstance(child, CT_Tbl):
        blocks.append(('T','<TABLE>','',False))
problems=[]
for idx,(kind,text,style,pic) in enumerate(blocks):
    if style=='表目录项' and text.startswith('表'):
        prev=blocks[idx-1] if idx>0 else None
        if prev is None or prev[0] != 'T':
            problems.append((idx,text,prev[:3] if prev else None))
        print(idx,text,'prev=',prev[:3] if prev else None,'next=',blocks[idx+1][:3] if idx+1<len(blocks) else None)
print('table_caption_position_problems', problems)
print('\nBAD_MARKERS')
joined='\n'.join(p.text for p in doc.paragraphs)
for m in ['锛','鈥','????','�']:
    print(m, m in joined)
