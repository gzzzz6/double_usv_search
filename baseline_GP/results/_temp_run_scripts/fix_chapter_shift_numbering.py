from __future__ import annotations

import json
import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(r"F:\pythonprojects")
DOCX_PATH = ROOT / "my_report.docx"
REPORT_PATH = (
    ROOT
    / "baseline_GP"
    / "results"
    / "_temp_run_scripts"
    / "chapter_shift_numbering_fix_report.json"
)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    """Replace paragraph text while keeping the paragraph style."""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def insert_paragraph_after(paragraph: Paragraph, text: str, style_name: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style_name:
        new_para.style = style_name
    new_para.add_run(text)
    return new_para


def replace_with_placeholders(text: str, mapping: list[tuple[str, str]]) -> tuple[str, int]:
    count = 0
    for old, placeholder in mapping:
        n = text.count(old)
        if n:
            text = text.replace(old, placeholder)
            count += n
    return text, count


def restore_placeholders(text: str, mapping: dict[str, str]) -> tuple[str, int]:
    count = 0
    for placeholder, new in mapping.items():
        n = text.count(placeholder)
        if n:
            text = text.replace(placeholder, new)
            count += n
    return text, count


def main() -> None:
    if not DOCX_PATH.exists():
        raise FileNotFoundError(DOCX_PATH)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = ROOT / f"my_report_backup_before_chapter_shift_numbering_fix_{timestamp}.docx"
    shutil.copy2(DOCX_PATH, backup_path)

    doc = Document(DOCX_PATH)

    # The document was shifted from chapters 3/4/5/6 to 2/3/4/5.
    # Use placeholders to avoid cascading replacements such as 图5-2 -> 图4-1 -> 图3-1.
    figure_to_placeholder = [
        ("图5-5", "__FIG_CH5_TRAJ__"),
        ("图5-4", "__FIG_CH5_MAPS__"),
        ("图5-3", "__FIG_CH4_RESIDUAL__"),
        ("图5-2", "__FIG_CH4_RESP__"),
        ("图4-5", "__FIG_CH3_SCORE__"),
        ("图4-4", "__FIG_CH3_ASTAR__"),
        ("图4-3", "__FIG_CH3_ANCHOR_VIEWPOINT__"),
        ("图4-1", "__FIG_CH3_FLOW__"),
        ("图3-1", "__FIG_CH2_FIELDS__"),
    ]
    placeholder_to_figure = {
        "__FIG_CH2_FIELDS__": "图2-1",
        "__FIG_CH3_FLOW__": "图3-1",
        "__FIG_CH3_ANCHOR_VIEWPOINT__": "图3-2",
        "__FIG_CH3_ASTAR__": "图3-3",
        "__FIG_CH3_SCORE__": "图3-4",
        "__FIG_CH4_RESP__": "图4-1",
        "__FIG_CH4_RESIDUAL__": "图4-2",
        "__FIG_CH5_MAPS__": "图5-1",
        "__FIG_CH5_TRAJ__": "图5-2",
    }

    text_replacements = {
        "第3章构建的综合信息价值场": "第2章构建的综合信息价值场",
        "第3章中构建的search_info_map": "第2章中构建的search_info_map",
        "第 6.2 节": "第 5.2 节",
        "第 6.3 节": "第 5.3 节",
        "第 6.4 节": "第 5.4 节",
        "第 5 章提出的双 USV 协同搜索决策机制": "第 4 章提出的双 USV 协同搜索决策机制",
    }

    changed_paragraphs: list[dict[str, str | int]] = []
    figure_replacement_count = 0
    text_replacement_count = 0

    for idx, paragraph in enumerate(doc.paragraphs):
        old_text = paragraph.text
        if not old_text:
            continue

        new_text, n1 = replace_with_placeholders(old_text, figure_to_placeholder)
        new_text, n2 = restore_placeholders(new_text, placeholder_to_figure)
        figure_replacement_count += n1 + n2

        for old, new in text_replacements.items():
            n = new_text.count(old)
            if n:
                new_text = new_text.replace(old, new)
                text_replacement_count += n

        if new_text != old_text:
            set_paragraph_text(paragraph, new_text)
            changed_paragraphs.append(
                {
                    "paragraph_index": idx,
                    "style": paragraph.style.name,
                    "before": old_text,
                    "after": new_text,
                }
            )

    # The manually maintained front figure directory did not yet include the two
    # experiment figures added in Chapter 5. Add them after the Chapter 4 figure entries.
    # Page numbers follow the current visible document pagination and can be refreshed in Word later.
    fig_dir_texts = [p.text.strip() for p in doc.paragraphs]
    inserted_front_entries: list[str] = []
    if not any(t.startswith("图5-1 实验采用的三类已知静态障碍地图") for t in fig_dir_texts):
        anchor = None
        for paragraph in doc.paragraphs:
            if paragraph.text.strip().startswith("图4-2 residual_map 顺序分配过程"):
                anchor = paragraph
                break
        if anchor is not None:
            first = insert_paragraph_after(
                anchor,
                "图5-1 实验采用的三类已知静态障碍地图\t33",
                style_name=anchor.style.name,
            )
            insert_paragraph_after(
                first,
                "图5-2 单 USV 在不同地图与目标运动模式下的 UCB 搜索轨迹\t35",
                style_name=anchor.style.name,
            )
            inserted_front_entries.extend(
                [
                    "图5-1 实验采用的三类已知静态障碍地图\t33",
                    "图5-2 单 USV 在不同地图与目标运动模式下的 UCB 搜索轨迹\t35",
                ]
            )

    doc.save(DOCX_PATH)

    report = {
        "docx": str(DOCX_PATH),
        "backup": str(backup_path),
        "changed_paragraph_count": len(changed_paragraphs),
        "figure_replacement_count_including_placeholders": figure_replacement_count,
        "text_replacement_count": text_replacement_count,
        "inserted_front_figure_directory_entries": inserted_front_entries,
        "changed_paragraphs": changed_paragraphs,
    }
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in report.items() if k != "changed_paragraphs"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
