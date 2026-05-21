"""Check XML of one TOC entry and one directory entry."""
import sys
sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

DOC = r"F:\pythonprojects\my_report2.docx"
doc = Document(DOC)

# Find a TOC entry with hyperlink (working)
print("=== TOC ENTRY XML (working hyperlink) ===\n")
for para in doc.paragraphs:
    text = para.text.strip()
    if "toc 1" in (para.style.name if para.style else ""):
        hl = para._element.findall(qn('w:hyperlink'))
        if hl:
            xml_str = etree.tostring(para._element, pretty_print=True, encoding="unicode")
            print(xml_str[:1500])
            break

print("\n=== 图目录 ENTRY XML (first entry) ===\n")
for para in doc.paragraphs:
    text = para.text.strip()
    if text.startswith("图2-1"):
        xml_str = etree.tostring(para._element, pretty_print=True, encoding="unicode")
        print(xml_str[:1500])
        break

print("\n=== 表目录 ENTRY XML (first entry) ===\n")
for para in doc.paragraphs:
    text = para.text.strip()
    if text.startswith("表5-1"):
        xml_str = etree.tostring(para._element, pretty_print=True, encoding="unicode")
        print(xml_str[:1500])
        break

print("\n=== BODY CAPTION 图目录项 XML ===\n")
for para in doc.paragraphs:
    style = para.style.name if para.style else ""
    if style == "图目录项":
        text = para.text.strip()
        if "图2-1" in text:
            xml_str = etree.tostring(para._element, pretty_print=True, encoding="unicode")
            print(xml_str[:1500])
            break
