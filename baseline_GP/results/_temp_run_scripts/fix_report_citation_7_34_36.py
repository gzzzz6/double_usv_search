from pathlib import Path
from datetime import datetime
import shutil
from docx import Document

root=Path(r'F:\pythonprojects')
docx=root/'my_report.docx'
backup=root/f'my_report_before_citation_7_34_36_fix_{datetime.now():%Y%m%d_%H%M%S}.docx'
shutil.copy2(docx, backup)
doc=Document(docx)
changed=0
for p in doc.paragraphs:
    if '[7,34,36]' in p.text:
        text=p.text.replace('[7,34,36]', '[7,35,37]')
        for r in p.runs:
            r.text=''
        if p.runs:
            p.runs[0].text=text
        else:
            p.add_run(text)
        changed += 1
doc.save(docx)
print('updated=', docx)
print('backup=', backup)
print('changed=', changed)
