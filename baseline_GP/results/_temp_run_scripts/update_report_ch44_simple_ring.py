from __future__ import annotations

import shutil
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt


ROOT = Path(r"F:\pythonprojects")
DOCX_PATH = ROOT / "my_report.docx"
BACKUP_PATH = ROOT / f"my_report_before_ch44_simple_ring_{datetime.now():%Y%m%d_%H%M%S}.docx"


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)
        paragraph._p = paragraph._element = None


def set_run_font(run, *, east_asia: str = "宋体", ascii_font: str = "Times New Roman", size_pt: float = 12.0) -> None:
    run.font.name = ascii_font
    run.font.size = Pt(size_pt)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", east_asia)
    r_fonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii", ascii_font)
    r_fonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi", ascii_font)


def format_paragraph(paragraph, *, size_pt: float = 12.0) -> None:
    for run in paragraph.runs:
        set_run_font(run, size_pt=size_pt)


def m_run(text: str):
    r = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    r.append(t)
    return r


def m_sub(base: str, sub: str):
    ssub = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    e.append(m_run(base))
    sub_el = OxmlElement("m:sub")
    sub_el.append(m_run(sub))
    ssub.append(e)
    ssub.append(sub_el)
    return ssub


def add_formula_before(anchor_paragraph, parts: list[object], number: str):
    paragraph = anchor_paragraph.insert_paragraph_before("", style="公式目录项")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath = OxmlElement("m:oMath")
    for part in parts:
        if isinstance(part, tuple) and len(part) == 2:
            omath.append(m_sub(part[0], part[1]))
        else:
            omath.append(m_run(str(part)))
    paragraph._p.append(omath)
    run = paragraph.add_run(f"    {number}")
    set_run_font(run, size_pt=10.5)
    return paragraph


def insert_paragraph_before(anchor_paragraph, text: str, style: str = "Normal"):
    paragraph = anchor_paragraph.insert_paragraph_before(text, style=style)
    format_paragraph(paragraph)
    return paragraph


def find_heading(doc: Document, prefix: str, start: int = 0, style_name: str | None = None) -> int:
    for idx, paragraph in enumerate(doc.paragraphs[start:], start=start):
        if paragraph.text.strip().startswith(prefix) and (
            style_name is None or paragraph.style.name == style_name
        ):
            return idx
    raise RuntimeError(f"Cannot find heading starting with {prefix!r}")


def main() -> None:
    if not DOCX_PATH.exists():
        raise FileNotFoundError(DOCX_PATH)
    shutil.copy2(DOCX_PATH, BACKUP_PATH)

    doc = Document(DOCX_PATH)
    # The document contains generated TOC display paragraphs near the front.
    # Restrict matching to the real chapter body by first finding the actual
    # Chapter 4 top-level heading and then the 4.4 body section.
    chapter4_idx = find_heading(doc, "4 信息型路径规划", style_name="一级标题")
    section44_idx = find_heading(doc, "4.4", start=chapter4_idx + 1, style_name="二级标题")
    start_idx = find_heading(doc, "4.4.2", start=section44_idx + 1, style_name="三级标题")
    middle_idx = find_heading(doc, "4.4.3", start=start_idx + 1, style_name="三级标题")
    end_idx = find_heading(doc, "4.5", start=middle_idx + 1, style_name="二级标题")

    old_paragraphs = list(doc.paragraphs[start_idx:end_idx])
    anchor = doc.paragraphs[start_idx]

    insert_paragraph_before(anchor, "4.4.2 简单环形候选视点生成", style="三级标题")
    insert_paragraph_before(
        anchor,
        "对某一热点簇的 anchor，本文当前主线采用 simple_ring_v1 生成候选 viewpoint。该方法不再使用多来源扰动采样候选池，而是直接围绕热点代表点构造一组距离适中、能够覆盖热点区域且位于自由空间中的环形候选视点。这样既保留了“从热点周围观察”的主动感知语义，也降低了候选生成过程的复杂度。",
    )
    insert_paragraph_before(
        anchor,
        "设热点簇代表点为 a，候选视点为 v，USV 传感器作用半径为 R_s。系统首先设置期望观测距离和最小观测距离，并要求候选视点仍处于能够覆盖 anchor 的传感器范围内：",
    )
    add_formula_before(
        anchor,
        [
            ("r", "pref"),
            " = 0.75 ",
            ("R", "s"),
            ",    ",
            ("r", "min"),
            " = 0.30 ",
            ("R", "s"),
            ",    ",
            ("r", "min"),
            " ≤ ||v - a|| ≤ ",
            ("R", "s"),
        ],
        "（45）",
    )
    insert_paragraph_before(
        anchor,
        "其中，r_pref 表示期望观测距离，r_min 用于避免视点过于贴近热点中心，R_s 则限定 viewpoint 必须位于传感器能够覆盖 anchor 的范围内。本文默认传感器半径为 5 个栅格，因此 r_pref 约为 3.75 个栅格，r_min 约为 1.5 个栅格。该设置使 USV 倾向于从热点周围合适距离进行观察，而不是直接驶入热点中心。",
    )
    insert_paragraph_before(
        anchor,
        "在实现中，系统以热点簇加权中心附近的局部自由栅格作为候选搜索范围，但候选距离约束仍以 anchor 为参照。对于满足距离条件的自由格点，系统首先按照其到 anchor 的距离与 r_pref 的偏差进行排序；偏差越小，说明该格点越接近期望观测环带。若多个候选点距离匹配程度相近，则优先选择邻近障碍更少的自由格点；若仍存在并列，则按照栅格坐标顺序稳定排序。",
    )
    insert_paragraph_before(
        anchor,
        "当前默认每个 anchor 最终保留 6 个 viewpoint 进入后续路径段生成。为了给预排序和可达性筛选留出冗余，simple_ring_v1 会先构造最多 4 倍于最终保留数量的候选池；在默认设置下，候选池上限为 24 个格点。若 anchor 周围不存在满足环形距离约束的自由格点，系统退化为在热点簇加权中心附近寻找最近自由格点作为兜底 viewpoint，从而避免局部障碍结构导致有效热点无法产生候选视点。",
    )

    insert_paragraph_before(anchor, "4.4.3 候选视点预排序", style="三级标题")
    insert_paragraph_before(
        anchor,
        "simple_ring_v1 生成的 viewpoint 只是局部几何候选，并不等同于最终执行目标。系统还需要根据候选视点的观测价值、几何匹配程度、当前位置距离和障碍邻近情况进行预排序，从中筛选出更值得进入路径生成阶段的候选点。该排序仍属于候选压缩步骤，最终执行路径段将在 4.6 节中根据路径段综合评分确定。",
    )
    insert_paragraph_before(
        anchor,
        "对于候选 viewpoint v，系统首先计算其传感器覆盖范围内综合信息价值场的平均值，记为 I_bar(v)。该项反映 USV 位于 v 时能够覆盖到的搜索信息价值。其次，系统计算 v 到 anchor 的距离匹配项 D(v)，当 ||v-a|| 越接近 r_pref 时，D(v) 越大。系统还引入当前位置到 v 的曼哈顿距离代理代价 C_d(v)，以及 v 邻近障碍数量对应的障碍惩罚 C_o(v)。候选 viewpoint 的预排序得分写为：",
    )
    add_formula_before(
        anchor,
        [
            ("S", "v"),
            " = 1.80 I_bar(v) + 0.60 D(v) - 0.05 ",
            ("C", "d"),
            "(v) - 0.10 ",
            ("C", "o"),
            "(v)",
        ],
        "（46）",
    )
    insert_paragraph_before(
        anchor,
        "式中，I_bar(v) 是主要收益项，用于鼓励候选视点覆盖综合信息价值较高的区域；D(v) 用于鼓励候选视点位于合适的观测距离；C_d(v) 用于抑制距离当前 USV 位置过远的候选点；C_o(v) 用于降低邻近障碍较多候选点的优先级。上述权重与当前代码中的 simple_ring_v1 预排序逻辑一致，其中搜索信息覆盖项权重最大，说明候选视点生成仍以任务信息价值为主要依据。",
    )
    insert_paragraph_before(
        anchor,
        "预排序完成后，系统会基于当前 USV 位置到各候选 viewpoint 的可达性进行筛选，并重构通向 viewpoint 的路径。只有可达候选点才会被保留进入后续 segment_path 生成阶段。需要强调的是，viewpoint 预排序不是最终路径评分；它只是减少候选数量并剔除低质量或不可达候选点。最终执行哪一段路径，仍由后续候选路径段的搜索信息收益、观测刷新收益、执行代价和掉头惩罚共同决定。",
    )

    for paragraph in old_paragraphs:
        delete_paragraph(paragraph)

    doc.save(DOCX_PATH)
    print(f"updated={DOCX_PATH}")
    print(f"backup={BACKUP_PATH}")


if __name__ == "__main__":
    main()
