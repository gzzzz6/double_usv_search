from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


ROOT = Path(r"F:\pythonprojects")
REPORT_DOCX = ROOT / "my_report.docx"
SPEC_DIR = ROOT / "readme"
OUT_JSON = ROOT / "baseline_GP" / "results" / "_temp_run_scripts" / "heading_indent_audit.json"
OUT_MD = ROOT / "baseline_GP" / "results" / "_temp_run_scripts" / "heading_indent_audit.md"


def twips_to_cm(value):
    if value is None:
        return None
    return round(value.twips / 567.0, 3)


def twips_to_pt(value):
    if value is None:
        return None
    return round(value.pt, 2)


def align_name(value):
    if value is None:
        return None
    mapping = {
        WD_ALIGN_PARAGRAPH.LEFT: "LEFT",
        WD_ALIGN_PARAGRAPH.CENTER: "CENTER",
        WD_ALIGN_PARAGRAPH.RIGHT: "RIGHT",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "JUSTIFY",
    }
    return mapping.get(value, str(value))


def style_info(doc, style_name):
    try:
        style = doc.styles[style_name]
    except Exception:
        return None
    pf = style.paragraph_format
    font = style.font
    return {
        "style_name": style.name,
        "base_style": style.base_style.name if style.base_style is not None else None,
        "left_indent_cm": twips_to_cm(pf.left_indent),
        "first_line_indent_cm": twips_to_cm(pf.first_line_indent),
        "space_before_pt": twips_to_pt(pf.space_before),
        "space_after_pt": twips_to_pt(pf.space_after),
        "line_spacing": pf.line_spacing,
        "alignment": align_name(pf.alignment),
        "font_name": font.name,
        "font_size_pt": twips_to_pt(font.size),
        "bold": font.bold,
    }


def para_info(idx, para):
    pf = para.paragraph_format
    style_pf = para.style.paragraph_format
    return {
        "index": idx,
        "text": para.text.strip(),
        "style": para.style.name,
        "direct_left_indent_cm": twips_to_cm(pf.left_indent),
        "direct_first_line_indent_cm": twips_to_cm(pf.first_line_indent),
        "direct_alignment": align_name(pf.alignment),
        "style_left_indent_cm": twips_to_cm(style_pf.left_indent),
        "style_first_line_indent_cm": twips_to_cm(style_pf.first_line_indent),
        "style_alignment": align_name(style_pf.alignment),
        "style_space_before_pt": twips_to_pt(style_pf.space_before),
        "style_space_after_pt": twips_to_pt(style_pf.space_after),
    }


def main():
    spec_files = sorted(SPEC_DIR.glob("*.docx"), key=lambda p: p.stat().st_mtime, reverse=True)
    spec_docx = next((p for p in spec_files if "规范" in p.name), spec_files[0] if spec_files else None)
    if spec_docx is None:
        raise FileNotFoundError(SPEC_DIR)

    spec = Document(spec_docx)
    report = Document(REPORT_DOCX)

    relevant_style_names = [
        "正文",
        "Normal",
        "一级标题",
        "二级标题",
        "三级标题",
        "标题 1",
        "标题 2",
        "标题 3",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "图目录项",
        "表目录项",
    ]

    spec_text = []
    for idx, para in enumerate(spec.paragraphs):
        text = para.text.strip()
        if text:
            spec_text.append({"index": idx, "style": para.style.name, "text": text})

    report_styles = {}
    for name in relevant_style_names:
        info = style_info(report, name)
        if info is not None:
            report_styles[name] = info

    spec_styles = {}
    for name in relevant_style_names:
        info = style_info(spec, name)
        if info is not None:
            spec_styles[name] = info

    # Capture headings by visible numbering pattern and by known heading styles.
    heading_pat = re.compile(r"^(?:[1-6](?:\.[1-9][0-9]*){0,2})\s+")
    report_heading_paras = []
    for idx, para in enumerate(report.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        if heading_pat.match(text) or para.style.name in {
            "一级标题",
            "二级标题",
            "三级标题",
            "标题 1",
            "标题 2",
            "标题 3",
            "Heading 1",
            "Heading 2",
            "Heading 3",
        }:
            report_heading_paras.append(para_info(idx, para))

    # Group a small sample per visible heading level.
    level_samples = {"level1": [], "level2": [], "level3": []}
    for item in report_heading_paras:
        text = item["text"]
        if re.match(r"^[1-6]\s+", text) and len(level_samples["level1"]) < 12:
            level_samples["level1"].append(item)
        elif re.match(r"^[1-6]\.[1-9][0-9]*\s+", text) and len(level_samples["level2"]) < 20:
            level_samples["level2"].append(item)
        elif re.match(r"^[1-6]\.[1-9][0-9]*\.[1-9][0-9]*\s+", text) and len(level_samples["level3"]) < 20:
            level_samples["level3"].append(item)

    suspicious = []
    for item in report_heading_paras:
        # Heading paragraphs should not carry a first-line indent like body text.
        if item["direct_first_line_indent_cm"] not in (None, 0):
            suspicious.append(
                {
                    **item,
                    "reason": "heading paragraph has direct first_line_indent",
                }
            )
        elif item["style_first_line_indent_cm"] not in (None, 0):
            suspicious.append(
                {
                    **item,
                    "reason": "heading style has first_line_indent",
                }
            )

    result = {
        "spec_docx": str(spec_docx),
        "report_docx": str(REPORT_DOCX),
        "spec_text": spec_text,
        "spec_styles": spec_styles,
        "report_styles": report_styles,
        "heading_count": len(report_heading_paras),
        "level_samples": level_samples,
        "suspicious_count": len(suspicious),
        "suspicious_heading_indent": suspicious[:80],
    }

    OUT_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    lines = []
    lines.append(f"# Heading indent audit")
    lines.append("")
    lines.append(f"- spec: `{spec_docx}`")
    lines.append(f"- report: `{REPORT_DOCX}`")
    lines.append(f"- heading_count: {len(report_heading_paras)}")
    lines.append(f"- suspicious_count: {len(suspicious)}")
    lines.append("")
    lines.append("## Report Styles")
    for name, info in report_styles.items():
        lines.append(f"- {name}: {info}")
    lines.append("")
    lines.append("## Spec Text")
    for item in spec_text:
        lines.append(f"- P{item['index']:03d} [{item['style']}] {item['text']}")
    lines.append("")
    lines.append("## Heading Samples")
    for level, items in level_samples.items():
        lines.append(f"### {level}")
        for item in items:
            lines.append(f"- P{item['index']:04d} [{item['style']}] {item['text']} | direct first={item['direct_first_line_indent_cm']} style first={item['style_first_line_indent_cm']} left={item['style_left_indent_cm']}")
    lines.append("")
    lines.append("## Suspicious")
    for item in suspicious[:80]:
        lines.append(f"- P{item['index']:04d} [{item['style']}] {item['text']} | {item['reason']} | direct first={item['direct_first_line_indent_cm']} style first={item['style_first_line_indent_cm']}")
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")

    print(json.dumps({
        "spec_docx": str(spec_docx),
        "report_docx": str(REPORT_DOCX),
        "heading_count": len(report_heading_paras),
        "suspicious_count": len(suspicious),
        "out_json": str(OUT_JSON),
        "out_md": str(OUT_MD),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
