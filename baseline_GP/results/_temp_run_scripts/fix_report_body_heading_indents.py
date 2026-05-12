from __future__ import annotations

import json
import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(r"F:\pythonprojects")
DOCX_PATH = ROOT / "my_report.docx"
OUT_JSON = (
    ROOT
    / "baseline_GP"
    / "results"
    / "_temp_run_scripts"
    / "fix_report_body_heading_indents_report.json"
)


def ensure_ppr(element):
    ppr = element.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        element.insert(0, ppr)
    return ppr


def clear_indent_attrs(ind):
    for attr in (
        "w:left",
        "w:leftChars",
        "w:right",
        "w:rightChars",
        "w:firstLine",
        "w:firstLineChars",
        "w:hanging",
        "w:hangingChars",
    ):
        qattr = qn(attr)
        if qattr in ind.attrib:
            del ind.attrib[qattr]


def set_first_line_chars_on_element(element, chars_hundredths: int | None) -> None:
    """Set first-line indent in hundredths of a character on a paragraph/style element.

    Word stores a 2-character first-line indent as w:firstLineChars="200".
    """
    ppr = ensure_ppr(element)
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    clear_indent_attrs(ind)
    if chars_hundredths is None:
        return
    ind.set(qn("w:firstLineChars"), str(chars_hundredths))


def set_no_first_line_indent_on_element(element) -> None:
    ppr = ensure_ppr(element)
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    for attr in ("w:firstLine", "w:firstLineChars", "w:hanging", "w:hangingChars"):
        qattr = qn(attr)
        if qattr in ind.attrib:
            del ind.attrib[qattr]
    ind.set(qn("w:firstLine"), "0")


def set_style_first_line_chars(style, chars_hundredths: int | None) -> None:
    set_first_line_chars_on_element(style.element, chars_hundredths)


def set_style_no_first_line_indent(style) -> None:
    set_no_first_line_indent_on_element(style.element)


def paragraph_set_first_line_chars(paragraph, chars_hundredths: int | None) -> None:
    set_first_line_chars_on_element(paragraph._p, chars_hundredths)


def paragraph_set_no_first_line_indent(paragraph) -> None:
    set_no_first_line_indent_on_element(paragraph._p)


def paragraph_clear_direct_indent(paragraph) -> None:
    ppr = paragraph._p.find(qn("w:pPr"))
    if ppr is None:
        return
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        return
    clear_indent_attrs(ind)


def style_or_none(doc, name):
    try:
        return doc.styles[name]
    except Exception:
        return None


def twips_to_cm(value):
    if value is None:
        return None
    return round(value.twips / 567.0, 3)


def para_direct_indent_xml(paragraph):
    ppr = paragraph._p.find(qn("w:pPr"))
    if ppr is None:
        return {}
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        return {}
    out = {}
    for attr in ("firstLine", "firstLineChars", "left", "leftChars", "hanging", "hangingChars"):
        value = ind.get(qn(f"w:{attr}"))
        if value is not None:
            out[attr] = value
    return out


def style_indent_xml(style):
    ppr = style.element.find(qn("w:pPr"))
    if ppr is None:
        return {}
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        return {}
    out = {}
    for attr in ("firstLine", "firstLineChars", "left", "leftChars", "hanging", "hangingChars"):
        value = ind.get(qn(f"w:{attr}"))
        if value is not None:
            out[attr] = value
    return out


def main() -> None:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = ROOT / f"my_report_backup_before_indent_fix_{timestamp}.docx"
    shutil.copy2(DOCX_PATH, backup_path)

    doc = Document(DOCX_PATH)

    changed_styles = []

    # Body text: first-line indent two Chinese characters.
    # In WordprocessingML this is represented by firstLineChars="200".
    for style_name in ("Normal", "正文"):
        style = style_or_none(doc, style_name)
        if style is not None:
            set_style_first_line_chars(style, 200)
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            style.paragraph_format.line_spacing = Pt(20)
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(0)
            changed_styles.append(style_name)

    # Chapter title: centered, no first-line indent.
    for style_name in ("一级标题", "Heading 1", "标题 1"):
        style = style_or_none(doc, style_name)
        if style is not None:
            set_style_no_first_line_indent(style)
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style.paragraph_format.line_spacing = 1
            style.paragraph_format.space_before = Pt(24)
            style.paragraph_format.space_after = Pt(18)
            style.font.name = "黑体"
            style.font.size = Pt(16)
            style.font.bold = True
            changed_styles.append(style_name)

    # Second-level title: left aligned, no first-line indent.
    for style_name in ("二级标题", "Heading 2", "标题 2"):
        style = style_or_none(doc, style_name)
        if style is not None:
            set_style_no_first_line_indent(style)
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style.paragraph_format.line_spacing = 1
            style.paragraph_format.space_before = Pt(24)
            style.paragraph_format.space_after = Pt(6)
            style.font.name = "黑体"
            style.font.size = Pt(14)
            style.font.bold = False
            changed_styles.append(style_name)

    # Third-level title: first-line indent two Chinese characters.
    for style_name in ("三级标题", "Heading 3", "标题 3"):
        style = style_or_none(doc, style_name)
        if style is not None:
            set_style_first_line_chars(style, 200)
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style.paragraph_format.line_spacing = 1
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(6)
            style.font.name = "黑体"
            style.font.size = Pt(12)
            style.font.bold = False
            changed_styles.append(style_name)

    # TOC entries should not inherit direct first-line indent from body editing.
    for style_name in ("TOC 1", "TOC 2", "TOC 3", "toc 1", "toc 2", "toc 3"):
        style = style_or_none(doc, style_name)
        if style is not None:
            set_style_no_first_line_indent(style)
            changed_styles.append(style_name)

    heading_pat = re.compile(r"^[1-6](?:\.[1-9][0-9]*){0,2}\s+")
    fixed_paragraphs = []
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        style_name = paragraph.style.name
        if not text:
            continue

        is_heading_text = bool(heading_pat.match(text))
        if style_name in {"一级标题", "Heading 1", "标题 1"} or re.match(r"^[1-6]\s+", text):
            before = para_direct_indent_xml(paragraph)
            paragraph_set_no_first_line_indent(paragraph)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fixed_paragraphs.append((idx, style_name, text, before, para_direct_indent_xml(paragraph)))
        elif style_name in {"二级标题", "Heading 2", "标题 2"} or re.match(r"^[1-6]\.[1-9][0-9]*\s+", text):
            before = para_direct_indent_xml(paragraph)
            paragraph_set_no_first_line_indent(paragraph)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fixed_paragraphs.append((idx, style_name, text, before, para_direct_indent_xml(paragraph)))
        elif style_name in {"三级标题", "Heading 3", "标题 3"} or re.match(r"^[1-6]\.[1-9][0-9]*\.[1-9][0-9]*\s+", text):
            before = para_direct_indent_xml(paragraph)
            paragraph_set_first_line_chars(paragraph, 200)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fixed_paragraphs.append((idx, style_name, text, before, para_direct_indent_xml(paragraph)))
        elif style_name.lower().startswith("toc"):
            before = para_direct_indent_xml(paragraph)
            paragraph_clear_direct_indent(paragraph)
            fixed_paragraphs.append((idx, style_name, text, before, para_direct_indent_xml(paragraph)))

    doc.save(DOCX_PATH)

    # Reopen and sample result.
    reopened = Document(DOCX_PATH)
    samples = []
    for idx, paragraph in enumerate(reopened.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        if (
            idx < 160
            and paragraph.style.name.lower().startswith("toc")
            or paragraph.style.name in {"一级标题", "二级标题", "三级标题"}
        ):
            samples.append(
                {
                    "index": idx,
                    "style": paragraph.style.name,
                    "text": text[:100],
                    "direct_indent_xml": para_direct_indent_xml(paragraph),
                    "style_indent_xml": style_indent_xml(paragraph.style),
                    "direct_first_line_cm": twips_to_cm(paragraph.paragraph_format.first_line_indent),
                }
            )
            if len(samples) >= 60:
                break

    result = {
        "docx": str(DOCX_PATH),
        "backup": str(backup_path),
        "changed_styles": sorted(set(changed_styles)),
        "fixed_paragraph_count": len(fixed_paragraphs),
        "sample_after": samples,
    }
    OUT_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in result.items() if k != "sample_after"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
