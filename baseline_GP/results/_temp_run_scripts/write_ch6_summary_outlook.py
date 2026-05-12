from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(r"F:\pythonprojects")
DOCX_PATH = ROOT / "my_report.docx"
OUT_JSON = (
    ROOT
    / "baseline_GP"
    / "results"
    / "_temp_run_scripts"
    / "write_ch6_summary_outlook_report.json"
)


SUMMARY_PARAGRAPHS = [
    "本文围绕已知静态障碍地图下的海上可疑目标搜索任务，研究了单 USV 与双 USV 条件下的信息驱动搜索决策方法。该任务中，障碍地图在任务开始前已知，但目标真实位置和线索状态并不直接可见，USV 需要依靠传感器观测不断修正搜索状态，并在有限步数内完成尽可能有效的目标搜索。",
    "针对这一问题，本文首先构建了面向搜索决策的多源场表示。高斯过程回归用于描述弱线索的空间分布，target intensity 场用于刻画未发现目标的空间可能性，recency 场用于反映历史观测的新旧程度。在此基础上，本文将 anomaly-aware GP clue 与 intensity 融合为综合信息价值场，使路径规划能够根据当前搜索状态选择更有观测价值的区域。",
    "在单 USV 搜索方法中，本文将信息场驱动的候选区域选择、环形观测点生成、安全 A* 路径生成和短时域路径段评分结合起来，形成了在线滚动执行的搜索路径规划框架。实验结果表明，该方法能够在 open_water、obstacle_field 和 peninsula_passage 三类已知静态地图中完成大部分目标发现任务，并在 static 与 random_walk 两类目标运动模式下保持基本稳定的搜索能力。",
    "在双 USV 协同搜索方面，本文在单艇路径规划基础上引入共享团队状态、软责任区偏置、residual_map 顺序分配以及 reservation 安全执行机制。与 independent 对照模式相比，coordinated 模式在多数实验条件下提高了搜索完成程度，并降低了重复观测和跨责任区选择比例。尤其在静态目标和通道约束较强的地图中，显式团队分配能够更有效地避免两艘艇集中搜索相同区域，从而提升双艇协同搜索的整体效果。",
    "本文还进一步考察了 anomaly_upper_tail 采集方式的作用。实验结果显示，该方法在部分地图和目标运动模式下能够提高搜索完成程度或改善发现时间，但其收益并不稳定。anomaly_upper_tail 的作用本质上是改变 GP clue 层中上尾异常区域的采样优先级，而不是直接改变目标真实存在概率或路径可达性。因此，其效果会受到地图拓扑、目标运动方式以及双艇协同分配机制的共同影响。",
    "综上，本文完成了从场建模、单艇信息驱动路径规划到双艇协同搜索决策的完整方法构建与实验验证。该方法不依赖大量训练数据，而是通过可解释的信息场、候选视点、短时域路径评分和团队级分配机制实现搜索决策，为后续更真实海上环境中的多 USV 协同搜索研究提供了基础。",
]


OUTLOOK_PARAGRAPHS = [
    "本文仍有进一步扩展的空间。首先，当前软责任区先验由双艇初始位置和已知静态地图可达距离预先构造，在运行过程中保持不变。该设计有助于形成稳定分工，但在 random_walk 目标或局部高价值区域快速变化时，固定责任区可能降低系统对当前高价值区域的响应速度。后续可考虑根据 USV 实时位置、历史观测结果和线索场变化动态调整责任区，使团队分工能够随搜索过程自适应变化。",
    "其次，anomaly_upper_tail 仍属于较直接的 anomaly-aware 采集变体。实验表明，它在部分场景中有效，但并不总是优于 UCB。后续可以研究条件触发式 anomaly-aware 采集机制，使异常线索只在分布足够集中、热点较稳定或已有目标发现信息支持时参与搜索决策，从而减少早期异常偏置可能带来的误导。",
    "再次，本文实验仍基于二维栅格仿真环境，USV 的运动被抽象为离散栅格路径段执行。该设置便于验证搜索决策逻辑，但尚未充分体现真实 USV 的连续动力学、水流扰动、控制误差和传感器噪声。后续可将本文方法迁移到 HoloOcean 等海洋机器人仿真平台中，将离散路径段转换为连续航点跟踪任务，进一步验证方法在更接近真实海上环境中的可执行性。",
    "此外，本文以双 USV 作为多艇协同搜索的初步研究对象。对于三艘及以上 USV，团队分配顺序、重复搜索抑制和艇间冲突处理都会更加复杂，计算代价也会随候选组合增加而上升。后续研究可在保持集中式共享状态假设的基础上，进一步探索更一般的多 USV 搜索分配机制，并分析其在不同地图结构和目标运动模式下的适用边界。",
    "对于更大规模的多 USV 协同搜索任务，若系统需要同时处理动态环境、不确定通信、异构平台和复杂协同行为，基于多智能体强化学习的决策方法也具有进一步研究价值。该类方法可以通过训练学习多艇之间的协作策略，但通常需要大量仿真数据、稳定的训练环境和较高的可解释性验证。因此，后续可将本文的信息场建模和安全路径生成机制作为可解释的状态表示与约束基础，再探索其与多智能体学习方法的结合，而不是直接以端到端学习替代本文的规划框架。",
]


def insert_paragraph_after(paragraph: Paragraph, text: str, style_name: str = "Normal") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = style_name
    new_para.add_run(text)
    return new_para


def remove_paragraph(paragraph: Paragraph) -> None:
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)
    paragraph._p = paragraph._element = None


def find_body_chapter6_paragraphs(doc: Document) -> tuple[Paragraph, Paragraph, Paragraph, Paragraph]:
    """Find Chapter 6 body headings, skipping TOC entries.

    The document also contains TOC paragraphs such as "6.1 ...\t40". Matching by
    text alone is unsafe. Locate the body chapter heading first, then find the
    body section headings after it.
    """
    paragraphs = list(doc.paragraphs)
    h6_idx = None
    for idx, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if text.startswith("6 总结与展望") and paragraph.style.name == "一级标题":
            h6_idx = idx
            break
    if h6_idx is None:
        raise RuntimeError("Could not find body heading: 6 总结与展望")

    h61 = h62 = refs = None
    for paragraph in paragraphs[h6_idx + 1 :]:
        text = paragraph.text.strip()
        style = paragraph.style.name
        if h61 is None and text.startswith("6.1 课题总结") and style == "二级标题":
            h61 = paragraph
            continue
        if h62 is None and text.startswith("6.2 后续工作展望") and style == "二级标题":
            h62 = paragraph
            continue
        if h62 is not None and text.startswith("参考文献"):
            refs = paragraph
            break

    if h61 is None or h62 is None or refs is None:
        raise RuntimeError(
            "Could not locate Chapter 6 body headings and following references "
            f"(h61={h61 is not None}, h62={h62 is not None}, refs={refs is not None})"
        )
    return paragraphs[h6_idx], h61, h62, refs


def paragraph_index(doc: Document, target: Paragraph) -> int:
    for idx, paragraph in enumerate(doc.paragraphs):
        if paragraph._element is target._element:
            return idx
    raise RuntimeError("paragraph not found")


def clear_existing_between(doc: Document, start: Paragraph, end: Paragraph) -> int:
    start_idx = paragraph_index(doc, start)
    end_idx = paragraph_index(doc, end)
    if end_idx <= start_idx:
        return 0
    to_remove = list(doc.paragraphs[start_idx + 1 : end_idx])
    for paragraph in to_remove:
        remove_paragraph(paragraph)
    return len(to_remove)


def insert_paragraphs(anchor: Paragraph, paragraphs: list[str]) -> Paragraph:
    current = anchor
    for text in paragraphs:
        current = insert_paragraph_after(current, text, style_name="Normal")
    return current


def main() -> None:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = ROOT / f"my_report_backup_before_write_ch6_summary_outlook_{timestamp}.docx"
    shutil.copy2(DOCX_PATH, backup_path)

    doc = Document(DOCX_PATH)
    _, h61, h62, refs = find_body_chapter6_paragraphs(doc)

    removed_61 = clear_existing_between(doc, h61, h62)
    # Re-find after removing paragraphs because doc.paragraphs is regenerated by python-docx.
    _, h61, h62, refs = find_body_chapter6_paragraphs(doc)
    removed_62 = clear_existing_between(doc, h62, refs)

    _, h61, h62, _ = find_body_chapter6_paragraphs(doc)
    insert_paragraphs(h61, SUMMARY_PARAGRAPHS)
    # Re-find because insertion changes paragraph list.
    _, _, h62, _ = find_body_chapter6_paragraphs(doc)
    insert_paragraphs(h62, OUTLOOK_PARAGRAPHS)

    doc.save(DOCX_PATH)

    # Read back a concise check.
    check_doc = Document(DOCX_PATH)
    section_texts = []
    recording = False
    for paragraph in check_doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("6 总结与展望"):
            recording = True
        if recording and text:
            section_texts.append({"style": paragraph.style.name, "text": text})
        if recording and text.startswith("参考文献"):
            break

    report = {
        "docx": str(DOCX_PATH),
        "backup": str(backup_path),
        "removed_between_6_1_and_6_2": removed_61,
        "removed_between_6_2_and_references": removed_62,
        "inserted_6_1_paragraphs": len(SUMMARY_PARAGRAPHS),
        "inserted_6_2_paragraphs": len(OUTLOOK_PARAGRAPHS),
        "section_check": section_texts,
    }
    OUT_JSON.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in report.items() if k != "section_check"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
