from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def shape_text(shape) -> str:
    parts: list[str] = []
    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text or ""
        if text.strip():
            parts.append(text.strip())
    if getattr(shape, "has_table", False):
        rows = []
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        table_text = "\n".join(row for row in rows if row.strip())
        if table_text.strip():
            parts.append(table_text.strip())
    if getattr(shape, "has_chart", False):
        try:
            parts.append(f"[chart: {shape.chart.chart_title.text_frame.text}]")
        except Exception:
            parts.append("[chart]")
    return "\n".join(parts).strip()


def font_sizes(shape) -> list[float]:
    sizes: list[float] = []
    if not getattr(shape, "has_text_frame", False):
        return sizes
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                sizes.append(round(run.font.size.pt, 1))
    return sizes


def summarize_pptx(pptx_path: Path, out_dir: Path) -> dict:
    prs = Presentation(str(pptx_path))
    slides = []
    for idx, slide in enumerate(prs.slides, start=1):
        texts = []
        shape_rows = []
        picture_count = 0
        table_count = 0
        chart_count = 0
        group_count = 0
        for shape in slide.shapes:
            text = shape_text(shape)
            if text:
                texts.append(text)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture_count += 1
            if getattr(shape, "has_table", False):
                table_count += 1
            if getattr(shape, "has_chart", False):
                chart_count += 1
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                group_count += 1
            shape_rows.append(
                {
                    "name": getattr(shape, "name", ""),
                    "type": str(shape.shape_type),
                    "left": round(shape.left / 914400, 3),
                    "top": round(shape.top / 914400, 3),
                    "width": round(shape.width / 914400, 3),
                    "height": round(shape.height / 914400, 3),
                    "text_preview": re.sub(r"\s+", " ", text)[:140],
                    "font_sizes": font_sizes(shape)[:10],
                }
            )
        full_text = "\n".join(texts)
        title = texts[0].splitlines()[0].strip() if texts else ""
        slides.append(
            {
                "index": idx,
                "title": title,
                "text": full_text,
                "text_chars": len(full_text),
                "text_blocks": len(texts),
                "pictures": picture_count,
                "tables": table_count,
                "charts": chart_count,
                "groups": group_count,
                "shapes": shape_rows,
            }
        )
    data = {
        "path": str(pptx_path),
        "slide_count": len(prs.slides),
        "slide_width_in": round(prs.slide_width / 914400, 3),
        "slide_height_in": round(prs.slide_height / 914400, 3),
        "slides": slides,
    }
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "ppt_summary.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    lines = [
        f"# PPT summary",
        f"path: {pptx_path}",
        f"slides: {len(prs.slides)}",
        f"size: {data['slide_width_in']} x {data['slide_height_in']} in",
        "",
    ]
    for slide in slides:
        lines.append(
            f"## Slide {slide['index']:02d}: {slide['title'] or '(no title)'}"
        )
        lines.append(
            f"chars={slide['text_chars']}, text_blocks={slide['text_blocks']}, "
            f"pictures={slide['pictures']}, tables={slide['tables']}, "
            f"charts={slide['charts']}, groups={slide['groups']}"
        )
        lines.append(slide["text"].strip())
        lines.append("")
    (out_dir / "ppt_text.md").write_text("\n".join(lines), encoding="utf-8")
    return data


def export_pngs_with_powerpoint(pptx_path: Path, out_dir: Path) -> bool:
    try:
        import win32com.client  # type: ignore
    except Exception as exc:
        print(f"win32com unavailable: {exc}", file=sys.stderr)
        return False
    out_dir.mkdir(parents=True, exist_ok=True)
    app = None
    presentation = None
    try:
        app = win32com.client.DispatchEx("PowerPoint.Application")
        presentation = app.Presentations.Open(str(pptx_path), True, False, False)
        presentation.Export(str(out_dir), "PNG", 1920, 1080)
        return True
    except Exception as exc:
        print(f"PowerPoint export failed: {exc}", file=sys.stderr)
        return False
    finally:
        if presentation is not None:
            presentation.Close()
        if app is not None:
            app.Quit()


def make_contact_sheets(png_dir: Path, out_dir: Path, cols: int = 3, rows: int = 2) -> list[Path]:
    files = sorted(
        png_dir.glob("*.PNG"),
        key=lambda p: int(re.search(r"(\d+)", p.stem).group(1)) if re.search(r"(\d+)", p.stem) else 0,
    )
    if not files:
        files = sorted(
            png_dir.glob("*.png"),
            key=lambda p: int(re.search(r"(\d+)", p.stem).group(1)) if re.search(r"(\d+)", p.stem) else 0,
        )
    out_dir.mkdir(parents=True, exist_ok=True)
    sheet_paths = []
    thumb_w, thumb_h = 480, 270
    label_h = 34
    gap = 16
    margin = 24
    per_sheet = cols * rows
    for sheet_idx, start in enumerate(range(0, len(files), per_sheet), start=1):
        batch = files[start : start + per_sheet]
        canvas_w = margin * 2 + cols * thumb_w + (cols - 1) * gap
        canvas_h = margin * 2 + rows * (thumb_h + label_h) + (rows - 1) * gap
        canvas = Image.new("RGB", (canvas_w, canvas_h), "white")
        for i, file in enumerate(batch):
            r, c = divmod(i, cols)
            x = margin + c * (thumb_w + gap)
            y = margin + r * (thumb_h + label_h + gap)
            img = Image.open(file).convert("RGB")
            img = ImageOps.contain(img, (thumb_w, thumb_h), Image.Resampling.LANCZOS)
            bg = Image.new("RGB", (thumb_w, thumb_h), (245, 247, 250))
            bg.paste(img, ((thumb_w - img.width) // 2, (thumb_h - img.height) // 2))
            canvas.paste(bg, (x, y))
            label = f"Slide {start + i + 1:02d}"
            # Draw label using PIL default font to avoid font dependencies.
            from PIL import ImageDraw

            draw = ImageDraw.Draw(canvas)
            draw.text((x, y + thumb_h + 8), label, fill=(30, 41, 59))
        out_path = out_dir / f"contact_sheet_{sheet_idx:02d}.jpg"
        canvas.save(out_path, quality=92)
        sheet_paths.append(out_path)
    return sheet_paths


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    pptx_path = args.pptx.resolve()
    out_dir = args.out.resolve()
    if not pptx_path.exists():
        print(f"missing: {pptx_path}", file=sys.stderr)
        return 2
    data = summarize_pptx(pptx_path, out_dir)
    png_dir = out_dir / "slides_png"
    exported = export_pngs_with_powerpoint(pptx_path, png_dir)
    sheet_paths = make_contact_sheets(png_dir, out_dir / "contact_sheets") if exported else []
    print(json.dumps(
        {
            "pptx": str(pptx_path),
            "out": str(out_dir),
            "slides": data["slide_count"],
            "png_exported": exported,
            "contact_sheets": [str(p) for p in sheet_paths],
        },
        ensure_ascii=False,
        indent=2,
    ))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
