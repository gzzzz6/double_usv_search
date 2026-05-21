import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

prs = Presentation(r"E:\毕设PPT\USV_PPT_gemini.pptx")
slide5 = prs.slides[4]  # Slide 5

with open(r"F:\pythonprojects\gemini\template_style.txt", "w", encoding="utf-8") as f:
    f.write(f"=== Slide 5 Shape Details ===\n")
    for i, shape in enumerate(slide5.shapes):
        try:
            f.write(f"\nShape {i}: '{shape.name}', type={shape.shape_type}\n")
            f.write(f"  Pos: L={shape.left}, T={shape.top}, W={shape.width}, H={shape.height}\n")
            if hasattr(shape, 'text') and shape.text.strip():
                f.write(f"  Text: {repr(shape.text.strip()[:80])}\n")
            try:
                if shape.fill and shape.fill.type:
                    fc = shape.fill.fore_color
                    if fc.type == 1:
                        f.write(f"  Fill: #{shape.fill.fore_color.rgb}\n")
            except Exception:
                pass
            try:
                if shape.line:
                    lc = shape.line.color
                    if lc.type == 1:
                        f.write(f"  Line: #{shape.line.color.rgb}\n")
                    w = shape.line.width
                    if w:
                        f.write(f"  Line width: {w/12700:.1f}pt\n")
            except Exception:
                pass
            try:
                if hasattr(shape, 'text_frame'):
                    tf = shape.text_frame
                    for para in tf.paragraphs[:1]:
                        if para.runs:
                            run = para.runs[0]
                            f.write(f"  Font: name={run.font.name}, size={run.font.size}, bold={run.font.bold}\n")
                            if run.font.color and run.font.color.type == 1:
                                f.write(f"  FontColor: #{run.font.color.rgb}\n")
            except Exception:
                pass
        except Exception as e:
            f.write(f"  ERROR: {e}\n")
print("Done")
