import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation(r"E:\毕设PPT\USV_PPT_gemini.pptx")
print(f"Total slides: {len(prs.slides)}")
for i, slide in enumerate(prs.slides):
    texts = []
    for s in slide.shapes:
        if hasattr(s, 'text') and s.text.strip():
            texts.append(s.text.strip())
    # Find navigation text or placeholders
    nav_header = ""
    sub_title = ""
    for s in slide.shapes:
        if s.is_placeholder:
            idx = s.placeholder_format.idx
            if idx == 16:
                nav_header = s.text.strip()
            elif idx == 14:
                sub_title = s.text.strip()
    
    print(f"\nSlide {i+1}:")
    print(f"  Nav: {repr(nav_header)}")
    print(f"  SubTitle: {repr(sub_title)}")
    print(f"  First 3 texts: {repr(texts[:3])}")
