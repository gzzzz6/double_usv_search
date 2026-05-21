import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation(r"E:\毕设PPT\USV_PPT_gemini.pptx")
slide10 = prs.slides[9] # Slide 10

print(f"=== Slide 10 Shapes ===")
for i, shape in enumerate(slide10.shapes):
    print(f"Shape {i}: '{shape.name}', type={shape.shape_type}")
    if shape.is_placeholder:
        print(f"  Placeholder idx={shape.placeholder_format.idx}, text={repr(shape.text)}")
    elif hasattr(shape, 'text') and shape.text.strip():
        print(f"  Text: {repr(shape.text.strip()[:60])}")
