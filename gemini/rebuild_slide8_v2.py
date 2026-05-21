import sys
import json
import shutil
sys.stdout.reconfigure(encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

# Load content strings
with open(r"F:\pythonprojects\gemini\slide8_content.json", "r", encoding="utf-8") as f:
    T = json.load(f)

PPT_PATH = r"E:\毕设PPT\USV_PPT_gemini.pptx"
BAK_PATH = r"E:\毕设PPT\USV_PPT_gemini_bak2.pptx"
shutil.copy2(PPT_PATH, BAK_PATH)
print("Backup done.")

prs = Presentation(PPT_PATH)
slide = prs.slides[7]

# Colors matching the template (Slide 5)
C_CONTAINER = RGBColor(0x64, 0xA0, 0xFF)
C_CARD_BD   = RGBColor(0x38, 0x58, 0x94)
C_CARD_TX   = RGBColor(0x00, 0x70, 0xC0)
C_BG_FILL   = RGBColor(0xEC, 0xF4, 0xFC)
C_BG_BD     = RGBColor(0xBD, 0xD1, 0xE5)
C_ARROW     = RGBColor(0x64, 0x96, 0xFF)
C_DARK      = RGBColor(0x20, 0x2A, 0x34)
C_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
C_BLUE_HDR  = RGBColor(0x00, 0x70, 0xC0)
C_GREEN_HDR = RGBColor(0x00, 0x7A, 0x50)
C_GREEN_BD  = RGBColor(0x2E, 0x7D, 0x50)
C_GREEN_TX  = RGBColor(0x00, 0x60, 0x3A)

F_BODY = "Microsoft YaHei"

def I(x): return int(x * 914400)

# Delete body shapes (keep 0-14)
orig = len(slide.shapes)
for j in range(orig - 1, 14, -1):
    slide.shapes._spTree.remove(slide.shapes[j]._element)
print(f"Retained {len(slide.shapes)} template shapes.")

# Update placeholders
for shape in slide.shapes:
    if shape.is_placeholder:
        idx = shape.placeholder_format.idx
        if idx == 16:
            shape.text = T["slide8_title"]
            for p in shape.text_frame.paragraphs:
                p.font.name = "Source Han Sans CN"
        elif idx == 14:
            shape.text = T["slide8_subtitle"]
            for p in shape.text_frame.paragraphs:
                p.font.name = "Microsoft YaHei"
        elif idx in (1, 2):
            shape.text = "8"

# ---- Helpers ----
def container(left, top, w, h):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    s.fill.background()
    s.line.color.rgb = C_CONTAINER
    s.line.width = Pt(1.5)

def hdr(text, left, top, w, h, bg):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = bg
    s.line.fill.background()
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = I(0.08)
    tf.margin_top = tf.margin_bottom = I(0.02)
    p = tf.paragraphs[0]
    p.text = text; p.alignment = PP_ALIGN.CENTER
    p.font.name = F_BODY; p.font.size = Pt(12); p.font.bold = True
    p.font.color.rgb = C_WHITE

def infobox(left, top, w, h):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = C_BG_FILL
    s.line.color.rgb = C_BG_BD; s.line.width = Pt(1.1)

def card(title, body_lines, left, top, w, h,
         bd=None, tx=None, title_sz=11, body_sz=9.5):
    bd = bd or C_CARD_BD; tx = tx or C_CARD_TX
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    s.fill.background(); s.line.color.rgb = bd; s.line.width = Pt(1.0)
    tb = slide.shapes.add_textbox(left+I(0.1), top+I(0.07), w-I(0.2), h-I(0.14))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = I(0)
    p = tf.paragraphs[0]
    p.text = title; p.font.name = F_BODY; p.font.size = Pt(title_sz)
    p.font.bold = True; p.font.color.rgb = tx; p.space_after = Pt(4)
    for line in body_lines:
        pb = tf.add_paragraph()
        pb.text = line; pb.font.name = F_BODY; pb.font.size = Pt(body_sz)
        pb.font.color.rgb = C_DARK; pb.space_before = Pt(3); pb.line_spacing = 1.2

def step_card(label, text, left, top, w, h, bd=None, tx=None):
    bd = bd or C_CARD_BD; tx = tx or C_CARD_TX
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    s.fill.background(); s.line.color.rgb = bd; s.line.width = Pt(1.0)
    tb = slide.shapes.add_textbox(left+I(0.08), top+I(0.06), w-I(0.16), h-I(0.12))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = I(0)
    p = tf.paragraphs[0]
    p.text = label; p.font.name = F_BODY; p.font.size = Pt(10)
    p.font.bold = True; p.font.color.rgb = tx; p.space_after = Pt(4)
    pb = tf.add_paragraph()
    pb.text = text; pb.font.name = F_BODY; pb.font.size = Pt(8.5)
    pb.font.color.rgb = C_DARK; pb.space_before = Pt(2); pb.line_spacing = 1.2

def formula_box(formula, left, top, w, h, label=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor(0xF0,0xF6,0xFF)
    s.line.color.rgb = C_BG_BD; s.line.width = Pt(1.1)
    tb = slide.shapes.add_textbox(left+I(0.1), top+I(0.06), w-I(0.2), h-I(0.12))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = I(0)
    if label:
        pl = tf.paragraphs[0]
        pl.text = label; pl.font.name = F_BODY; pl.font.size = Pt(8.5)
        pl.font.bold = True; pl.font.color.rgb = C_CARD_TX; pl.space_after = Pt(3)
        pf = tf.add_paragraph()
    else:
        pf = tf.paragraphs[0]
    pf.text = formula; pf.font.name = "Consolas"
    pf.font.size = Pt(10.5); pf.font.bold = True; pf.font.color.rgb = C_DARK

def txtbox(text, left, top, w, h, sz=10, bold=False, color=None, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(left, top, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = I(0)
    p = tf.paragraphs[0]
    p.text = text; p.alignment = align; p.font.name = F_BODY
    p.font.size = Pt(sz); p.font.bold = bold
    p.font.color.rgb = color or C_DARK

def arrow_r(left, top):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, I(0.18), I(0.1))
    a.fill.solid(); a.fill.fore_color.rgb = C_ARROW; a.line.fill.background()

# ===============================================================
# Build content
# ===============================================================

# Intro box
infobox(I(0.45), I(1.28), I(12.42), I(0.38))
txtbox(T["intro"], I(0.57), I(1.33), I(12.18), I(0.30), sz=9.8, color=C_DARK)

T_COL = I(1.76); H_COL = I(4.35)

# LEFT: Miss Update container
L_L = I(0.45); W_L = I(5.6)
container(L_L, T_COL, W_L, H_COL)
hdr(T["miss_header"], L_L+I(0.08), T_COL+I(0.08), W_L-I(0.16), I(0.36), C_BLUE_HDR)

card(T["miss_event_title"], T["miss_event_body"],
     L_L+I(0.12), T_COL+I(0.54), W_L-I(0.24), I(0.75))

formula_box(T["miss_formula"],
            L_L+I(0.12), T_COL+I(1.38), W_L-I(0.24), I(0.52),
            label=T["miss_formula_label"])

card(T["miss_props_title"], T["miss_props_body"],
     L_L+I(0.12), T_COL+I(2.0), W_L-I(0.24), I(1.5), body_sz=9)

txtbox(T["miss_compare"],
       L_L+I(0.12), T_COL+I(3.63), W_L-I(0.24), I(0.24),
       sz=8.5, color=RGBColor(0x33,0x77,0x33))

# RIGHT: Hit Update container
L_R = I(6.2); W_R = I(6.55)
container(L_R, T_COL, W_R, H_COL)
hdr(T["hit_header"], L_R+I(0.08), T_COL+I(0.08), W_R-I(0.16), I(0.36), C_GREEN_HDR)

card(T["hit_event_title"], T["hit_event_body"],
     L_R+I(0.12), T_COL+I(0.54), W_R-I(0.24), I(0.78),
     bd=C_GREEN_BD, tx=C_GREEN_TX)

# 3-step flow
T_ST = T_COL + I(1.44)
SW3  = (W_R - I(0.24) - I(0.44)) / 3
step_card(T["step1_label"], T["step1_body"],
          L_R+I(0.12), T_ST, SW3, I(1.55), bd=C_GREEN_BD, tx=C_GREEN_TX)

arrow_r(L_R+I(0.12)+SW3+I(0.02), T_ST+I(0.65))

step_card(T["step2_label"], T["step2_body"],
          L_R+I(0.12)+SW3+I(0.22), T_ST, SW3, I(1.55), bd=C_GREEN_BD, tx=C_GREEN_TX)

arrow_r(L_R+I(0.12)+SW3*2+I(0.24), T_ST+I(0.65))

step_card(T["step3_label"], T["step3_body"],
          L_R+I(0.12)+SW3*2+I(0.44), T_ST, SW3, I(1.55), bd=C_GREEN_BD, tx=C_GREEN_TX)

formula_box(T["hit_formula"],
            L_R+I(0.12), T_COL+I(3.1), W_R-I(0.24), I(0.65),
            label=T["hit_formula_label"])

# Design intent: simple body-only note box
s_note = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
    L_R+I(0.12), T_COL+I(3.87), W_R-I(0.24), I(0.36))
s_note.fill.background(); s_note.line.color.rgb = C_GREEN_BD; s_note.line.width = Pt(1.0)
tb_n = slide.shapes.add_textbox(L_R+I(0.2), T_COL+I(3.91), W_R-I(0.4), I(0.30))
tf_n = tb_n.text_frame; tf_n.word_wrap = True
tf_n.margin_left = tf_n.margin_right = tf_n.margin_top = tf_n.margin_bottom = I(0)
pn = tf_n.paragraphs[0]
pn.text = T["hit_intent_body"][0]
pn.font.name = F_BODY; pn.font.size = Pt(8.5); pn.font.color.rgb = C_DARK; pn.line_spacing = 1.2

# Bottom summary bar
infobox(I(0.45), I(6.18), I(12.42), I(0.55))
txtbox(T["summary"], I(0.57), I(6.22), I(12.18), I(0.47), sz=9.2, color=C_DARK)

prs.save(PPT_PATH)
print(f"Saved: {PPT_PATH}")
print(f"Total shapes on slide 8: {len(slide.shapes)}")
