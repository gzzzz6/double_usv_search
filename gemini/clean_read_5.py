import sys
from pptx import Presentation

try:
    prs = Presentation(r"E:\毕设PPT\USV_PPT.pptx")
    with open(r"F:\pythonprojects\gemini\slide_texts_utf8.txt", "w", encoding="utf-8") as f:
        for i in range(5):
            f.write(f"=== Slide {i+1} ===\n")
            slide = prs.slides[i]
            for shape in slide.shapes:
                if hasattr(shape, 'text') and shape.text.strip():
                    f.write(f"[{shape.name}] {shape.text.strip()}\n")
            f.write("\n")
    print("Successfully wrote slide texts to slide_texts_utf8.txt")
except Exception as e:
    print(f"Error: {e}")
