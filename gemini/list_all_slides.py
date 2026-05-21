from pptx import Presentation
prs = Presentation(r"E:\毕设PPT\USV_PPT_gemini.pptx")
print(f"Total slides: {len(prs.slides)}")
for i, slide in enumerate(prs.slides):
    title = ""
    # Find title shape if possible
    for shape in slide.shapes:
        if shape.is_placeholder and shape.placeholder_format.idx == 16:
            title = shape.text.strip() if hasattr(shape, 'text') else ""
            break
    if not title:
        # try search shapes for title-like shapes
        for shape in slide.shapes:
            if hasattr(shape, 'text') and "多源搜索状态建模" in shape.text:
                title = "多源搜索状态建模 (Custom)"
                break
    print(f"Slide {i+1}: title='{title}'")
