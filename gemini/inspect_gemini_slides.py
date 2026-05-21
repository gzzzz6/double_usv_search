import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation(r"E:\毕设PPT\USV_PPT_gemini.pptx")
print(f"Total slides in gemini: {len(prs.slides)}")
for i, slide in enumerate(prs.slides):
    title = ""
    for shape in slide.shapes:
        if shape.is_placeholder and shape.placeholder_format.idx == 14:
            title = shape.text
            break
        if hasattr(shape, 'text') and '单艇' in shape.text:
            title = shape.text[:50]
    print(f"Slide {i+1}: {repr(title)}")
