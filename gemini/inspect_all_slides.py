import sys
from pptx import Presentation

prs = Presentation(r"E:\毕设PPT\USV_PPT_gemini.pptx")
with open(r"F:\pythonprojects\gemini\check_all_slides.txt", "w", encoding="utf-8") as f:
    f.write(f"Total slides: {len(prs.slides)}\n\n")
    for idx, slide in enumerate(prs.slides):
        title = ""
        # Find title shape if possible
        for shape in slide.shapes:
            if shape.is_placeholder and shape.placeholder_format.idx == 16:
                title = shape.text.strip() if hasattr(shape, 'text') else ""
                break
        if not title:
            for shape in slide.shapes:
                if hasattr(shape, 'text') and shape.text.strip():
                    txt = shape.text.strip()
                    if len(txt) < 50:
                        title = txt
                        break
        f.write(f"Slide {idx+1}: title='{title}'\n")
print("Done")
