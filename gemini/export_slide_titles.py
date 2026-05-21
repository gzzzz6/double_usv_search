import sys
from pptx import Presentation

prs = Presentation(r"E:\毕设PPT\USV_PPT_gemini.pptx")
with open(r"F:\pythonprojects\gemini\slide_titles_utf8.txt", "w", encoding="utf-8") as f:
    f.write(f"Total slides: {len(prs.slides)}\n")
    for i, slide in enumerate(prs.slides):
        title = ""
        # Look for placeholder title (idx 16)
        for shape in slide.shapes:
            if shape.is_placeholder and shape.placeholder_format.idx == 16:
                title = shape.text.strip() if hasattr(shape, 'text') else ""
                break
        if not title:
            for shape in slide.shapes:
                if hasattr(shape, 'text') and shape.text.strip():
                    txt = shape.text.strip()
                    if len(txt) < 50 and any(c in txt for c in "多源搜索状态建模背景和意义总体技术路线"):
                        title = txt
                        break
        f.write(f"Slide {i+1}: title='{title}'\n")
print("Done")
