from pptx import Presentation
prs = Presentation(r"E:\毕设PPT\USV_PPT.pptx")
slide = prs.slides[6] # slide 7 is index 6
print("=== Slide 7 Shapes ===")
for j, shape in enumerate(slide.shapes):
    print(f"Shape {j}: type={shape.shape_type}, name={shape.name}")
    if hasattr(shape, 'text') and shape.text.strip():
        print(f"  Text: {shape.text.strip()}")
    if hasattr(shape, 'width'):
        print(f"  Geometry: left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}")
