"""
Build premium academic-style slides for Part 2 '多源搜索状态建模' in E:\\毕设PPT\\USV_PPT_gemini.pptx.
Uses frontend-design guidelines:
- Clear typography hierarchy (Microsoft YaHei + Consolas).
- Harmonious color scheme (Steel Blue for GP Clue, Forest Green for Intensity, Amber for Recency).
- Unique layered layouts with custom offset shadows for tactile card realism.
- Mathematical formulas presented inside clean highlight blocks.
- 100% editable native shapes, preserving navigation template and slide backgrounds.
"""
import sys
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

# Reconfigure stdout to use UTF-8 just in case
sys.stdout.reconfigure(encoding='utf-8')

PPT_PATH = r"E:\毕设PPT\USV_PPT_gemini.pptx"
BAK_PATH = r"E:\毕设PPT\USV_PPT_gemini_bak.pptx"

# 1. Back up original file
shutil.copy2(PPT_PATH, BAK_PATH)
print(f"[Backup] Saved original file to {BAK_PATH}")

# Load presentation
prs = Presentation(PPT_PATH)
print(f"[Load] Presentation loaded. Total slides: {len(prs.slides)}")

# ================================================================
# Design System Tokens
# ================================================================
# Fonts
F_TITLE = "Microsoft YaHei"
F_BODY = "Microsoft YaHei"
F_MATH = "Consolas"

# Color Palettes
# Blues (GP Clue theme)
C_BLUE_DEEP   = RGBColor(28, 58, 140)    # Dominant title blue
C_BLUE_MID    = RGBColor(43, 108, 235)   # Vibrant accent blue
C_BLUE_LIGHT  = RGBColor(240, 246, 255)  # Soft background tint
C_BLUE_BORDER = RGBColor(191, 219, 254)  # Light border blue

# Greens (Intensity theme)
C_GREEN_DEEP   = RGBColor(15, 76, 50)     # Dominant title green
C_GREEN_MID    = RGBColor(22, 163, 74)    # Vibrant accent green
C_GREEN_LIGHT  = RGBColor(240, 253, 244)  # Soft background tint
C_GREEN_BORDER = RGBColor(187, 247, 208)  # Light border green

# Ambers (Recency theme)
C_AMBER_DEEP   = RGBColor(124, 45, 18)    # Dominant title amber
C_AMBER_MID    = RGBColor(217, 119, 6)    # Vibrant accent amber
C_AMBER_LIGHT  = RGBColor(254, 252, 232)  # Soft background tint
C_AMBER_BORDER = RGBColor(254, 240, 138)  # Light border amber

# Grays & Slate (Neutral theme)
C_SLATE_DEEP   = RGBColor(30, 41, 59)     # Deep dark charcoal for text
C_SLATE_MID    = RGBColor(71, 85, 105)    # Medium gray for body
C_SLATE_LIGHT  = RGBColor(248, 250, 252)  # Soft light neutral card
C_SLATE_BORDER = RGBColor(226, 232, 240)  # Neutral border
C_SLATE_SHADOW = RGBColor(226, 232, 240)  # Shadow color

C_WHITE        = RGBColor(255, 255, 255)

# ================================================================
# Elegant Shape Drawing Helpers
# ================================================================
def draw_premium_card(slide, title, bullets, left, top, width, height,
                      theme_color, bg_color, border_color, tag_text=None,
                      is_math=False):
    """
    Draws a premium double-layered card with tactile shadow, left-accent bar,
    and a top accent tag. All text sizes are optimized for academic slide density.
    """
    # 1. Shadow card (offset bottom-right by 0.04")
    shadow = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.04), top + Inches(0.04), width, height)
    shadow.fill.solid()
    shadow.fill.fore_color.rgb = C_SLATE_SHADOW
    shadow.line.fill.background()
    
    # 2. Main card
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.5)
    
    # 3. Left accent bar
    bar_w = Inches(0.04)
    bar_h = height - Inches(0.2)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.02), top + Inches(0.1), bar_w, bar_h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = theme_color
    bar.line.fill.background()
    
    # 4. Top Tag (if tag_text is provided)
    if tag_text:
        tag_w = Inches(1.3)
        tag_h = Inches(0.24)
        tag = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.15), top - Inches(0.12), tag_w, tag_h)
        tag.fill.solid()
        tag.fill.fore_color.rgb = theme_color
        tag.line.fill.background()
        
        tf_t = tag.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = Inches(0)
        p_t = tf_t.paragraphs[0]
        p_t.alignment = PP_ALIGN.CENTER
        p_t.text = tag_text
        p_t.font.name = F_TITLE
        p_t.font.size = Pt(8.5)
        p_t.font.bold = True
        p_t.font.color.rgb = C_WHITE

    # 5. Text container
    tb = slide.shapes.add_textbox(left + Inches(0.16), top + Inches(0.16), width - Inches(0.28), height - Inches(0.28))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = F_TITLE
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = theme_color
    p.space_after = Pt(8)
    
    for idx, b in enumerate(bullets):
        p_b = tf.add_paragraph()
        # Handle bullet text
        p_b.text = "•  " + b
        p_b.font.name = F_BODY
        p_b.font.size = Pt(9.5)
        p_b.font.color.rgb = C_SLATE_DEEP
        p_b.space_before = Pt(4)
        p_b.line_spacing = 1.15

def draw_formula_box(slide, formula_str, left, top, width, height, label=None, theme_color=C_SLATE_MID):
    """
    Draws a highly polished high-contrast math formula container with a label.
    """
    # Card back
    shadow = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.03), top + Inches(0.03), width, height)
    shadow.fill.solid()
    shadow.fill.fore_color.rgb = C_SLATE_SHADOW
    shadow.line.fill.background()

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = C_SLATE_LIGHT
    card.line.color.rgb = C_SLATE_BORDER
    card.line.width = Pt(1)

    tb = slide.shapes.add_textbox(left + Inches(0.12), top + Inches(0.08), width - Inches(0.24), height - Inches(0.16))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)

    if label:
        p_l = tf.paragraphs[0]
        p_l.text = label
        p_l.font.name = F_TITLE
        p_l.font.size = Pt(8.5)
        p_l.font.bold = True
        p_l.font.color.rgb = theme_color
        p_l.space_after = Pt(4)
        
        p_f = tf.add_paragraph()
        p_f.text = formula_str
        p_f.font.name = F_MATH
        p_f.font.size = Pt(11.5)
        p_f.font.bold = True
        p_f.font.color.rgb = C_SLATE_DEEP
    else:
        p_f = tf.paragraphs[0]
        p_f.text = formula_str
        p_f.font.name = F_MATH
        p_f.font.size = Pt(12)
        p_f.font.bold = True
        p_f.font.color.rgb = C_SLATE_DEEP

def draw_arrow_right(slide, left, top, width=Inches(0.3), height=Inches(0.15), color=C_SLATE_BORDER):
    """Draws a neat solid right arrow shape."""
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = color
    arrow.line.fill.background()

def draw_arrow_down(slide, left, top, width=Inches(0.15), height=Inches(0.3), color=C_SLATE_BORDER):
    """Draws a neat solid down arrow shape."""
    arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, left, top, width, height)
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = color
    arrow.line.fill.background()

# ================================================================
# Slide Processing Workflow
# ================================================================
# Slide indexes 5, 6, 7, 8 in presentation correspond to Slide 6, 7, 8, 9

SLIDE_DATA = [
    # ------------------------------------------------------------
    # Slide 6: 总体框架与要素定义
    # ------------------------------------------------------------
    {
        "index": 5,
        "title": "多源搜索状态建模",
        "subtitle": "三源状态融合驱动搜索价值判断",
        "builder": lambda slide: [
            # Background intro note
            draw_premium_card(
                slide,
                "三维互补的动态信息场体系",
                [
                    "在搜索任务中，USV的路径决策不能仅依赖静态的几何障碍信息，必须感知目标的动态线索分布和搜索覆盖历史。",
                    "本工作构建了线索估计场(GP Clue)、目标存在信念场(Intensity)和观测时效场(Recency)三维并行的建模框架，并融合成综合信息价值场指导决策。"
                ],
                Inches(0.6), Inches(1.5), Inches(12.13), Inches(1.0),
                C_SLATE_MID, C_SLATE_LIGHT, C_SLATE_BORDER, tag_text="核心设计思路"
            ),
            
            # Column 1: GP Clue
            draw_premium_card(
                slide,
                "① 线索估计场 GP Clue",
                [
                    "采用高斯过程回归(GPR)对传感器反馈的带噪局部线索进行空间外推与连续建模。",
                    "引入置信上限(UCB)理论，权衡'开发'(高先验线索区域)与'探索'(高认知不确定性区域)。",
                    "为路径规划提供前瞻性的线索引导，驱动USV主动追踪可疑残留迹象。"
                ],
                Inches(0.6), Inches(2.7), Inches(3.8), Inches(2.3),
                C_BLUE_MID, C_BLUE_LIGHT, C_BLUE_BORDER, tag_text="高斯空间推断"
            ),
            
            # Column 2: Intensity
            draw_premium_card(
                slide,
                "② 目标信念场 Intensity",
                [
                    "基于不完美传感器扫描模型，对每个离散网格维护目标存在的概率信念。",
                    "利用贝叶斯法则，在探测到疑似信号时上调概率，在空扫未检出时下调概率。",
                    "信念场随扫描进行而逐步收敛，为规避重复覆盖提供科学的'剩余价值'度量。"
                ],
                Inches(4.76), Inches(2.7), Inches(3.8), Inches(2.3),
                C_GREEN_MID, C_GREEN_LIGHT, C_GREEN_BORDER, tag_text="贝叶斯概率信念"
            ),
            
            # Column 3: Recency
            draw_premium_card(
                slide,
                "③ 观测时效场 Recency",
                [
                    "基于时间戳累加的递推机制，统计每个自由空间网格自上次被扫描所经过的相对时间步数。",
                    "通过指数衰减函数将步数映射为[0,1]的刷新状态权值，控制系统对旧信息的遗忘速率。",
                    "不直接参与信念更新，作为独立的评分项引导路径规划，确保高概率区被周期性复扫。"
                ],
                Inches(8.93), Inches(2.7), Inches(3.8), Inches(2.3),
                C_AMBER_MID, C_AMBER_LIGHT, C_AMBER_BORDER, tag_text="时效衰减机制"
            ),
            
            # Down Arrows to Fusion
            draw_arrow_down(slide, Inches(2.5), Inches(5.08), Inches(0.12), Inches(0.24), C_BLUE_MID),
            draw_arrow_down(slide, Inches(6.66), Inches(5.08), Inches(0.12), Inches(0.24), C_GREEN_MID),
            
            # Bottom Fusion Area
            draw_formula_box(
                slide,
                "search_info_map(x) = 0.5 × norm( GP_Clue(x) ) + 0.5 × norm( Intensity(x) )",
                Inches(0.6), Inches(5.4), Inches(8.0), Inches(0.6),
                label="▼ 决策底座：双场线性归一化加权融合", theme_color=C_SLATE_DEEP
            ),
            
            # Recency explanation card
            draw_premium_card(
                slide,
                "ℹ️ Recency 独立解耦设计",
                [
                    "时效维度不进入信息场融合，而是作为独立的偏置乘积项参与规划评分，避免稀释和污染核心信念浓度。"
                ],
                Inches(8.93), Inches(5.2), Inches(3.8), Inches(0.8),
                C_AMBER_MID, C_WHITE, C_AMBER_BORDER, tag_text=None
            )
        ]
    },
    
    # ------------------------------------------------------------
    # Slide 7: 高斯过程线索推断
    # ------------------------------------------------------------
    {
        "index": 6,
        "title": "多源搜索状态建模",
        "subtitle": "线索估计场 GP Clue —— 基于高斯过程的空间推断",
        "builder": lambda slide: [
            # Left column: input & kernel
            draw_premium_card(
                slide,
                "系统输入与空间关联核",
                [
                    "系统输入：历史沿途网格的线索稀疏观测样本集 D = {(x_i, y_i)}，由于传感器干扰，局部样本存在带噪特性。",
                    "核函数选择：采用径向基函数(RBF)来描述空间任意两点间的线索关联度：\n   k(x, x') = exp( - ||x - x'||² / (2 * l²) )\n其中长度尺度参数 l 调控空间相关性的衰减距离。"
                ],
                Inches(0.6), Inches(1.5), Inches(3.8), Inches(3.5),
                C_BLUE_MID, C_BLUE_LIGHT, C_BLUE_BORDER, tag_text="建模输入"
            ),
            
            # Middle column: formulas
            draw_premium_card(
                slide,
                "高斯过程后验推断",
                [
                    "利用已获取的观测，回归推断全搜索地图上任意测试点 x* 处的线索强度均值及其不确定性(方差)："
                ],
                Inches(4.76), Inches(1.5), Inches(3.8), Inches(3.5),
                C_BLUE_DEEP, C_WHITE, C_BLUE_BORDER, tag_text="数学推断"
            ),
            # Formula boxes inside middle column
            draw_formula_box(
                slide,
                "μ(x*) = k*ᵀ (K + σn² I)⁻¹ y",
                Inches(4.9), Inches(2.4), Inches(3.5), Inches(0.55),
                label="后验均值(预测强度)", theme_color=C_BLUE_MID
            ),
            draw_formula_box(
                slide,
                "σ²(x*) = k(x*, x*) - k*ᵀ (K + σn² I)⁻¹ k*",
                Inches(4.9), Inches(3.25), Inches(3.5), Inches(0.55),
                label="后验方差(不确定度)", theme_color=C_BLUE_MID
            ),
            # Description under formulas
            draw_premium_card(
                slide,
                "物理意义",
                [
                    "均值场 μ(x*) 指引USV追踪高强度迹象，方差场 σ²(x*) 反映环境认知的空缺程度，是自适应主动搜索的科学度量。"
                ],
                Inches(4.9), Inches(4.05), Inches(3.5), Inches(0.8),
                C_SLATE_MID, C_SLATE_LIGHT, C_SLATE_BORDER, tag_text=None
            ),
            
            # Right column: UCB
            draw_premium_card(
                slide,
                "探索与开发的科学权衡 (UCB)",
                [
                    "为了防止USV陷入局部最优导致漏检目标，路径决策引入置信上限(Upper Confidence Bound)作为动作选择准则："
                ],
                Inches(8.93), Inches(1.5), Inches(3.8), Inches(3.5),
                C_BLUE_DEEP, C_WHITE, C_BLUE_BORDER, tag_text="自适应权衡"
            ),
            draw_formula_box(
                slide,
                "GP_Clue(x) = μ(x) + β · σ(x)",
                Inches(9.07), Inches(2.35), Inches(3.5), Inches(0.55),
                label="▼ 置信上限综合评分公式", theme_color=C_BLUE_MID
            ),
            draw_premium_card(
                slide,
                "μ(x) 开发 (Exploitation)",
                [
                    "根据现有估计，优先探索最有可能残留线索的区域，提升搜索时效。"
                ],
                Inches(9.07), Inches(3.1), Inches(3.5), Inches(0.8),
                C_BLUE_MID, C_BLUE_LIGHT, C_BLUE_BORDER, tag_text=None
            ),
            draw_premium_card(
                slide,
                "β · σ(x) 探索 (Exploration)",
                [
                    "优先探索估计方差大(认知空缺)的陌生未知区域，增加全局覆盖度。"
                ],
                Inches(9.07), Inches(4.05), Inches(3.5), Inches(0.8),
                C_AMBER_MID, C_AMBER_LIGHT, C_AMBER_BORDER, tag_text=None
            ),
            
            # Bottom warning/engineering note
            draw_premium_card(
                slide,
                "⚠️ 工程计算复杂度优化",
                [
                    "高斯过程后验均值与方差求逆的计算复杂度为 O(N³)。为保证USV在线运动决策的实时性，本系统设计了滑动观测点采样上限(滑窗机制)，在满足鲁棒度的前提下大幅降低计算开销。"
                ],
                Inches(0.6), Inches(5.15), Inches(12.13), Inches(0.7),
                C_SLATE_MID, C_SLATE_LIGHT, C_SLATE_BORDER, tag_text=None
            )
        ]
    },
    
    # ------------------------------------------------------------
    # Slide 8: 目标存在信念场 Intensity
    # ------------------------------------------------------------
    {
        "index": 7,
        "title": "多源搜索状态建模",
        "subtitle": "目标存在信念场 Intensity —— 贝叶斯递推更新",
        "builder": lambda slide: [
            # Top: Sensor model
            draw_premium_card(
                slide,
                "不完美传感器探测概率模型",
                [
                    "物理约束：USV携带的声纳/光电载荷并非绝对完美，存在一定漏检率和虚警率：\n  • 检测概率 P_d = P(检测到目标 | 目标存在) < 1.0  (存在漏检风险)\n  • 虚警概率 P_f = P(检测到目标 | 目标不存在) > 0.0  (存在假信号误报风险)",
                    "本工作构建贝叶斯递推估计网络，利用时序上的多次连续扫描更新，逐步消减噪声与误检，获取稳健的目标存在后验概率。"
                ],
                Inches(0.6), Inches(1.5), Inches(12.13), Inches(1.15),
                C_SLATE_DEEP, C_SLATE_LIGHT, C_SLATE_BORDER, tag_text="核心观测假设"
            ),
            
            # Left: Hit update
            draw_premium_card(
                slide,
                "命中更新 (Hit Update) —— 目标概率上调",
                [
                    "当传感器在某网格 x 报告探测到疑似目标信号时，说明存在目标概率增高，利用贝叶斯法则上调其目标存在概率信念："
                ],
                Inches(0.6), Inches(2.8), Inches(5.8), Inches(2.9),
                C_GREEN_DEEP, C_WHITE, C_GREEN_BORDER, tag_text="探测到疑似信号"
            ),
            draw_formula_box(
                slide,
                "P(T=1 | hit) = P_d · P(T=1) / [ P_d · P(T=1) + P_f · (1 - P(T=1)) ]",
                Inches(0.76), Inches(3.65), Inches(5.48), Inches(0.6),
                label="▼ 命中贝叶斯后验更新公式 (Probability Escalation)", theme_color=C_GREEN_MID
            ),
            draw_premium_card(
                slide,
                "决策响应",
                [
                    "后验概率上调将大幅增加该网格的当前搜索价值，吸引USV快速前来进行局部精细化探测与持续跟踪。"
                ],
                Inches(0.76), Inches(4.45), Inches(5.48), Inches(1.1),
                C_GREEN_MID, C_GREEN_LIGHT, C_GREEN_BORDER, tag_text=None
            ),
            
            # Right: Miss update
            draw_premium_card(
                slide,
                "未命中更新 (Miss Update) —— 目标概率下调",
                [
                    "当传感器完整覆盖扫描了网格 x 但并未探测到目标信号时，说明目标在此的概率降低，向下衰减其信念概率："
                ],
                Inches(6.93), Inches(2.8), Inches(5.8), Inches(2.9),
                C_AMBER_DEEP, C_WHITE, C_AMBER_BORDER, tag_text="空扫未检出信号"
            ),
            draw_formula_box(
                slide,
                "P(T=1 | miss) = (1 - P_d)·P(T=1) / [ (1 - P_d)·P(T=1) + (1 - P_f)·(1 - P(T=1)) ]",
                Inches(7.09), Inches(3.65), Inches(5.48), Inches(0.6),
                label="▼ 未命中贝叶斯后验更新公式 (Probability Decay)", theme_color=C_AMBER_MID
            ),
            draw_premium_card(
                slide,
                "物理内涵",
                [
                    "因为存在漏检可能(P_d < 1)，所以空扫一次后的后验概率并不直接归零。这既防止了漏检导致盲目漏掉目标，又有效降低了该网格近期被重复扫描的无意义消耗。"
                ],
                Inches(7.09), Inches(4.45), Inches(5.48), Inches(1.1),
                C_AMBER_MID, C_AMBER_LIGHT, C_AMBER_BORDER, tag_text=None
            ),
            
            # Bottom summary: convergence
            draw_premium_card(
                slide,
                "🎯 动态收敛属性",
                [
                    "在连续扫描下，若某处存在目标，随着多次命中信号的叠加，其目标信念概率会迅速收敛并稳定在极高的置信上限(P ≈ 1.0)；反之则快速趋于 0。这为规划层提供了实时的'剩余搜索净收益'评估。"
                ],
                Inches(0.6), Inches(5.85), Inches(12.13), Inches(0.7),
                C_GREEN_MID, C_SLATE_LIGHT, C_GREEN_BORDER, tag_text=None
            )
        ]
    },
    
    # ------------------------------------------------------------
    # Slide 9: Recency 与综合信息价值场构建
    # ------------------------------------------------------------
    {
        "index": 8,
        "title": "多源搜索状态建模",
        "subtitle": "观测时效场 Recency 与综合信息价值场构建",
        "builder": lambda slide: [
            # Left: Recency
            draw_premium_card(
                slide,
                "观测时效场 Recency 的状态递推",
                [
                    "背景动机：海上搜索任务具有时效不确定性，目标可能会漂移或进入已被扫描过的安全区域。若USV从不重访，系统将失去环境的时效掌控力。",
                    "机制定义：系统为网格 x 维护相对时间戳 T_last(x)。\n  • 当网格处于当前传感器扫描视场内：T_last(x) 归零重置。\n  • 否则每步增加 1：T_last(x, t+1) = T_last(x, t) + 1",
                    "时效权重计算：通过负指数衰减模型将步数映射为 [0,1] 的时效衰减系数，控制遗忘曲线斜率：\n  R_weight(x) = 1.0 - exp( - T_last(x) / τ )\n其中 τ 为时间常数，控制系统周期性复扫的频率偏好。"
                ],
                Inches(0.6), Inches(1.5), Inches(5.8), Inches(3.2),
                C_AMBER_MID, C_AMBER_LIGHT, C_AMBER_BORDER, tag_text="时效衰减机制"
            ),
            
            # Right: Fusion
            draw_premium_card(
                slide,
                "综合信息价值场构建 (Linear Fusion)",
                [
                    "在单艇与双艇路径规划中，我们需要一个统一的指标来定量评估每个候选空间网格的瞬时搜索价值：",
                    "由于 GP Clue(x) 描绘线索连续强度，Intensity(x) 描绘目标离散存在概率。两者数值区间与量纲完全不同，我们首先对两场执行全局归一化(Min-Max)处理，随后执行线性加权融合构建决策底座："
                ],
                Inches(6.93), Inches(1.5), Inches(5.8), Inches(3.2),
                C_BLUE_DEEP, C_WHITE, C_BLUE_BORDER, tag_text="信息价值融合"
            ),
            draw_formula_box(
                slide,
                "search_info_map(x) = w₁ · norm( GP_Clue(x) ) + w₂ · norm( Intensity(x) )",
                Inches(7.09), Inches(3.25), Inches(5.48), Inches(0.65),
                label="▼ 综合搜索信息价值场公式", theme_color=C_BLUE_MID
            ),
            draw_premium_card(
                slide,
                "解耦设计决策：时效性与信念强度的分离",
                [
                    "注意：本工作并未将 Recency 时效合并入统一价值场公式，目的是防止时效权重过度稀释线索与目标存在信念的空间强度结构。时效衰减作为独立的乘积偏置项，作用于候选路径评分中，实现了结构上的科学解耦。"
                ],
                Inches(7.09), Inches(4.0), Inches(5.48), Inches(1.15),
                C_AMBER_MID, C_WHITE, C_AMBER_BORDER, tag_text=None
            ),
            
            # Bottom: Chapter summary
            draw_premium_card(
                slide,
                "本章小结 —— 多源互补的信息底座体系",
                [
                    "本章构建了由 GP Clue（线索外推估算）、Intensity（贝叶斯存在信念）、Recency（时间刷新时效）三维互补的搜索建模底座。\n"
                    "其中 GP Clue 负责寻找蛛丝马迹，Intensity 负责压实搜索价值和防止重复扫描，Recency 独立生效确保全局长期覆盖度。\n"
                    "三者通过科学的层级架构与线性融合，构建出兼顾动态追踪与持续覆扫的搜索状态评估平台，为后续路径规划打下坚实基础。"
                ],
                Inches(0.6), Inches(4.85), Inches(12.13), Inches(1.6),
                C_SLATE_DEEP, C_SLATE_LIGHT, C_SLATE_BORDER, tag_text="Part 2 小结"
            )
        ]
    }
]

# ================================================================
# Main Execution Loop
# ================================================================
for page in SLIDE_DATA:
    idx = page["index"]
    title = page["title"]
    subtitle = page["subtitle"]
    builder = page["builder"]
    
    slide = prs.slides[idx]
    print(f"\n[Processing] Overwriting Slide {idx+1}...")
    
    # 1. Update Title and Subtitle placeholder
    for shape in slide.shapes:
        if shape.is_placeholder:
            ph_idx = shape.placeholder_format.idx
            if ph_idx == 16:
                shape.text = title
                for para in shape.text_frame.paragraphs:
                    para.font.name = F_TITLE
                    para.font.size = Pt(18)
                    para.font.bold = True
                print(f"  -> Title placeholder updated to: '{title}'")
            elif ph_idx == 14:
                shape.text = subtitle
                for para in shape.text_frame.paragraphs:
                    para.font.name = F_TITLE
                    para.font.size = Pt(13)
                    para.font.color.rgb = C_SLATE_MID
                print(f"  -> Subtitle placeholder updated to: '{subtitle}'")
            elif ph_idx in (1, 2):
                shape.text = str(idx + 1)
                print(f"  -> Slide number updated to: '{idx + 1}'")
                
    # 2. Safely delete original body placeholder shapes from index 15 onwards
    # Deleting in reverse order ensures shape index alignment is preserved
    original_shape_count = len(slide.shapes)
    print(f"  -> Original shape count: {original_shape_count}")
    for j in range(original_shape_count - 1, 14, -1):
        shape = slide.shapes[j]
        slide.shapes._spTree.remove(shape._element)
    print(f"  -> Deleted {original_shape_count - len(slide.shapes)} body shapes. Retained {len(slide.shapes)} template shapes.")
    
    # 3. Build premium editable layout
    builder(slide)
    print(f"  -> Successfully generated premium layout. Total shapes: {len(slide.shapes)}")

# Save final presentation
prs.save(PPT_PATH)
print(f"\n[Success] PPT successfully written to: {PPT_PATH}")
print(f"[Done] Cleaned placeholders and structured academic presentation.")
