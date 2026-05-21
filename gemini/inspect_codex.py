import sys
import os
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

codex_path = r"E:\毕设PPT\USV_PPT_codex.pptx"
if not os.path.exists(codex_path):
    print(f"Error: File not found at {codex_path}")
    sys.exit(1)

prs = Presentation(codex_path)
print(f"Total slides in codex: {len(prs.slides)}")

# Let's inspect slide 6, 7, 8, 9 (indices 5, 6, 7, 8)
target_indices = [5, 6, 7, 8]
for idx in target_indices:
    if idx >= len(prs.slides):
        continue
    slide = prs.slides[idx]
    print(f"\n=========================================")
    print(f"Slide Index {idx} (Slide {idx + 1})")
    print(f"=========================================")
    print(f"Total shapes: {len(slide.shapes)}")
    
    # List all shapes and their brief details
    for i, shape in enumerate(slide.shapes):
        print(f"\nShape {i}: '{shape.name}', type={shape.shape_type}")
        print(f"  Pos: L={shape.left}, T={shape.top}, W={shape.width}, H={shape.height}")
        if hasattr(shape, 'text') and shape.text.strip():
            print(f"  Text: {repr(shape.text.strip()[:100])}")
        
        # Check fill
        try:
            if shape.fill and shape.fill.type:
                fc = shape.fill.fore_color
                if fc.type == 1:
                    print(f"  Fill Color: #{fc.rgb}")
                elif fc.type == 2:
                    print(f"  Fill Theme: idx={fc.theme_color}")
        except Exception:
            pass
            
        # Check border/line
        try:
            if shape.line:
                lc = shape.line.color
                if lc.type == 1:
                    print(f"  Line Color: #{lc.rgb}")
                elif lc.type == 2:
                    print(f"  Line Theme: idx={lc.theme_color}")
                w = shape.line.width
                if w:
                    print(f"  Line width: {w/12700:.2f}pt")
        except Exception:
            pass
            
        # Check first paragraph font info if any
        try:
            if hasattr(shape, 'text_frame') and shape.text_frame.paragraphs:
                p = shape.text_frame.paragraphs[0]
                if p.runs:
                    run = p.runs[0]
                    print(f"  Font: name={run.font.name}, size={run.font.size}, bold={run.font.bold}")
                    if run.font.color and run.font.color.type == 1:
                        print(f"  Font Color: #{run.font.color.rgb}")
        except Exception:
            pass
