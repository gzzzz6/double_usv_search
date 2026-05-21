import os
import copy
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def build_slide_5():
    src_file = r"E:\毕设PPT\USV_PPT.pptx"
    backup_file = r"E:\毕设PPT\USV_PPT.pptx.bak"
    
    # Back up the file just in case
    shutil.copy2(src_file, backup_file)
    print(f"Created backup of presentation to {backup_file}")
    
    prs = Presentation(src_file)
    slide_5 = prs.slides[4] # index 4
    slide_7 = prs.slides[6] # index 6
    
    # 1. Delete all existing shapes in Slide 5
    shapes_to_delete = list(slide_5.shapes)
    for s in shapes_to_delete:
        el = s.element
        el.getparent().remove(el)
        
    print("Cleared existing Slide 5 shapes.")
    
    # 2. Copy the slide template/header shapes 0 to 12 from Slide 7 to Slide 5
    for i in range(13):
        if i < len(slide_7.shapes):
            shape = slide_7.shapes[i]
            el = shape.element
            new_el = copy.deepcopy(el)
            slide_5.shapes._spTree.append(new_el)
            
    print("Copied slide headers and navigation template from Slide 7.")
    
    # 3. Update the page number text box (typically Shape 11) to '5'
    if len(slide_5.shapes) > 11:
        s = slide_5.shapes[11]
        if hasattr(s, 'text') and s.text.strip() == "7":
            s.text = "5"
            # Ensure proper font formatting
            tf = s.text_frame
            tf.paragraphs[0].font.size = Pt(10)
            tf.paragraphs[0].font.name = "Arial"
            print("Updated slide number to 5.")
            
    # 4. Add the Slide Subtitle (Title on the content area)
    # Shape 13 on Slide 7 was the subtitle. Let's add it manually on Slide 5.
    sub_title_box = slide_5.shapes.add_textbox(Inches(0.63), Inches(1.2), Inches(11.5), Inches(0.4))
    tf = sub_title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "总体技术路线：系统分层在线滚动决策与安全执行流程"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = RGBColor(30, 80, 160)
    print("Added slide content title.")

    # Helper function to style a rounded rectangle card
    def add_card(slide, title, subtitle, left, top, width, height, bg_color, border_color, title_color=RGBColor(0,0,0), sub_color=RGBColor(100,100,100), is_bold=True, border_width_pt=1):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width_pt)
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.08)
        tf.margin_right = Inches(0.08)
        tf.margin_top = Inches(0.05)
        tf.margin_bottom = Inches(0.05)
        
        # Clear default paragraph
        p_title = tf.paragraphs[0]
        p_title.alignment = PP_ALIGN.CENTER
        p_title.text = title
        p_title.font.name = "Microsoft YaHei"
        p_title.font.size = Pt(9.5)
        p_title.font.bold = is_bold
        p_title.font.color.rgb = title_color
        
        if subtitle:
            p_sub = tf.add_paragraph()
            p_sub.alignment = PP_ALIGN.CENTER
            p_sub.text = subtitle
            p_sub.font.name = "Microsoft YaHei"
            p_sub.font.size = Pt(7.5)
            p_sub.font.color.rgb = sub_color
            
        return card

    # Helper function to add arrows
    def add_arrow(slide, left, top, width, height, color=RGBColor(160, 160, 160)):
        arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = color
        arrow.line.fill.background() # No border
        return arrow

    def add_down_arrow(slide, left, top, width, height, color=RGBColor(160, 160, 160)):
        arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, left, top, width, height)
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = color
        arrow.line.fill.background()
        return arrow

    # ==================== LAYER 1: 感知与信息场构建 ====================
    # Background container
    ly1_bg = slide_5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.65), Inches(12.33), Inches(1.4))
    ly1_bg.fill.solid()
    ly1_bg.fill.fore_color.rgb = RGBColor(240, 246, 255) # Light Alice Blue
    ly1_bg.line.color.rgb = RGBColor(200, 222, 255)
    ly1_bg.line.width = Pt(1.5)
    
    # Layer Title
    ly1_title = slide_5.shapes.add_textbox(Inches(0.6), Inches(1.7), Inches(4.0), Inches(0.3))
    ly1_title.text_frame.paragraphs[0].text = "第一层：感知与信息场构建 (Perception & Belief Modeling)"
    ly1_title.text_frame.paragraphs[0].font.name = "Microsoft YaHei"
    ly1_title.text_frame.paragraphs[0].font.size = Pt(10.5)
    ly1_title.text_frame.paragraphs[0].font.bold = True
    ly1_title.text_frame.paragraphs[0].font.color.rgb = RGBColor(20, 80, 160)

    # Input Cards
    # Card 1: known_map
    add_card(slide_5, "已知静态地图 (known_map)", "障碍空间先验分布结构\n(nav_map_prior)", 
             Inches(0.8), Inches(2.1), Inches(2.4), Inches(0.75), 
             RGBColor(255,255,255), RGBColor(220,220,220))
    # Card 2: observation
    add_card(slide_5, "局部带噪观测 (observation)", "实时传感器视场内 clue 采样\n(detect_targets & sample)", 
             Inches(3.4), Inches(2.1), Inches(2.4), Inches(0.75), 
             RGBColor(255,255,255), RGBColor(220,220,220))
    # Card 3: detect feedback
    add_card(slide_5, "目标探测反馈 (detect)", "命中贝叶斯更新 / 未命中衰减\n(miss_update_intensity)", 
             Inches(6.0), Inches(2.1), Inches(2.4), Inches(0.75), 
             RGBColor(255,255,255), RGBColor(220,220,220))

    # Core Value Field Card on the right
    # Small component labels above
    add_card(slide_5, "GP Clue", "后验均值/方差", Inches(8.8), Inches(1.75), Inches(1.05), Inches(0.4), RGBColor(255,255,255), RGBColor(200,200,200), is_bold=False)
    add_card(slide_5, "Intensity", "目标分布信念", Inches(9.95), Inches(1.75), Inches(1.05), Inches(0.4), RGBColor(255,255,255), RGBColor(200,200,200), is_bold=False)
    add_card(slide_5, "Recency", "时效衰减因子", Inches(11.1), Inches(1.75), Inches(1.05), Inches(0.4), RGBColor(255,255,255), RGBColor(200,200,200), is_bold=False)

    # Big search_info_map card
    add_card(slide_5, "综合信息价值场 search_info_map", 
             "= 0.5 × norm(GP Clue) + 0.5 × norm(Intensity)", 
             Inches(8.8), Inches(2.25), Inches(3.35), Inches(0.65), 
             RGBColor(215, 235, 255), RGBColor(140, 185, 255), 
             title_color=RGBColor(0, 50, 150), border_width_pt=1.5)
             
    # Connecting arrow from input to belief
    add_arrow(slide_5, Inches(8.45), Inches(2.32), Inches(0.3), Inches(0.12))


    # ==================== LAYER 2: 路径规划与决策 ====================
    # Background container
    ly2_bg = slide_5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(3.2), Inches(12.33), Inches(1.6))
    ly2_bg.fill.solid()
    ly2_bg.fill.fore_color.rgb = RGBColor(240, 252, 245) # Light Mint
    ly2_bg.line.color.rgb = RGBColor(200, 240, 215)
    ly2_bg.line.width = Pt(1.5)

    # Layer Title
    ly2_title = slide_5.shapes.add_textbox(Inches(0.6), Inches(3.25), Inches(4.5), Inches(0.3))
    ly2_title.text_frame.paragraphs[0].text = "第二层：单艇路径规划与决策 (Single-USV Planning & Decision)"
    ly2_title.text_frame.paragraphs[0].font.name = "Microsoft YaHei"
    ly2_title.text_frame.paragraphs[0].font.size = Pt(10.5)
    ly2_title.text_frame.paragraphs[0].font.bold = True
    ly2_title.text_frame.paragraphs[0].font.color.rgb = RGBColor(10, 110, 60)

    # Single USV Planning Sequence
    # Card 1: Anchor
    add_card(slide_5, "Anchor 种子点提取", "提取 top_k=6 个\n高搜索价值自由格点", 
             Inches(0.8), Inches(3.68), Inches(1.75), Inches(0.7), 
             RGBColor(255,255,255), RGBColor(210,210,210))
    # Card 2: Viewpoint
    add_card(slide_5, "viewpoint 邻域采样", "各 Anchor 环形采 6 候选点\n(viewpoints_per_anchor)", 
             Inches(2.8), Inches(3.68), Inches(1.75), Inches(0.7), 
             RGBColor(255,255,255), RGBColor(210,210,210))
    # Card 3: A* pathfinding
    add_card(slide_5, "A* 安全寻路", "软膨胀避障代价计算\n(soft_clearance_astar)", 
             Inches(4.8), Inches(3.68), Inches(1.75), Inches(0.7), 
             RGBColor(255,255,255), RGBColor(210,210,210))
    # Card 4: Cut segment
    add_card(slide_5, "路径段截取 segment_path", "仅截取前 horizon=8 步\n作为滚动执行备选段", 
             Inches(6.8), Inches(3.68), Inches(1.75), Inches(0.7), 
             RGBColor(255,255,255), RGBColor(210,210,210))

    # Connectors for planning chain
    add_arrow(slide_5, Inches(2.6), Inches(3.92), Inches(0.16), Inches(0.1))
    add_arrow(slide_5, Inches(4.60), Inches(3.92), Inches(0.16), Inches(0.1))
    add_arrow(slide_5, Inches(6.60), Inches(3.92), Inches(0.16), Inches(0.1))
    add_arrow(slide_5, Inches(8.60), Inches(3.92), Inches(0.16), Inches(0.1))

    # Core scoring box on the right
    score_card = slide_5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(3.35), Inches(3.35), Inches(1.3))
    score_card.fill.solid()
    score_card.fill.fore_color.rgb = RGBColor(225, 248, 232)
    score_card.line.color.rgb = RGBColor(160, 220, 180)
    score_card.line.width = Pt(1.5)
    
    tf = score_card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)
    
    p = tf.paragraphs[0]
    p.text = "候选路径段综合评分决策 (Path Evaluation)"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = RGBColor(10, 80, 40)
    
    p_body = tf.add_paragraph()
    p_body.text = "• 信息增益 (search_info_gain) × 权重比 6.5\n• 刷新增益 (recency_bias) × 权重比 2.0\n• 移动代价 (execution_cost) × 权重比 -0.35\n• 动力约束 (U-turn penalty) 负向转向惩罚"
    p_body.font.name = "Microsoft YaHei"
    p_body.font.size = Pt(7.5)
    p_body.font.color.rgb = RGBColor(50, 50, 50)


    # ==================== LAYER 3: 双艇协同与安全执行 ====================
    # Background container
    ly3_bg = slide_5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(4.95), Inches(12.33), Inches(1.8))
    ly3_bg.fill.solid()
    ly3_bg.fill.fore_color.rgb = RGBColor(255, 248, 242) # Light Apricot
    ly3_bg.line.color.rgb = RGBColor(255, 225, 200)
    ly3_bg.line.width = Pt(1.5)

    # Layer Title
    ly3_title = slide_5.shapes.add_textbox(Inches(0.6), Inches(5.0), Inches(5.0), Inches(0.3))
    ly3_title.text_frame.paragraphs[0].text = "第三层：双艇协同分配与安全执行层 (Coordinated & Safe Execution)"
    ly3_title.text_frame.paragraphs[0].font.name = "Microsoft YaHei"
    ly3_title.text_frame.paragraphs[0].font.size = Pt(10.5)
    ly3_title.text_frame.paragraphs[0].font.bold = True
    ly3_title.text_frame.paragraphs[0].font.color.rgb = RGBColor(140, 60, 10)

    # Coordinated Section
    # Coordinated main entry card
    add_card(slide_5, "Coordinated 序列化分配", "遍历规划顺序 (0,1) 与 (1,0)，取最优", 
             Inches(0.8), Inches(5.4), Inches(2.2), Inches(0.5), 
             RGBColor(255,255,255), RGBColor(230,210,190))
             
    # Branch 1: USV 0 Lead
    add_card(slide_5, "领艇 USV 0 独立决策", "基于全局价值场 search_info_map", 
             Inches(0.8), Inches(6.05), Inches(2.2), Inches(0.5), 
             RGBColor(255,255,255), RGBColor(220,200,180))
             
    # Branch 2: USV 1 Follower (with residual_map)
    add_card(slide_5, "艇 USV 1 协同规划", "基于扣除视场重叠后的 residual_map\n(带 overlap_penalty 重叠惩罚项)", 
             Inches(3.3), Inches(6.05), Inches(2.6), Inches(0.5), 
             RGBColor(255,255,255), RGBColor(220,200,180))
             
    # Small sequence arrow between USV 0 and USV 1
    add_arrow(slide_5, Inches(3.05), Inches(6.2), Inches(0.2), Inches(0.08))

    # Responsibility Regularizer annotation
    add_card(slide_5, "Soft Responsibility Bias", 
             "本地评分类偏置偏好：\nplanner_map × (1 + 0.25 × score)\n只引导搜索，不强行割裂区域，允许跨区", 
             Inches(6.1), Inches(5.4), Inches(2.5), Inches(1.15), 
             RGBColor(255,255,255), RGBColor(230,200,180), 
             title_color=RGBColor(120,60,0), border_width_pt=1, is_bold=True)

    # Right side: reservation and execution
    # Card 1: reservation
    add_card(slide_5, "reservation_v1 时空冲突校验", "检测格冲突/换位冲突，施加软halo时空代价", 
             Inches(8.8), Inches(5.35), Inches(3.35), Inches(0.55), 
             RGBColor(255, 235, 220), RGBColor(255, 190, 150), 
             title_color=RGBColor(140, 50, 0), border_width_pt=1.5)
             
    # Card 2: execution & loop back
    add_card(slide_5, "双艇动作滚动执行 (Replanning)", "执行当前 step 路径点，获取新局部观测并更新", 
             Inches(8.8), Inches(6.15), Inches(3.35), Inches(0.5), 
             RGBColor(255, 235, 220), RGBColor(255, 190, 150), 
             title_color=RGBColor(140, 50, 0), border_width_pt=1.5)

    # Arrow to reservation
    add_arrow(slide_5, Inches(8.65), Inches(5.72), Inches(0.12), Inches(0.08))
    # Down arrow from reservation to execution
    add_down_arrow(slide_5, Inches(10.45), Inches(5.95), Inches(0.08), Inches(0.15))

    # Notes under Layer 3
    notes = slide_5.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(5.2), Inches(0.2))
    notes.text_frame.margin_top = 0
    notes.text_frame.margin_bottom = 0
    p = notes.text_frame.paragraphs[0]
    p.text = "* 注：避障规划 (A* soft clearance) 与协同避撞 (reservation_v1) 均在安全执行层实现，不污染上层信念状态。"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(7.2)
    p.font.color.rgb = RGBColor(120, 120, 120)

    # Save to the actual target file
    prs.save(src_file)
    print(f"Successfully modified Slide 5 and saved to {src_file}!")

if __name__ == "__main__":
    build_slide_5()
