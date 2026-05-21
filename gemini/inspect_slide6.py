from pptx import Presentation
prs = Presentation(r"E:\毕设PPT\USV_PPT.pptx")
slide = prs.slides[5] # slide 6 is index 5
print("=== Slide 6 Shapes ===")
for j, shape in enumerate(slide.shapes):
    print(f"Shape {j}: type={shape.shape_type}, name={shape.name}")
    if hasattr(shape, 'text') and shape.text.strip():
        print(f"  Text: {shape.text.strip()}")
