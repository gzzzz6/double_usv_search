import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def build_native_slide_5():
    src_file = r"E:\毕设PPT\USV_PPT.pptx"
    backup_file = r"E:\毕设PPT\USV_PPT_native_bak.pptx"
    
    # Back up the file just in case
    shutil.copy2(src_file, backup_file)
    print(f"Created backup of presentation to {backup_file}")
    
    prs = Presentation(src_file)
    slide_5 = prs.slides[4] # index 4
    
    # Delete all existing shapes in Slide 5
    shapes_to_delete = list(slide_5.shapes)
    for s in shapes_to_delete:
        el = s.element
        el.getparent().remove(el)
    print("Cleared existing Slide 5 shapes.")
    
    # Define Palette (Premium Dark Tech Theme)
    c_bg = RGBColor(15, 32, 68)           # Deep Tech Blue background
    c_panel_bg = RGBColor(22, 42, 85)     # Lighter dark blue for panels
    c_card_bg = RGBColor(32, 58, 108)     # Card fill
    c_card_border = RGBColor(56, 88, 148)  # Card border
    
    c_white = RGBColor(255, 255, 255)
    c_light_blue = RGBColor(180, 210, 255)
    c_green = RGBColor(110, 231, 183)
    c_orange = RGBColor(251, 191, 36)
    
    # Panel Borders
    b_p1 = RGBColor(100, 160, 255)
    b_p2 = RGBColor(100, 200, 160)
    b_p3 = RGBColor(180, 180, 255)
    b_p4 = RGBColor(255, 160, 80)
    
    # Draw Background Canvas
    bg_canvas = slide_5.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.33), Inches(7.5))
    bg_canvas.fill.solid()
    bg_canvas.fill.fore_color.rgb = c_bg
    bg_canvas.line.fill.background()
    print("Added deep blue canvas background.")
    
    # Helper to style headers
    def add_header(title, subtitle):
        tb = slide_5.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.33), Inches(0.95))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_top = 0
        tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = title
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = c_white
        
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.text = subtitle
        p2.font.name = "Microsoft YaHei"
        p2.font.size = Pt(13)
        p2.font.color.rgb = c_light_blue
        
    add_header("总体技术路线：信息场驱动的滚动搜索决策", "known static map  +  local observation  +  segment replanning")
    
    # Helper to add standard panels
    def add_panel(title, num, left, top, width, height, border_color, num_bg):
        panel = slide_5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        panel.fill.solid()
        panel.fill.fore_color.rgb = c_panel_bg
        panel.line.color.rgb = border_color
        panel.line.width = Pt(1.5)
        
        # Add badge circle
        circle = slide_5.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.12), top + Inches(0.12), Inches(0.25), Inches(0.25))
        circle.fill.solid()
        circle.fill.fore_color.rgb = num_bg
        circle.line.fill.background()
        tf_c = circle.text_frame
        tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0
        p_c = tf_c.paragraphs[0]
        p_c.alignment = PP_ALIGN.CENTER
        p_c.text = str(num)
        p_c.font.name = "Microsoft YaHei"
        p_c.font.size = Pt(10)
        p_c.font.bold = True
        p_c.font.color.rgb = c_white
        
        # Add Title text next to it
        tb = slide_5.shapes.add_textbox(left + Inches(0.42), top + Inches(0.08), width - Inches(0.55), Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = c_white
        
        return panel

    # Helper to add inner editable cards
    def add_card(title, desc, left, top, width, height, custom_bg=None, custom_border=None, font_size_title=9, font_size_desc=7.5, custom_title_color=c_white):
        card = slide_5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = custom_bg if custom_bg else c_card_bg
        card.line.color.rgb = custom_border if custom_border else c_card_border
        card.line.width = Pt(1)
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.08)
        tf.margin_right = Inches(0.08)
        tf.margin_top = Inches(0.05)
        tf.margin_bottom = Inches(0.05)
        
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.text = title
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(font_size_title)
        p.font.bold = True
        p.font.color.rgb = custom_title_color
        
        if desc:
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.LEFT
            p2.text = desc
            p2.font.name = "Microsoft YaHei"
            p2.font.size = Pt(font_size_desc)
            p2.font.color.rgb = c_light_blue
            
        return card

    # Helper for simple thin arrow markers
    def add_right_arrow(left, top, width=Inches(0.18), height=Inches(0.1)):
        arrow = slide_5.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = RGBColor(100, 150, 255)
        arrow.line.fill.background()

    def add_down_arrow(left, top, width=Inches(0.1), height=Inches(0.18)):
        arrow = slide_5.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, left, top, width, height)
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = RGBColor(100, 150, 255)
        arrow.line.fill.background()

    # Draw the 4 main Panels
    add_panel("感知与状态更新", 1, Inches(0.5), Inches(1.3), Inches(2.6), Inches(4.5), b_p1, RGBColor(59, 130, 246))
    add_panel("信息价值场构建", 2, Inches(3.4), Inches(1.3), Inches(2.6), Inches(4.5), b_p2, RGBColor(16, 185, 129))
    add_panel("单艇路径规划与评分", 3, Inches(6.3), Inches(1.3), Inches(3.5), Inches(4.5), b_p3, RGBColor(139, 92, 246))
    add_panel("双艇协同与安全执行", 4, Inches(10.1), Inches(1.3), Inches(2.73), Inches(4.5), b_p4, RGBColor(245, 158, 11))
    
    # Draw horizontal big arrows between panels
    add_right_arrow(Inches(3.16), Inches(3.2))
    add_right_arrow(Inches(6.06), Inches(3.2))
    add_right_arrow(Inches(9.86), Inches(3.2))
    
    # ================= PANEL 1 CARDS =================
    add_card("known_map 已知静态地图", "地图障碍空间先验分布结构\n(nav_map_prior)", Inches(0.65), Inches(1.8), Inches(2.3), Inches(0.65))
    add_card("observation 局部带噪观测", "局部传感器视场内 clue 采样\n(detect_targets & sample)", Inches(0.65), Inches(2.5), Inches(2.3), Inches(0.65))
    add_card("detect 目标探测反馈", "命中贝叶斯更新 / 未命中衰减\n(miss_update_intensity)", Inches(0.65), Inches(3.2), Inches(2.3), Inches(0.65))
    
    # arrow & label
    add_down_arrow(Inches(1.75), Inches(3.9))
    tb_out = slide_5.shapes.add_textbox(Inches(0.65), Inches(4.02), Inches(2.3), Inches(0.2))
    tb_out.text_frame.margin_top = tb_out.text_frame.margin_bottom = 0
    tb_out.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    tb_out.text_frame.paragraphs[0].text = "输出三状态"
    tb_out.text_frame.paragraphs[0].font.name = "Microsoft YaHei"
    tb_out.text_frame.paragraphs[0].font.size = Pt(8.5)
    tb_out.text_frame.paragraphs[0].font.color.rgb = c_light_blue
    
    # triplet small cards
    add_card("GP Clue", "后验均值方差\n形成 UCB", Inches(0.65), Inches(4.25), Inches(0.72), Inches(1.4), font_size_title=8.5, font_size_desc=7)
    add_card("Intensity", "目标分布信念\n探测反馈更新", Inches(1.44), Inches(4.25), Inches(0.72), Inches(1.4), font_size_title=8.5, font_size_desc=7)
    add_card("Recency", "时效衰减因子\nτ=12 遗忘机制", Inches(2.23), Inches(4.25), Inches(0.72), Inches(1.4), font_size_title=8.5, font_size_desc=7)
    
    # ================= PANEL 2 CARDS =================
    # Hero Card (Green)
    c_hero_bg = RGBColor(6, 78, 59)
    c_hero_border = RGBColor(52, 211, 153)
    add_card("search_info_map 综合信息价值场", 
             "search_info_map =\n0.5 × norm(GP Clue)\n+ 0.5 × norm(Intensity)", 
             Inches(3.55), Inches(1.8), Inches(2.3), Inches(1.3), 
             custom_bg=c_hero_bg, custom_border=c_hero_border, font_size_title=10, custom_title_color=c_green)
             
    # Note card
    c_note_bg = RGBColor(30, 30, 20)
    c_note_border = RGBColor(251, 191, 36)
    add_card("⚠️ Recency 时效项说明", "观测时效因子不直接代入 search_info_map，而是单独进入路径段评分中的 recency_bias 乘积项。", 
             Inches(3.55), Inches(3.25), Inches(2.3), Inches(1.1), 
             custom_bg=c_note_bg, custom_border=c_note_border, font_size_title=8.5, font_size_desc=7, custom_title_color=c_orange)
             
    # Bottom GP limitation card
    add_card("gp_max_points = 400", "高斯过程拟合上限；通过设置每 gp_fit_every = 5 步重新拟合，控制单步计算资源开销。", 
             Inches(3.55), Inches(4.5), Inches(2.3), Inches(1.15), font_size_title=9, font_size_desc=7.5)

    # ================= PANEL 3 CARDS =================
    # Flow elements (4 cards + 3 small arrows)
    w_flow = Inches(0.7)
    h_flow = Inches(0.9)
    add_card("Anchor 提取", "提取 top_k=6\n高价值自由点", Inches(6.45), Inches(1.8), w_flow, h_flow, font_size_title=8.5, font_size_desc=7)
    add_right_arrow(Inches(7.18), Inches(2.2), width=Inches(0.12), height=Inches(0.08))
    
    add_card("viewpoint 采样", "simple_ring_v1\n各种子点采6点", Inches(7.33), Inches(1.8), w_flow, h_flow, font_size_title=8.5, font_size_desc=7)
    add_right_arrow(Inches(8.06), Inches(2.2), width=Inches(0.12), height=Inches(0.08))
    
    add_card("A* 寻路", "soft_clearance\n_astar_v1 避障", Inches(8.21), Inches(1.8), w_flow, h_flow, font_size_title=8.5, font_size_desc=7)
    add_right_arrow(Inches(8.94), Inches(2.2), width=Inches(0.12), height=Inches(0.08))
    
    add_card("截取段", "horizon = 8\n滚动截取前8步", Inches(9.09), Inches(1.8), w_flow, h_flow, font_size_title=8.5, font_size_desc=7)

    # Scoring Panel card
    score_p = slide_5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.45), Inches(2.85), Inches(3.2), Inches(2.8))
    score_p.fill.solid()
    score_p.fill.fore_color.rgb = RGBColor(26, 40, 75)
    score_p.line.color.rgb = c_card_border
    score_p.line.width = Pt(1.2)
    
    tf_s = score_p.text_frame
    tf_s.margin_left = tf_s.margin_right = Inches(0.1)
    tf_s.margin_top = Inches(0.08)
    p_st = tf_s.paragraphs[0]
    p_st.alignment = PP_ALIGN.CENTER
    p_st.text = "🏆 路径段综合评分机制"
    p_st.font.name = "Microsoft YaHei"
    p_st.font.size = Pt(10)
    p_st.font.bold = True
    p_st.font.color.rgb = c_white
    
    # We will draw 3 rows inside manually using nested cards for rows to look absolutely brilliant and editable
    def add_score_row(title, coef, desc, top, bg_color, border_color, badge_color):
        row = slide_5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.55), top, Inches(2.15), Inches(0.55))
        row.fill.solid()
        row.fill.fore_color.rgb = bg_color
        row.line.color.rgb = border_color
        row.line.width = Pt(1)
        
        tf = row.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = Inches(0.04)
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = c_white
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Microsoft YaHei"
        p2.font.size = Pt(7)
        p2.font.color.rgb = c_light_blue
        
        # Badge card
        badge = slide_5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.75), top + Inches(0.08), Inches(0.8), Inches(0.38))
        badge.fill.solid()
        badge.fill.fore_color.rgb = badge_color
        badge.line.fill.background()
        tf_b = badge.text_frame
        tf_b.margin_top = tf_b.margin_bottom = 0
        p_b = tf_b.paragraphs[0]
        p_b.alignment = PP_ALIGN.CENTER
        p_b.text = coef
        p_b.font.name = "Courier New"
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = c_white
        
    add_score_row("search_info_gain", "× 6.5", "信息收益：沿线累加综合信息价值场", Inches(3.25), RGBColor(30, 45, 95), RGBColor(50, 100, 200), RGBColor(29, 78, 216))
    add_score_row("recency_bias", "× 2.0", "刷新收益：时效场衰减因子累积补偿", Inches(3.88), RGBColor(20, 50, 80), RGBColor(40, 160, 100), RGBColor(4, 120, 87))
    add_score_row("exec_cost + U-turn", "×(-0.35)", "执行代价：长度阻尼以及航向掉头惩罚", Inches(4.51), RGBColor(45, 30, 60), RGBColor(160, 60, 60), RGBColor(185, 28, 28))

    # ================= PANEL 4 CARDS =================
    add_card("coordinated 分配顺序 (0,1) / (1,0)", "双顺序遍历枚举，计算最大联合总分", Inches(10.2), Inches(1.8), Inches(2.53), Inches(0.55))
    add_card("residual_map 后规划扣减", "后规划艇扣除先规划艇视场已覆盖价值", Inches(10.2), Inches(2.42), Inches(2.53), Inches(0.55))
    add_card("overlap / same-viewpoint", "强视场重叠惩罚，同视点给予额外抵消", Inches(10.2), Inches(3.04), Inches(2.53), Inches(0.55))
    add_card("soft responsibility ranking", "仅作评分类排序偏置，不阻断跨区跨界", Inches(10.2), Inches(3.66), Inches(2.53), Inches(0.55))
    
    # reservation gets a yellow highlight border/bg
    add_card("reservation_v1 冲突校验 (核心)", "同格碰撞校验 / 对向换位冲突 / 时空软 halo 代价规避", Inches(10.2), Inches(4.28), Inches(2.53), Inches(0.65),
             custom_bg=RGBColor(48, 44, 40), custom_border=RGBColor(245, 158, 11), font_size_title=9.5, font_size_desc=7, custom_title_color=c_orange)
             
    add_card("双艇动作按 segment 执行", "双艇执行所截取的 horizon = 8 路径点", Inches(10.2), Inches(5.0), Inches(2.53), Inches(0.55))


    # ================= BOTTOM REPLANNING LOOP =================
    # Background panel for the loop
    loop_panel = slide_5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(5.95), Inches(12.33), Inches(0.85))
    loop_panel.fill.solid()
    loop_panel.fill.fore_color.rgb = RGBColor(20, 36, 76)
    loop_panel.line.color.rgb = RGBColor(40, 70, 120)
    loop_panel.line.width = Pt(1)
    
    # Replanning label
    tb_rl = slide_5.shapes.add_textbox(Inches(0.6), Inches(6.05), Inches(1.5), Inches(0.6))
    tf_rl = tb_rl.text_frame
    tf_rl.word_wrap = True
    p_rl = tf_rl.paragraphs[0]
    p_rl.text = "🔄 在线滚动重规划\nReplanning Loop"
    p_rl.font.name = "Microsoft YaHei"
    p_rl.font.size = Pt(10)
    p_rl.font.bold = True
    p_rl.font.color.rgb = c_orange
    
    # 4 Steps of the Loop
    def add_loop_step(title, desc, left):
        card = slide_5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(6.02), Inches(2.2), Inches(0.7))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card_bg
        card.line.color.rgb = c_card_border
        card.line.width = Pt(1)
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.04)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = title
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = c_white
        
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.text = desc
        p2.font.name = "Microsoft YaHei"
        p2.font.size = Pt(7.5)
        p2.font.color.rgb = c_light_blue
        
    add_loop_step("1. 执行短时域路径段", "执行当前 step 的移动段", Inches(2.3))
    add_right_arrow(Inches(4.55), Inches(6.32), width=Inches(0.15), height=Inches(0.1))
    
    add_loop_step("2. 获取局部新观测", "获取传感器覆盖内Clue信息", Inches(4.75))
    add_right_arrow(Inches(7.0), Inches(6.32), width=Inches(0.15), height=Inches(0.1))
    
    add_loop_step("3. 动态更新多源信息场", "拟合GP并进行信念贝叶斯更新", Inches(7.2))
    add_right_arrow(Inches(9.45), Inches(6.32), width=Inches(0.15), height=Inches(0.1))
    
    add_loop_step("4. 全局联合重新决策", "重新进行最优路径段分配规划", Inches(9.65))
    
    # Loop back arrow indicator
    back_arrow = slide_5.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, Inches(11.9), Inches(6.25), Inches(0.7), Inches(0.25))
    back_arrow.fill.solid()
    back_arrow.fill.fore_color.rgb = RGBColor(100, 160, 255)
    back_arrow.line.fill.background()
    tf_ba = back_arrow.text_frame
    tf_ba.margin_top = tf_ba.margin_bottom = 0
    p_ba = tf_ba.paragraphs[0]
    p_ba.alignment = PP_ALIGN.CENTER
    p_ba.text = "循环"
    p_ba.font.name = "Microsoft YaHei"
    p_ba.font.size = Pt(8.5)
    p_ba.font.bold = True
    p_ba.font.color.rgb = c_white

    # ================= FOOTER NOTE =================
    tb_foot = slide_5.shapes.add_textbox(Inches(0.5), Inches(6.92), Inches(12.33), Inches(0.4))
    tf_f = tb_foot.text_frame
    tf_f.word_wrap = True
    p_f = tf_f.paragraphs[0]
    p_f.text = "🛡️ 注：A* soft clearance 避障寻路与 reservation_v1 路径时空冲突规避属于安全执行层，只负责运动路径纠偏，不污染上层的状态估计（intensity / GP clue 均不受影响）。"
    p_f.font.name = "Microsoft YaHei"
    p_f.font.size = Pt(9)
    p_f.font.color.rgb = c_light_blue

    # Save to the actual target file
    prs.save(src_file)
    print(f"Successfully modified Slide 5 to native flat-vector PPT edit format and saved to {src_file}!")

if __name__ == "__main__":
    build_native_slide_5()
