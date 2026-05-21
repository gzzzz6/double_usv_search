from pptx import Presentation
prs = Presentation(r"E:\毕设PPT\USV_PPT.pptx")
print("=== Slide 1 Shapes ===")
for j, shape in enumerate(prs.slides[0].shapes):
    print(f"Shape {j}: type={shape.shape_type}, name={shape.name}")
    if hasattr(shape, 'text') and shape.text.strip():
        print(f"  Text: {shape.text.strip()}")
