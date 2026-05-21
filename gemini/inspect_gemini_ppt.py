from pptx import Presentation
prs = Presentation(r"E:\毕设PPT\USV_PPT_gemini.pptx")
with open(r"F:\pythonprojects\gemini\gemini_ppt_texts.txt", "w", encoding="utf-8") as f:
    f.write(f"Total slides: {len(prs.slides)}\n\n")
    for i in range(min(8, len(prs.slides))):
        f.write(f"=== Slide {i+1} ===\n")
        slide = prs.slides[i]
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text.strip():
                f.write(f"[{shape.name}] {shape.text.strip()}\n")
        f.write("\n")
print("Done")
