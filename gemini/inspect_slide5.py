from pptx import Presentation
prs = Presentation(r"E:\毕设PPT\USV_PPT.pptx")
slide = prs.slides[4] # 5th slide is index 4
print(f"=== Slide 5 Shape Details ===")
for j, shape in enumerate(slide.shapes):
    print(f"Shape {j}: type={shape.shape_type}, name={shape.name}, has_text={hasattr(shape, 'text')}")
    if hasattr(shape, 'text') and shape.text.strip():
        print(f"  Text: {shape.text.strip()}")
