from pptx import Presentation
prs = Presentation(r"E:\毕设PPT\USV_PPT_gemini.pptx")
slide = prs.slides[5] # Index 5 (Slide 6 before we ran build_part2_slides.py)
print("=== Slide 6 (Index 5) Shapes ===")
for i, shape in enumerate(slide.shapes):
    print(f"Shape {i}: '{shape.name}', type={shape.shape_type}")
    if hasattr(shape, 'text') and shape.text.strip():
        print(f"  Text: {shape.text.strip()}")
