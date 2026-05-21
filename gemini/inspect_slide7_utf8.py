import sys
from pptx import Presentation

# set output encoding to utf-8 just in case
sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation(r"E:\毕设PPT\USV_PPT_gemini.pptx")
slide = prs.slides[6] # Slide 7 is index 6
print(f"=== Slide 7 Shapes (Total: {len(slide.shapes)}) ===")
for i, shape in enumerate(slide.shapes):
    txt = shape.text.strip() if hasattr(shape, 'text') else ""
    print(f"Shape {i}: '{shape.name}', type={shape.shape_type}, text='{txt}'")
