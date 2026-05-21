import copy
from pptx import Presentation

try:
    prs = Presentation(r"E:\毕设PPT\USV_PPT.pptx")
    slide_5 = prs.slides[4] # Index 4 is slide 5
    slide_7 = prs.slides[6] # Index 6 is slide 7
    
    print("Initial slide 5 shapes:")
    for i, s in enumerate(slide_5.shapes):
        print(f"Shape {i}: type={s.shape_type}, name={s.name}")
        
    # Delete all shapes in slide 5
    # To delete, we iterate in reverse or collect elements to delete
    shapes_to_delete = list(slide_5.shapes)
    for s in shapes_to_delete:
        el = s.element
        el.getparent().remove(el)
        
    print("Slide 5 shapes after deletion:")
    print(len(slide_5.shapes))
    
    # Copy shapes 0 to 12 from Slide 7 to Slide 5
    for i in range(13):
        if i < len(slide_7.shapes):
            shape = slide_7.shapes[i]
            el = shape.element
            new_el = copy.deepcopy(el)
            slide_5.shapes._spTree.append(new_el)
            
    print("Slide 5 shapes after copying:")
    for i, s in enumerate(slide_5.shapes):
        print(f"Shape {i}: type={s.shape_type}, name={s.name}")
        if hasattr(s, 'text') and s.text.strip():
            print(f"  Text: {s.text.strip()}")
            
    # Update page number to '5' if it's copied
    # The page number was Shape 11 on Slide 7
    if len(slide_5.shapes) > 11:
        s = slide_5.shapes[11]
        if hasattr(s, 'text') and s.text.strip() == "7":
            s.text = "5"
            print("Successfully updated page number to 5!")

    # Let's save to a temp file to verify first
    prs.save(r"C:\Users\32022\.gemini\antigravity\brain\d0acdc8b-e6fd-413b-a1c3-b9364460abb5\scratch\test_copy.pptx")
    print("Saved test copy successfully!")
except Exception as e:
    print(f"Error: {e}")
