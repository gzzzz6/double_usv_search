from pptx import Presentation
prs = Presentation(r"E:\毕设PPT\USV_PPT_gemini.pptx")
slide = prs.slides[4] # Slide 5 is index 4
print("=== Slide 5 (Index 4) Shapes ===")
for i, shape in enumerate(slide.shapes):
    print(f"Shape {i}: '{shape.name}', type={shape.shape_type}")
    if hasattr(shape, 'text') and shape.text.strip():
        print(f"  Text: {shape.text.strip()}")
    try:
        if shape.fill and shape.fill.type == 1: # Solid
            print(f"  Fill color: {shape.fill.fore_color.rgb if shape.fill.fore_color.type == 1 else shape.fill.fore_color}")
    except AttributeError:
        pass
    try:
        if shape.line:
            print(f"  Line color: {shape.line.color.rgb if shape.line.color.type == 1 else shape.line.color}")
    except AttributeError:
        pass
    print(f"  Pos: left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}")
