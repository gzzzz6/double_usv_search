from pptx import Presentation
prs = Presentation(r"E:\毕设PPT\USV_PPT.pptx")
found = False
for i, slide in enumerate(prs.slides):
    texts = []
    for s in slide.shapes:
        if hasattr(s, 'text') and s.text.strip():
            texts.append(s.text.strip())
    full_text = " ".join(texts)
    if "信息场驱动的滚动搜索决策" in full_text or "reservation_v1" in full_text or "recency_bias" in full_text:
        print(f"FOUND our slide on Slide {i+1}!")
        found = True
if not found:
    print("Our custom slide was NOT found in the PPT. It seems the file was replaced or restored.")
