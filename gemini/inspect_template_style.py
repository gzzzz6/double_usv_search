import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

prs = Presentation(r"E:\毕设PPT\USV_PPT_gemini.pptx")
slide5 = prs.slides[4]  # Index 4 = Slide 5 (the reference slide with content)

print(f"=== Slide 5 Shape Details ===")
for i, shape in enumerate(slide5.shapes[:30]):  # First 30 shapes
    print(f"\nShape {i}: '{shape.name}', type={shape.shape_type}")
    print(f"  Pos: L={shape.left}, T={shape.top}, W={shape.width}, H={shape.height}")
    if hasattr(shape, 'text') and shape.text.strip():
        print(f"  Text: {repr(shape.text.strip()[:60])}")
    try:
        if shape.fill and shape.fill.type:
            fc = shape.fill.fore_color
            if fc.type == 1:
                print(f"  Fill: #{shape.fill.fore_color.rgb}")
            elif fc.type == 2:
                print(f"  Fill: theme color idx={fc.theme_color}")
    except Exception:
        pass
    try:
        if shape.line:
            lc = shape.line.color
            if lc.type == 1:
                print(f"  Line: #{shape.line.color.rgb}")
            elif lc.type == 2:
                print(f"  Line: theme color idx={lc.theme_color}")
            w = shape.line.width
            if w:
                print(f"  Line width: {w} ({w/12700:.1f}pt)")
    except Exception:
        pass
    try:
        if hasattr(shape, 'text_frame'):
            tf = shape.text_frame
            for para in tf.paragraphs[:1]:
                run = para.runs[0] if para.runs else None
                if run:
                    print(f"  Font: name={run.font.name}, size={run.font.size}, bold={run.font.bold}")
                    if run.font.color.type == 1:
                        print(f"  FontColor: #{run.font.color.rgb}")
    except Exception:
        pass
